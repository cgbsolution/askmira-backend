"""Salaries in bulk from the Employee Salary Master sheet: a template, a check, then all at once.

The sheet is laid out the way HR already keeps it (the "EmployeeSalaryMaster"
template of the HRMS the user moved from): Employee Code, WEF, Pay Basis,
Replace all future salary revisions, CTC TOTAL, MONTHLY GROSS, then one column
per wage type. A file exported from that system uploads without editing.

How a row is read:

* Pay Basis (Yearly, Monthly, Weekly, Daily or Hourly) applies to every amount
  on the row.
* A row with amounts is taken exactly as typed, the way the Salary master's
  editor saves: each wage type given is paid at that amount, one left empty is
  not paid, and CTC TOTAL is kept as given. The CTC also holds employer PF, ESI
  and admin charges, which payroll adds on top, so the wage types need not add
  up to it; the difference is shown on the row.
* A row with only CTC TOTAL is split by the salary structure.
* The statutory MANUAL columns (TDS, PROF. TAX, MPF, ESIC, LWF) are deducted
  instead of what payroll would work out, for that person; empty, payroll works
  them out as usual.
* MONTHLY GROSS is a check: a different gross is noted on the row.
* "Replace all future salary revisions" set to Yes removes the person's
  revisions from WEF on (only ones payroll has not used) and puts this one in
  their place. Otherwise WEF must be after the latest revision.

An upload is checked first and nothing is saved. Saving takes every ready row
in one transaction, and only when no row has an error.
"""

from __future__ import annotations

import contextlib
import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any

from fastapi import HTTPException, Request, status
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import tenant_scoped
from app.models.audit import AuditAction
from app.models.employee import Employee, EmploymentStatus
from app.models.payroll import (
    ComponentType,
    EmployeeSalary,
    SalaryComponent,
    SalaryStructure,
    WageGroup,
    wage_group_of,
)
from app.models.user import User
from app.schemas.payroll import (
    EmployeeSalaryCreate,
    SalaryComponentAmountInput,
    SalaryImportRead,
    SalaryImportRowRead,
)
from app.services import audit
from app.services.employee_import.fields import read_date, text
from app.services.payroll.salary_master import revision_locks, revision_snapshot
from app.services.payroll.service import (
    assign_salary,
    get_structure,
    resolve_amounts,
    to_period_amounts,
)
from app.services.storage import StorageError, get_storage

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000

# ── The columns ─────────────────────────────────────────────────────────────

CODE = "Employee Code"
WEF = "WEF"
BASIS = "Pay Basis"
REPLACE = "Replace all future salary revisions"
CTC = "CTC TOTAL"
GROSS = "MONTHLY GROSS"
STRUCTURE = "Salary Structure"
REASON = "Reason"
NAME = "Employee Name"

#: The Employee Salary Master layout, column for column. Wage-type headers are
#: matched to the workspace's own wage types by name (see `_OLD_NAMES`).
LAYOUT: tuple[str, ...] = (
    CODE, WEF, BASIS, REPLACE, CTC, GROSS,
    "GRATUITY AMOUNT MANUAL", "PROF TAX CTC",
    "BASIC SALARY", "HOUSE RENT ALLOWANCE", "CONVEYANCE ALLOWANCE", "SPECIAL ALLOWANCE",
    "EDUCATION ALLOWANCE", "BONUS", "LEAVE ENCASHMENT", "MEDICAL REIMBURSEMENT",
    "FUEL REIMBURSEMENT", "BONUS PROVISION", "GRATUITY PROVISION",
    "TDS MANUAL", "PROF. TAX MANUAL", "MPF CONTRIBUTION MANUAL",
    "ESIC EMPLOYEE CONTRIBUTION MANUAL", "LWF MANUAL",
    "INSURANCE", "Uniform Allowance", "Telephone Allowance",
)  # fmt: skip


def _key(header: Any) -> str:
    """A header compared loosely: case, punctuation and spacing do not matter."""
    return re.sub(r"[^a-z0-9]+", " ", str(header or "").lower()).strip()


#: Other ways the fixed columns are written.
_FIXED: dict[str, str] = {
    "employee code": CODE, "emp code": CODE, "empcode": CODE, "code": CODE,
    "employee no": CODE, "emp no": CODE,
    "wef": WEF, "w e f": WEF, "with effect from": WEF, "effective from": WEF,
    "effective date": WEF,
    "pay basis": BASIS, "basis": BASIS,
    "replace all future salary revisions": REPLACE, "replace future revisions": REPLACE,
    "replace future salary revisions": REPLACE,
    "ctc total": CTC, "ctc": CTC, "total ctc": CTC,
    "monthly gross": GROSS, "gross": GROSS, "gross salary": GROSS,
    "salary structure": STRUCTURE, "structure": STRUCTURE,
    "reason": REASON, "remarks": REASON,
    "employee name": NAME, "name": NAME,
}  # fmt: skip

#: The old system's wage-type names, mapped to AskMira's standard codes.
_OLD_NAMES: dict[str, str] = {
    "basic salary": "BASIC", "basic": "BASIC",
    "house rent allowance": "HRA", "hra": "HRA",
    "conveyance allowance": "CONVEYANCE", "special allowance": "SPECIAL",
    "education allowance": "EDUCATION", "bonus": "BONUS",
    "leave encashment": "LEAVE_ENCASH", "medical reimbursement": "MEDICAL",
    "fuel reimbursement": "FUEL", "bonus provision": "BONUS_PROVISION",
    "gratuity amount manual": "GRATUITY", "gratuity": "GRATUITY",
    "gratuity provision": "GRATUITY_PROVISION", "insurance": "INSURANCE",
    "uniform allowance": "UNIFORM", "telephone allowance": "TELEPHONE",
    "tds manual": "TDS_MANUAL", "prof tax manual": "PT_MANUAL",
    "professional tax manual": "PT_MANUAL", "mpf contribution manual": "PF_MANUAL",
    "pf contribution manual": "PF_MANUAL", "esic employee contribution manual": "ESIC_MANUAL",
    "lwf manual": "LWF_MANUAL", "prof tax ctc": "PT_CTC",
}  # fmt: skip

#: Template columns a workspace set up before they were added lacks.
_TEMPLATE_ONLY = frozenset(
    key for key, code in _OLD_NAMES.items()
    if code in {"TDS_MANUAL", "PT_MANUAL", "PF_MANUAL", "ESIC_MANUAL", "LWF_MANUAL", "PT_CTC",
                "GRATUITY_PROVISION"}
)  # fmt: skip

_HEADER_FILL = PatternFill("solid", fgColor="EB061D")  # AskMira Red
_REQUIRED_FILL = PatternFill("solid", fgColor="F8A01A")  # Marigold
_INFO_FILL = PatternFill("solid", fgColor="E9E1DC")  # Cream, darker: not read
_GROUP_ORDER = {group: index for index, group in enumerate(WageGroup)}


def _unprocessable(detail: str) -> HTTPException:
    return HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=detail)


async def _components(session: AsyncSession, user: User) -> list[SalaryComponent]:
    rows = (
        await session.scalars(tenant_scoped(select(SalaryComponent), SalaryComponent, user))
    ).all()
    return sorted(
        rows,
        key=lambda c: (
            _GROUP_ORDER[wage_group_of(c.component_type, c.wage_group)],
            c.display_order,
            c.code,
        ),
    )


def _wage_code(header: Any, components: list[SalaryComponent]) -> str | None:
    """The workspace wage type a header names: by its name, its code, or the old name."""
    key = _key(header)
    by_name = {_key(c.name): c.code for c in components}
    codes = {c.code for c in components}
    if key in by_name:
        return by_name[key]
    as_code = key.upper().replace(" ", "_")
    if as_code in codes:
        return as_code
    legacy = re.search(r"\[([A-Z][A-Z0-9_]*)\]", str(header or ""))
    if legacy and legacy.group(1) in codes:
        return legacy.group(1)
    old = _OLD_NAMES.get(key)
    return old if old in codes else None


# ── The template ────────────────────────────────────────────────────────────


async def template(session: AsyncSession, user: User) -> bytes:
    """The Employee Salary Master sheet, with this workspace's people listed."""
    components = await _components(session, user)
    active = [c for c in components if c.active]
    structures = (
        await session.scalars(
            tenant_scoped(
                select(SalaryStructure)
                .where(SalaryStructure.active.is_(True))
                .order_by(SalaryStructure.name),
                SalaryStructure,
                user,
            )
        )
    ).all()
    people = (
        await session.scalars(
            tenant_scoped(
                select(Employee)
                .where(Employee.deleted_at.is_(None), Employee.status != EmploymentStatus.EXITED)
                .order_by(Employee.employee_code),
                Employee,
                user,
            )
        )
    ).all()
    current: dict[uuid.UUID, EmployeeSalary] = {}
    if people:
        for salary in (
            await session.scalars(
                tenant_scoped(
                    select(EmployeeSalary).where(
                        EmployeeSalary.employee_id.in_([p.id for p in people]),
                        EmployeeSalary.effective_to.is_(None),
                    ),
                    EmployeeSalary,
                    user,
                )
            )
        ).all():
            known = current.get(salary.employee_id)
            if known is None or salary.effective_from > known.effective_from:
                current[salary.employee_id] = salary

    covered = {_wage_code(h, active) for h in LAYOUT} - {None}
    extras = [c.name.upper() for c in active if c.code not in covered]
    headers = [*LAYOUT, *extras, STRUCTURE, REASON, NAME]
    required = {CODE, WEF, BASIS, CTC}
    not_read = {_key(NAME)}

    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    _header_row(sheet, headers, required, not_read)

    for person in people:
        salary = current.get(person.id)
        row = [None] * len(headers)
        row[0] = person.employee_code
        if salary and salary.structure and salary.structure.active:
            row[headers.index(STRUCTURE)] = salary.structure.name
        row[headers.index(NAME)] = person.full_name
        sheet.append(row)

    last = max(len(people) + 1, 2) + 500
    for row in range(2, last + 1):
        sheet.cell(row=row, column=headers.index(WEF) + 1).number_format = "DD-MMM-YYYY"

    def choices(column: str, values: list[str]) -> None:
        listed = ",".join(v.replace(",", " ") for v in values)
        if not values or len(listed) >= 250:
            return
        validation = DataValidation(
            type="list", formula1=f'"{listed}"', allow_blank=True, showErrorMessage=False
        )
        letter = get_column_letter(headers.index(column) + 1)
        validation.add(f"{letter}2:{letter}{last}")
        sheet.add_data_validation(validation)

    choices(BASIS, ["Yearly", "Monthly", "Weekly", "Daily", "Hourly"])
    choices(REPLACE, ["Yes", "No"])
    choices(STRUCTURE, [s.name for s in structures])

    sample = book.create_sheet("Sample")
    sample.append(headers)
    example = {
        CODE: "EMP001", WEF: date(2026, 4, 1), BASIS: "Monthly", REPLACE: "No",
        CTC: Decimal("51666.67"), "BASIC SALARY": 25833, "HOUSE RENT ALLOWANCE": 10334,
        "GRATUITY AMOUNT MANUAL": 1242, "INSURANCE": 239, REASON: "Annual increment",
    }  # fmt: skip
    sample.append([example.get(h) for h in headers])
    sample.cell(row=2, column=headers.index(WEF) + 1).number_format = "DD-MMM-YYYY"
    for index, header in enumerate(headers, start=1):
        sample.cell(row=1, column=index).font = Font(bold=True)
        sample.column_dimensions[get_column_letter(index)].width = max(12, min(len(header) + 3, 28))

    notes = book.create_sheet("Instructions")
    lines = [
        "Employee Salary Master: how to fill Sheet1",
        "",
        "One row is one salary revision. Rows without CTC TOTAL or any amount are skipped, so",
        "the listed people can stay: fill only the rows you want to change.",
        "",
        f"{CODE} (required): as in AskMira.",
        f"{WEF} (required): the date the salary applies from, e.g. 01-Apr-2026.",
        f"{BASIS} (required): Yearly, Monthly, Weekly, Daily or Hourly. Every amount on the",
        "    row is read in it (a week is 1/52 of a year; a month has 26 working days of",
        "    8 hours).",
        f"{REPLACE}: Yes removes the person's revisions from {WEF} on and puts this one in their",
        f"    place (only ones payroll has not used). No or empty: {WEF} must be after the latest",
        "    revision.",
        f"{CTC}: the CTC, kept as given.",
        "Wage-type columns: the amounts, in the row's Pay Basis. A row with amounts is saved",
        "    exactly as typed; a wage type left empty is not paid. A row with only CTC TOTAL",
        "    is split by the salary structure.",
        "TDS / PROF. TAX / MPF / ESIC / LWF MANUAL: deducted instead of what payroll works out;",
        "    leave empty to let payroll work it out.",
        f"{GROSS}: kept as typed; the row notes it when the earnings add up to another figure.",
        f"{STRUCTURE}: empty keeps the person's present structure (or the workspace's only one).",
        f"{REASON}: why the salary changes. {NAME} is for reference and is not read.",
        "",
        "Upload on Payroll > Salary master > Bulk upload. Every row is checked and shown first;",
        "nothing is saved until you confirm, and nothing at all while any row has an error.",
        "",
        "Structures in this workspace: " + (", ".join(s.name for s in structures) or "none yet"),
        "See the Sample sheet for a filled row.",
    ]
    for line in lines:
        notes.append([line])
    notes["A1"].font = Font(bold=True, size=13)
    notes.column_dimensions["A"].width = 100

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _header_row(sheet: Any, headers: list[str], required: set[str], not_read: set[str]) -> None:
    """The Employee Salary Master's header row: required columns marigold, columns
    an upload passes over cream, the rest red."""
    sheet.append(headers)
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        if header in required:
            cell.fill, cell.font = _REQUIRED_FILL, Font(bold=True, color="100404")
        elif _key(header) in not_read:
            cell.fill, cell.font = _INFO_FILL, Font(bold=True, color="55524B")
        elif _OLD_NAMES.get(_key(header), "").endswith("_MANUAL"):
            cell.fill, cell.font = _HEADER_FILL, Font(bold=True, color="FFFFFF")
            cell.comment = Comment(
                "Deducted instead of what payroll works out, for this person. "
                "Leave empty to let payroll work it out.",
                "AskMira",
            )
        else:
            cell.fill, cell.font = _HEADER_FILL, Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = max(12, min(len(header) + 3, 28))
    sheet.row_dimensions[1].height = 42
    sheet.freeze_panes = "C2"


# ── Exporting the present salaries ──────────────────────────────────────────


def _present(series: list[EmployeeSalary], today: date) -> EmployeeSalary | None:
    """The revision in force today; the latest one when all start later."""
    in_force = [
        s
        for s in series
        if s.effective_from <= today and (s.effective_to is None or s.effective_to >= today)
    ]
    if in_force:
        return max(in_force, key=lambda s: s.effective_from)
    return max(series, key=lambda s: s.effective_from) if series else None


#: The Employee Salary Master as the old HRMS gives it out, so a downloaded sheet
#: looks exactly like theirs: plain bold headers, these notes on the first four
#: (the other names each column is known by), its column widths.
_MASTER_NOTES: dict[str, str] = {
    CODE: "Employee,EmpCode,Emp Code,Code,Emp-Code,EmployeeCode",
    WEF: "Date,From,Revision Date,Salary Revision Date,EffectiveStartdate,Effective Start date",
    BASIS: "RateUnit,PayBasis,Unit,Frequency",
    REPLACE: "Update all future entries",
}
_MASTER_WIDTHS: tuple[float, ...] = (
    16.39, 20.0, 10.61, 34.21, 11.94, 18.65, 30.8, 15.27, 15.19, 27.03, 27.92, 22.24, 26.2,
    9.14, 21.96, 28.47, 24.27, 20.09, 22.99, 14.74, 21.28, 31.14, 41.93, 15.21, 12.95,
    20.26, 22.37,
)  # fmt: skip
_MASTER_DATE = "dd-MMM-yy"
_SAMPLE_FILL = PatternFill("solid", fgColor="90EE90")


def _master_header(sheet: Any, headers: list[str]) -> None:
    """The Employee Salary Master's header row, styled as the old HRMS styles it."""
    sheet.append(headers)
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(horizontal="center")
        if header in _MASTER_NOTES:
            note = Comment(_MASTER_NOTES[header], "AskMira")
            note.width, note.height = 144, 79
            cell.comment = note
        sheet.column_dimensions[get_column_letter(index)].width = (
            _MASTER_WIDTHS[index - 1] if index <= len(_MASTER_WIDTHS) else max(12, len(header) + 4)
        )
    sheet.freeze_panes = "A2"


def _row_basis(salary: EmployeeSalary, gross: tuple[Decimal, Decimal]) -> str:
    """Monthly when every amount on the row is a whole-paisa twelfth of its year,
    so the sheet uploads back to the very same figures; Yearly otherwise."""
    ctc = Decimal(str(salary.ctc_annual))
    pairs = [(ctc, (ctc / 12).quantize(Decimal("0.01"))), gross]
    pairs += [
        (Decimal(str(c.annual_amount)), Decimal(str(c.monthly_amount))) for c in salary.components
    ]
    return "Monthly" if all(monthly * 12 == annual for annual, monthly in pairs) else "Yearly"


async def export_present(
    session: AsyncSession,
    user: User,
    employee_ids: list[uuid.UUID] | None,
    request: Request | None = None,
) -> tuple[bytes, int]:
    """Everyone's (or the chosen people's) present salary, as the Employee Salary Master.

    The sheet is the user's own EmployeeSalaryMaster file, column for column and
    styled the same, with one row per person filled from the salary in force
    today (the latest one when all start later). It uploads back unchanged, so a
    revision is: export, change WEF and the amounts, upload. A person with no
    salary gets a row with only the code. A wage type the workspace added beyond
    the sheet's columns gets a column after them. Audited as an export.
    """
    components = await _components(session, user)
    active = [c for c in components if c.active]
    stmt = select(Employee).where(Employee.deleted_at.is_(None))
    if employee_ids is not None:
        stmt = stmt.where(Employee.id.in_(employee_ids))
    people = (
        await session.scalars(tenant_scoped(stmt.order_by(Employee.employee_code), Employee, user))
    ).all()
    series: dict[uuid.UUID, list[EmployeeSalary]] = {}
    if people:
        for salary in (
            await session.scalars(
                tenant_scoped(
                    select(EmployeeSalary).where(
                        EmployeeSalary.employee_id.in_([p.id for p in people])
                    ),
                    EmployeeSalary,
                    user,
                )
            )
        ).all():
            series.setdefault(salary.employee_id, []).append(salary)

    covered = {_wage_code(h, components) for h in LAYOUT} - {None}
    extras = [c.name.upper() for c in active if c.code not in covered]
    headers = [*LAYOUT, *extras]
    column_code = {h: _wage_code(h, components) for h in headers}

    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    _master_header(sheet, headers)
    sheet.cell(row=1, column=headers.index(WEF) + 1).number_format = _MASTER_DATE

    today = date.today()
    for person in people:
        row: dict[str, Any] = {CODE: person.employee_code}
        salary = _present(series.get(person.id, []), today)
        if salary is not None:
            earnings = [c for c in salary.components if c.component_type is ComponentType.EARNING]
            if salary.gross_amount is not None:
                gross = to_period_amounts(
                    Decimal(str(salary.gross_amount)), salary.gross_basis or "monthly"
                )
            else:
                gross = (
                    sum((Decimal(str(c.annual_amount)) for c in earnings), Decimal(0)),
                    sum((Decimal(str(c.monthly_amount)) for c in earnings), Decimal(0)),
                )
            basis = _row_basis(salary, gross)
            monthly = basis == "Monthly"
            ctc_annual = Decimal(str(salary.ctc_annual))
            row.update(
                {
                    WEF: salary.effective_from,
                    BASIS: basis,
                    CTC: (ctc_annual / 12).quantize(Decimal("0.01")) if monthly else ctc_annual,
                    GROSS: gross[1] if monthly else gross[0],
                }
            )
            amounts = {
                c.component_code: Decimal(str(c.monthly_amount if monthly else c.annual_amount))
                for c in salary.components
            }
            for header, code in column_code.items():
                if code and code in amounts:
                    row[header] = amounts[code]
        sheet.append([row.get(h) for h in headers])
        sheet.cell(row=sheet.max_row, column=headers.index(WEF) + 1).number_format = _MASTER_DATE

    wef = get_column_letter(headers.index(WEF) + 1)
    dates = DataValidation(
        type="date",
        formula1="1",
        formula2="730486",
        allow_blank=False,
        showInputMessage=True,
        showErrorMessage=True,
        errorTitle="An invalid date was entered",
        error="Enter a valid date",
        prompt="Enter a date",
    )
    dates.add(f"{wef}2:{wef}1048576")
    sheet.add_data_validation(dates)
    pay = get_column_letter(headers.index(BASIS) + 1)
    bases = DataValidation(
        type="list",
        formula1='"Yearly,Monthly,Weekly,Daily,Hourly"',
        allow_blank=False,
        showInputMessage=True,
        showErrorMessage=True,
        errorTitle="An invalid Pay Basis was entered",
        error="Select a Pay Basis from the list",
        prompt="Select or enter a Pay Basis from the list",
    )
    bases.add(f"{pay}2:{pay}1048576")
    sheet.add_data_validation(bases)

    sample = book.create_sheet("Sample")
    _master_header(sample, headers)
    sample.append([None if h in (CODE, WEF, BASIS, REPLACE) else h for h in headers])
    for index in range(1, len(headers) + 1):
        cell = sample.cell(row=2, column=index)
        cell.fill = _SAMPLE_FILL
        cell.alignment = Alignment(wrap_text=True)
    sample.column_dimensions[wef].width = 9.14

    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.EXPORT,
        entity_type="salary_export",
        actor=user,
        changes={
            "employees": len(people),
            "chosen": employee_ids is not None,
        },
        summary=f"Exported the present salary of {len(people)} employee(s)",
        request=request,
    )
    await session.commit()
    out = io.BytesIO()
    book.save(out)
    return out.getvalue(), len(people)


# ── Reading an upload ───────────────────────────────────────────────────────


def _rows(data: bytes, filename: str) -> list[tuple[int, list[Any]]]:
    """Every row of the sheet to read, numbered as Excel does.

    The first sheet with an Employee Code header, passing over a Sample or an
    Instructions sheet.
    """
    if len(data) > MAX_BYTES:
        raise _unprocessable("The file is larger than 5 MB.")
    if filename.lower().endswith(".csv"):
        try:
            decoded = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = data.decode("latin-1")
        return list(enumerate(csv.reader(io.StringIO(decoded)), start=1))
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds for a bad file
        raise _unprocessable("This is not an Excel (.xlsx) or CSV file AskMira can read.") from exc
    for sheet in book.worksheets:
        if _key(sheet.title) in {"sample", "instructions"}:
            continue
        rows = [
            (number, list(values))
            for number, values in enumerate(sheet.iter_rows(values_only=True), start=1)
        ]
        if any(_FIXED.get(_key(v)) == CODE for _, values in rows[:15] for v in values):
            return rows
    raise _unprocessable(f"No sheet has an '{CODE}' column. Start from the template.")


def _money(raw: Any, what: str) -> Decimal | None:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if isinstance(raw, bool):
        raise ValueError(f"{what}: '{raw}' is not an amount")
    cleaned = str(raw).strip().replace(",", "").replace("₹", "").strip()
    try:
        amount = Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"{what}: '{raw}' is not an amount") from exc
    if amount < 0:
        raise ValueError(f"{what} cannot be negative")
    if amount != amount.quantize(Decimal("0.01")):
        raise ValueError(f"{what}: {cleaned} has more than two decimals")
    return amount


_BASES = {
    "annual": {"yearly", "annual", "annually", "year", "per annum", "pa"},
    "monthly": {"monthly", "month", "pm", "per month"},
    "weekly": {"weekly", "week", "per week"},
    "daily": {"daily", "day", "per day"},
    "hourly": {"hourly", "hour", "per hour"},
}


def _basis(raw: Any) -> str | None:
    value = _key(raw)
    if not value:
        return None
    for basis, words in _BASES.items():
        if value in words:
            return basis
    raise ValueError(f"{BASIS} '{raw}': use Yearly, Monthly, Weekly, Daily or Hourly")


def _yes(raw: Any) -> bool:
    return _key(raw) in {"yes", "y", "true", "1"}


def _rupees(amount: Decimal) -> str:
    whole, _, paise = f"{amount:.2f}".partition(".")
    grouped = whole if len(whole) <= 3 else f"{_indian(whole[:-3])},{whole[-3:]}"
    return f"₹{grouped}" + ("" if paise == "00" else f".{paise}")


def _indian(digits: str) -> str:
    parts = []
    while len(digits) > 2:
        parts.insert(0, digits[-2:])
        digits = digits[:-2]
    if digits:
        parts.insert(0, digits)
    return ",".join(parts)


@dataclass
class _Plan:
    employee: Employee
    payload: EmployeeSalaryCreate
    replace: bool


@dataclass
class _Outcome:
    rows: list[SalaryImportRowRead] = field(default_factory=list)
    plans: list[_Plan] = field(default_factory=list)


async def _check(session: AsyncSession, user: User, data: bytes, filename: str) -> _Outcome:
    table = _rows(data, filename)
    header_at = next(
        (
            index
            for index, (_, values) in enumerate(table[:15])
            if any(_FIXED.get(_key(v)) == CODE for v in values)
        ),
        None,
    )
    if header_at is None:
        raise _unprocessable(f"No '{CODE}' column found. Start from the template.")
    header_values = table[header_at][1]
    body = table[header_at + 1 :]
    if len(body) > MAX_ROWS:
        raise _unprocessable(f"A sheet can hold up to {MAX_ROWS} rows.")

    components = await _components(session, user)
    by_code = {c.code: c for c in components}
    fixed: dict[str, int] = {}
    wage_columns: list[tuple[int, str, str]] = []  # (index, header, code)
    unknown_columns: list[tuple[int, str]] = []
    for index, header in enumerate(header_values):
        key = _key(header)
        if not key:
            continue
        if key in _FIXED and _FIXED[key] not in fixed:
            fixed[_FIXED[key]] = index
        elif (code := _wage_code(header, components)) is not None:
            wage_columns.append((index, str(header), code))
        else:
            unknown_columns.append((index, str(header)))

    def cell(values: list[Any], name: str) -> Any:
        index = fixed.get(name)
        return values[index] if index is not None and index < len(values) else None

    def at(values: list[Any], index: int) -> Any:
        return values[index] if index < len(values) else None

    codes = {code for _, values in body if (code := text(cell(values, CODE))) is not None}
    employees: dict[str, Employee] = {}
    if codes:
        employees = {
            e.employee_code.upper(): e
            for e in (
                await session.scalars(
                    tenant_scoped(
                        select(Employee).where(
                            Employee.deleted_at.is_(None),
                            func.upper(Employee.employee_code).in_([c.upper() for c in codes]),
                        ),
                        Employee,
                        user,
                    )
                )
            ).all()
        }
    structures = {
        _key(s.name): s
        for s in (
            await session.scalars(
                tenant_scoped(
                    select(SalaryStructure).where(SalaryStructure.active.is_(True)),
                    SalaryStructure,
                    user,
                )
            )
        ).all()
    }

    # Each person's revisions as dates, so rows can be checked in order, with
    # the ones payroll has used marked: those can never be replaced.
    series: dict[uuid.UUID, list[EmployeeSalary]] = {}
    if employees:
        for salary in (
            await session.scalars(
                tenant_scoped(
                    select(EmployeeSalary)
                    .where(EmployeeSalary.employee_id.in_([e.id for e in employees.values()]))
                    .order_by(EmployeeSalary.effective_from),
                    EmployeeSalary,
                    user,
                )
            )
        ).all():
            series.setdefault(salary.employee_id, []).append(salary)
    locks = await revision_locks(session, user, [s.id for ss in series.values() for s in ss])
    dates: dict[uuid.UUID, list[date]] = {
        k: [s.effective_from for s in v] for k, v in series.items()
    }
    locked: dict[uuid.UUID, set[date]] = {
        k: {s.effective_from for s in v if s.id in locks} for k, v in series.items()
    }

    outcome = _Outcome()
    for number, values in body:
        if all(text(v) is None for v in values):
            continue
        code = text(cell(values, CODE))
        employee = employees.get(code.upper()) if code else None
        row = SalaryImportRowRead(
            row=number,
            employee_code=code,
            employee_name=employee.full_name if employee else text(cell(values, NAME)),
            effective_from=None,
            structure_name=None,
            ctc_annual=None,
            gross_monthly=None,
            status="error",
        )
        outcome.rows.append(row)
        notes: list[str] = []
        try:
            ctc = _money(cell(values, CTC), CTC)
            given: dict[str, tuple[str, Decimal]] = {}
            for index, header, wage_code in wage_columns:
                amount = _money(at(values, index), header)
                if amount is None:
                    continue
                if wage_code in given and given[wage_code][1] != amount:
                    raise ValueError(
                        f"{given[wage_code][0]} and {header} are both {by_code[wage_code].name} "
                        "but hold different amounts"
                    )
                given[wage_code] = (header, amount)
            for index, header in unknown_columns:
                if text(at(values, index)) is not None:
                    if _key(header) in _TEMPLATE_ONLY:
                        raise ValueError(
                            f"'{header}' is not set up in this workspace yet. Press "
                            "'Update wage types' on the Salary master, then upload again."
                        )
                    raise ValueError(
                        f"'{header}' is not a wage type of this workspace. Add it under "
                        "Payroll setup → Salary structures, or clear the column."
                    )
            if ctc is None and not given:
                row.status, row.message = "skipped", "No CTC or amount on this row: left as it is."
                continue
            if code is None:
                raise ValueError(f"{CODE} is empty")
            if employee is None:
                raise ValueError(f"No employee with code {code}")
            basis = _basis(cell(values, BASIS))
            if basis is None:
                raise ValueError(f"{BASIS} is empty: Yearly, Monthly, Weekly, Daily or Hourly")
            on = read_date(cell(values, WEF))
            if on is None:
                raise ValueError(f"{WEF} is empty")
            row.effective_from = on

            # ── Where the revision goes ─────────────────────────────────────
            replace = _yes(cell(values, REPLACE))
            person_dates = dates.setdefault(employee.id, [])
            if replace:
                later = sorted(d for d in person_dates if d >= on)
                stuck = [d for d in later if d in locked.get(employee.id, set())]
                if stuck:
                    raise ValueError(
                        f"The revision from {stuck[0]:%d-%b-%Y} cannot be replaced: payroll "
                        "has been approved on it"
                    )
                if later:
                    notes.append(f"Replaces {len(later)} revision(s) from {later[0]:%d-%b-%Y} on")
            elif person_dates and on <= max(person_dates):
                raise ValueError(
                    f"{WEF} must be after the latest revision ({max(person_dates):%d-%b-%Y}). "
                    f"Put Yes in '{REPLACE}' to replace revisions from that date on."
                )

            # ── Which structure ─────────────────────────────────────────────
            named = text(cell(values, STRUCTURE))
            if named:
                structure = structures.get(_key(named))
                if structure is None:
                    raise ValueError(f"No active salary structure called '{named}'")
            else:
                history = series.get(employee.id, [])
                present = (
                    await get_structure(session, user, history[-1].structure_id)
                    if history
                    else None
                )
                if present is not None and present.active:
                    structure = present
                elif len(structures) == 1:
                    structure = next(iter(structures.values()))
                else:
                    raise ValueError(f"{STRUCTURE} is empty: name the salary structure to use")
            row.structure_name = structure.name

            # ── The amounts ─────────────────────────────────────────────────
            for wage_code in given:
                component = by_code[wage_code]
                if not component.active:
                    raise ValueError(f"{component.name} is switched off in payroll setup")
            reason = (text(cell(values, REASON)) or "Bulk upload")[:255]
            if given:
                # Amounts given: saved exactly as typed, like the Salary master editor.
                if ctc is not None:
                    ctc_fields: dict[str, Any] = {"ctc_amount": ctc, "ctc_basis": basis}
                else:
                    in_ctc = sum(
                        (
                            to_period_amounts(amount, basis)[0]
                            for c, (_, amount) in given.items()
                            if by_code[c].part_of_ctc
                        ),
                        Decimal(0),
                    )
                    if in_ctc <= 0:
                        raise ValueError(f"{CTC} is empty")
                    ctc_fields = {"ctc_annual": in_ctc}
                    notes.append(f"{CTC} empty: set to the wage types' total, {_rupees(in_ctc)}")
                gross_typed = _money(cell(values, GROSS), GROSS)
                payload = EmployeeSalaryCreate(
                    entry="manual",
                    gross_amount=gross_typed,
                    gross_basis=basis,
                    structure_id=structure.id,
                    effective_from=on,
                    revision_reason=reason,
                    components=[
                        SalaryComponentAmountInput(component_code=c, amount=amount, basis=basis)
                        for c, (_, amount) in given.items()
                    ],
                    **ctc_fields,
                )
            else:
                # Only a CTC: the structure splits it.
                assert ctc is not None
                payload = EmployeeSalaryCreate(
                    structure_id=structure.id,
                    effective_from=on,
                    revision_reason=reason,
                    ctc_amount=ctc,
                    ctc_basis=basis,
                )
            resolved = await resolve_amounts(session, user, structure, payload)
        except ValidationError as exc:
            row.message = exc.errors()[0].get("msg", "Invalid row").removeprefix("Value error, ")
            continue
        except HTTPException as exc:
            row.message = str(exc.detail)
            continue
        except ValueError as exc:
            row.message = str(exc)
            continue

        earnings = [
            c for c in resolved.breakdown.components if c.component_type is ComponentType.EARNING
        ]
        gross_monthly = sum((c.monthly_amount for c in earnings), Decimal(0))
        gross_given = None
        with contextlib.suppress(ValueError):
            gross_given = _money(cell(values, GROSS), GROSS)
        if gross_given is not None:
            typed_monthly = to_period_amounts(gross_given, basis)[1]
            if typed_monthly != gross_monthly:
                kept = " (kept as typed)" if payload.entry == "manual" else ""
                notes.append(
                    f"{GROSS} {_rupees(typed_monthly)} a month{kept}; the earnings add up to "
                    f"{_rupees(gross_monthly)}"
                )
        if payload.entry == "manual":
            in_ctc = sum(
                (c.annual_amount for c in resolved.breakdown.components if c.part_of_ctc),
                Decimal(0),
            )
            gap = resolved.ctc_annual - in_ctc
            if gap > 0:
                notes.append(
                    f"{CTC} {_rupees(resolved.ctc_annual)} a year; the wage types in it add up "
                    f"to {_rupees(in_ctc)}. {_rupees(gap)} is outside them (employer PF, ESI "
                    "and admin charges, which payroll adds)"
                )
            elif gap < 0:
                notes.append(
                    f"The wage types in CTC add up to {_rupees(in_ctc)} a year, "
                    f"{_rupees(-gap)} more than {CTC} {_rupees(resolved.ctc_annual)}"
                )
        row.status = "ready"
        row.ctc_annual = resolved.ctc_annual
        row.gross_monthly = gross_monthly
        row.typed = len(given)
        row.notes = notes
        if replace:
            dates[employee.id] = [d for d in dates[employee.id] if d < on] + [on]
        else:
            dates[employee.id].append(on)
        outcome.plans.append(_Plan(employee=employee, payload=payload, replace=replace))
    return outcome


async def _replace_from(
    session: AsyncSession,
    user: User,
    employee: Employee,
    on: date,
    request: Request | None,
) -> list[str]:
    """Remove a person's revisions from `on` (none payroll has used), audited.

    The revision before `on` is reopened, so the upload's revision follows it.
    Returns the storage keys of files the removed revisions carried.
    """
    rows = (
        await session.scalars(
            tenant_scoped(
                select(EmployeeSalary)
                .where(EmployeeSalary.employee_id == employee.id)
                .order_by(EmployeeSalary.effective_from),
                EmployeeSalary,
                user,
            )
        )
    ).all()
    removing = [s for s in rows if s.effective_from >= on]
    if not removing:
        return []
    locks = await revision_locks(session, user, [s.id for s in removing])
    if locks:
        raise _unprocessable(
            f"A revision of {employee.employee_code} was paid by payroll meanwhile; "
            "nothing was saved."
        )
    keys: list[str] = []
    for salary in removing:
        await audit.record(
            session,
            tenant_id=user.tenant_id,
            action=AuditAction.DELETE,
            entity_type="employee_salary",
            entity_id=salary.id,
            actor=user,
            changes={
                "employee_id": str(employee.id),
                "deleted": revision_snapshot(
                    salary, salary.structure.name if salary.structure else None
                ),
                "replaced_by_upload": True,
            },
            summary=(
                f"Salary revision from {salary.effective_from:%d %b %Y} replaced by a bulk "
                f"upload for {employee.employee_code}"
            ),
            request=request,
        )
        if salary.attachment_key:
            keys.append(salary.attachment_key)
        await session.delete(salary)
    await session.flush()
    before = [s for s in rows if s.effective_from < on]
    if before:
        before[-1].effective_to = None
        await session.flush()
    return keys


async def upload(
    session: AsyncSession,
    user: User,
    data: bytes,
    filename: str,
    *,
    apply: bool,
    request: Request | None = None,
) -> SalaryImportRead:
    """Check a filled sheet; with `apply`, save every ready row in one transaction."""
    outcome = await _check(session, user, data, filename)
    errors = sum(1 for r in outcome.rows if r.status == "error")
    result = SalaryImportRead(
        filename=filename,
        rows=outcome.rows,
        ready=len(outcome.plans),
        errors=errors,
        skipped=sum(1 for r in outcome.rows if r.status == "skipped"),
        noted=sum(1 for r in outcome.rows if r.status == "ready" and r.notes),
    )
    if not apply:
        return result
    if errors:
        raise _unprocessable(f"Fix the {errors} row(s) marked as errors first; nothing was saved.")
    if not outcome.plans:
        raise _unprocessable("No row has a CTC or an amount, so there is nothing to save.")

    stale: list[str] = []
    for plan in outcome.plans:
        if plan.replace:
            stale += await _replace_from(
                session, user, plan.employee, plan.payload.effective_from, request
            )
        await assign_salary(
            session, user, plan.employee.id, plan.payload, request=request, commit=False
        )
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.CREATE,
        entity_type="salary_import",
        actor=user,
        changes={
            "filename": filename,
            "revisions": len(outcome.plans),
            "employees": sorted({p.employee.employee_code for p in outcome.plans}),
        },
        summary=f"Bulk salary upload: {len(outcome.plans)} revision(s) from {filename}",
        request=request,
    )
    await session.commit()
    for key in stale:
        # The rows are gone; a file left behind is only storage, never data.
        with contextlib.suppress(StorageError):
            await get_storage().delete(key)
    result.applied = True
    return result
