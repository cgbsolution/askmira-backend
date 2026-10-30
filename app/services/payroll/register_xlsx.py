"""A run's salary register as an Excel workbook: the month's payroll data, per person.

One row per payslip, one column per wage type, the way an Indian salary
register reads: who (code, name, department, PAN, UAN, bank), days, each
earning, gross, each deduction, net pay, and what the employer pays on top.

Every row adds up exactly as the payslip does:

    net pay      = gross earnings + other earnings - total deductions
    deductions   = statutory columns + structure deductions + other deductions

"Other earnings" and "Other deductions" are what other modules put on the
payslip (claims, overtime, adjustments, loan recoveries); the second sheet
lists each of them. Taking a copy of everybody's pay is audited.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any

from fastapi import HTTPException, Request, status
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import tenant_scoped
from app.models.audit import AuditAction
from app.models.employee import Employee
from app.models.org import Department, LegalEntity, Location
from app.models.payroll import Payslip
from app.models.user import User
from app.services import audit
from app.services.payroll.payslips import (
    payslip_project_id,
    project_names,
    register_client_text,
    run_read,
)
from app.services.payroll.run import (
    CLOSED_STATUSES,
    get_run,
    imported_from,
    is_imported,
    other_earnings_of,
    period_label,
)

_ZERO = Decimal(0)
_MONEY = "#,##0.00"
_DATE = "dd-mmm-yyyy"

_TITLE_FONT = Font(bold=True, size=14, color="1F1B16")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="7D0712")  # AskMira maroon
_GROUP_FILL = PatternFill("solid", fgColor="F3ECE6")  # Cream
_TOTAL_FILL = PatternFill("solid", fgColor="F3ECE6")
_LOCKED_FONT = Font(bold=True, color="1E6B3A")
_UNLOCKED_FONT = Font(bold=True, color="9A5B00")
_THIN = Side(style="thin", color="D9CFC7")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

#: Statutory deductions, always shown, in the order a payslip lists them.
#: Their line codes are skipped when the structure's own deductions are read.
_STATUTORY_DEDUCTIONS: tuple[tuple[str, str], ...] = (
    ("pf_employee", "PF (employee)"),
    ("vpf", "VPF"),
    ("esi_employee", "ESI (employee)"),
    ("professional_tax", "Professional Tax"),
    ("lwf_employee", "LWF (employee)"),
    ("tds", "TDS (income tax)"),
)
_STATUTORY_DEDUCTION_CODES = frozenset({"PF_EMPLOYEE", "VPF", "ESI_EMPLOYEE", "PT", "LWF_EMPLOYEE"})
_STATUTORY_EMPLOYER: tuple[tuple[str, str], ...] = (
    ("pf_employer", "PF (employer)"),
    ("esi_employer", "ESI (employer)"),
    ("lwf_employer", "LWF (employer)"),
)
_STATUTORY_EMPLOYER_CODES = frozenset({"PF_EMPLOYER", "ESI_EMPLOYER", "LWF_EMPLOYER"})

_INPUT_KIND_LABEL = {
    "earning": "Earning",
    "reimbursement": "Reimbursement",
    "deduction": "Deduction",
    "recovery": "Recovery",
    "perquisite": "Perquisite (taxable, not paid in cash)",
    "employer_contribution": "Employer contribution",
}


def _dec(value: Any) -> Decimal:
    return Decimal(str(value)) if value is not None else _ZERO


def _text(cell: Cell, value: str | None) -> None:
    """Write people's text as text: openpyxl would treat a leading '=' as a formula."""
    cell.value = value or ""
    cell.data_type = "s"


def _ordered_lines(
    payslips: list[Payslip], kind: str, skip: frozenset[str]
) -> list[tuple[str, str]]:
    """(code, name) of every line of `kind` on the run, in the order they first appear."""
    seen: dict[str, str] = {}
    for payslip in payslips:
        for line in payslip.line_items or []:
            code = str(line.get("code") or "")
            if line.get("type") != kind or not code or code in skip or code in seen:
                continue
            seen[code] = str(line.get("name") or code)
    return list(seen.items())


def _line_amounts(payslip: Payslip, kind: str) -> dict[str, Decimal]:
    amounts: dict[str, Decimal] = {}
    for line in payslip.line_items or []:
        if line.get("type") == kind:
            code = str(line.get("code") or "")
            amounts[code] = amounts.get(code, _ZERO) + _dec(line.get("amount"))
    return amounts


async def _names_by_id(
    session: AsyncSession, user: User, model: Any, ids: set[uuid.UUID]
) -> dict[uuid.UUID, str]:
    ids = {i for i in ids if i is not None}
    if not ids:
        return {}
    rows = await session.execute(
        tenant_scoped(select(model.id, model.name).where(model.id.in_(ids)), model, user)
    )
    return {row_id: name for row_id, name in rows.all()}


def _widths(sheet: Worksheet, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


async def register_xlsx(
    session: AsyncSession, user: User, run_id: uuid.UUID, request: Request | None = None
) -> tuple[str, bytes]:
    run = await get_run(session, user, run_id)
    rows = list(
        (
            await session.execute(
                tenant_scoped(
                    select(Payslip, Employee)
                    .join(Employee, Payslip.employee_id == Employee.id)
                    .where(Payslip.run_id == run.id)
                    .order_by(Employee.employee_code),
                    Payslip,
                    user,
                )
            )
        ).all()
    )
    label = period_label(run.period_year, run.period_month)
    if not rows:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"The {label} run has no payslips yet. Compute it first, then download.",
        )

    summary = await run_read(session, user, run)
    imported = is_imported(run)
    entity = await session.scalar(
        tenant_scoped(
            select(LegalEntity).where(LegalEntity.id == run.legal_entity_id), LegalEntity, user
        )
    )
    employees = [employee for _, employee in rows]
    departments = await _names_by_id(
        session, user, Department, {e.department_id for e in employees}
    )
    locations = await _names_by_id(session, user, Location, {e.location_id for e in employees})
    projects = await project_names(
        session,
        user,
        {payslip_project_id(payslip, employee, imported=imported) for payslip, employee in rows},
    )

    def project_of(payslip: Payslip, employee: Employee) -> str | None:
        # An imported payslip names only the project its register row named; when
        # that matched no AskMira project, the register's own text is shown.
        project_id = payslip_project_id(payslip, employee, imported=imported)
        if project_id is not None:
            return projects.get(project_id)
        return register_client_text(payslip) if imported else None

    payslips = [payslip for payslip, _ in rows]

    earnings = _ordered_lines(payslips, "earning", frozenset())
    structure_deductions = _ordered_lines(payslips, "deduction", _STATUTORY_DEDUCTION_CODES)
    employer_lines = _ordered_lines(payslips, "employer_contribution", _STATUTORY_EMPLOYER_CODES)

    book = Workbook()
    sheet = book.active
    sheet.title = "Salary register"

    # ── Title block ─────────────────────────────────────────────────────────
    entity_name = entity.name if entity is not None else (summary.legal_entity_name or "")
    run_kind = (
        "" if run.run_type.value == "regular" else f" ({run.run_type.value.replace('_', '-')})"
    )
    sheet["A1"] = f"Salary register — {label}{run_kind}"
    sheet["A1"].font = _TITLE_FONT
    _text(sheet["A2"], entity_name)
    sheet["A2"].font = Font(bold=True)
    if imported:
        # The import time is not when the month was approved or paid: say what it is.
        source = summary.import_summary
        files = ", ".join(source.files) if source and source.files else ""
        when = f" on {source.imported_at:%d %b %Y}" if source and source.imported_at else ""
        who = f" by {source.imported_by_name}" if source and source.imported_by_name else ""
        sheet["A3"] = (
            f"Imported from {imported_from(run)} as paid there"
            f"{' (' + files + ')' if files else ''}; figures as in its register, never "
            f"recomputed. Locked on import{when}{who}."
        )
        sheet["A3"].font = _LOCKED_FONT
    elif run.status in CLOSED_STATUSES:
        when = f"{summary.approved_at:%d %b %Y}" if summary.approved_at else ""
        who = f" by {summary.approved_by_name}" if summary.approved_by_name else ""
        state = f"Locked: approved {when}{who}".strip()
        if run.paid_at:
            state += f"; paid {run.payment_date:%d %b %Y}" if run.payment_date else "; paid"
        sheet["A3"] = state + "."
        sheet["A3"].font = _LOCKED_FONT
    else:
        sheet["A3"] = (
            "Not locked: these figures can still change until the run is approved and locked."
        )
        sheet["A3"].font = _UNLOCKED_FONT
    sheet["A4"] = (
        f"{len(rows)} employee(s). Downloaded {datetime.now(UTC):%d %b %Y %H:%M} UTC"
        f" by {user.full_name}."
    )
    sheet["A4"].font = Font(italic=True, color="6B6259")

    # ── Columns ─────────────────────────────────────────────────────────────
    # (group, heading, kind) — kind drives the number format and the totals.
    columns: list[tuple[str, str, str]] = [
        ("Employee", "S.No", "int"),
        ("Employee", "Employee code", "text"),
        ("Employee", "Employee name", "text"),
        ("Employee", "Department", "text"),
        ("Employee", "Project", "text"),
        ("Employee", "Designation", "text"),
        ("Employee", "Location", "text"),
        ("Employee", "Date of joining", "date"),
        ("Employee", "PAN", "text"),
        ("Employee", "UAN", "text"),
        ("Employee", "ESI number", "text"),
        ("Employee", "Bank account", "text"),
        ("Employee", "IFSC", "text"),
        ("Days", "Working days", "days"),
        ("Days", "Paid days", "days"),
        ("Days", "LOP days", "days"),
        ("Days", "Not employed days", "days"),
    ]
    columns += [("Earnings", name, "money") for _, name in earnings]
    columns += [
        ("Earnings", "Gross earnings", "money"),
        ("Earnings", "Other earnings", "money"),
        ("Earnings", "Total earnings", "money"),
    ]
    columns += [("Deductions", name, "money") for _, name in _STATUTORY_DEDUCTIONS]
    columns += [("Deductions", name, "money") for _, name in structure_deductions]
    columns += [
        ("Deductions", "Other deductions", "money"),
        ("Deductions", "Total deductions", "money"),
        ("Net", "Net pay", "money"),
    ]
    columns += [("Employer", name, "money") for _, name in _STATUTORY_EMPLOYER]
    columns += [("Employer", name, "money") for _, name in employer_lines]
    columns += [("Employer", "Employer cost", "money")]

    group_row, head_row, first_data = 6, 7, 8
    # Group band over each run of columns, then the headings.
    start = 1
    for index in range(1, len(columns) + 2):
        group = columns[index - 1][0] if index <= len(columns) else None
        if group != columns[start - 1][0]:
            if index - 1 > start:
                sheet.merge_cells(
                    start_row=group_row, start_column=start, end_row=group_row, end_column=index - 1
                )
            cell = sheet.cell(row=group_row, column=start, value=columns[start - 1][0])
            cell.font, cell.fill = Font(bold=True, color="1F1B16"), _GROUP_FILL
            cell.alignment = Alignment(horizontal="center")
            start = index
    for index, (_, heading, kind) in enumerate(columns, start=1):
        cell = sheet.cell(row=head_row, column=index, value=heading)
        cell.font, cell.fill, cell.border = _HEADER_FONT, _HEADER_FILL, _BOX
        cell.alignment = Alignment(
            horizontal="right" if kind in ("money", "days") else "left",
            vertical="center",
            wrap_text=True,
        )
    sheet.row_dimensions[head_row].height = 32

    # ── Rows ────────────────────────────────────────────────────────────────
    totals: list[Decimal | None] = [_ZERO if kind == "money" else None for _, _, kind in columns]
    for offset, (payslip, employee) in enumerate(rows):
        earning_amounts = _line_amounts(payslip, "earning")
        deduction_amounts = _line_amounts(payslip, "deduction")
        employer_amounts = _line_amounts(payslip, "employer_contribution")

        gross = _dec(payslip.gross_earnings)
        other_earnings = other_earnings_of(payslip)
        statutory = [_dec(getattr(payslip, field)) for field, _ in _STATUTORY_DEDUCTIONS]
        structure = [deduction_amounts.get(code, _ZERO) for code, _ in structure_deductions]
        total_deductions = _dec(payslip.total_deductions)
        # Whatever is left is what other modules deducted: loan recoveries,
        # advances, adjustments. Derived, so the row always adds up.
        other_deductions = total_deductions - sum(statutory, _ZERO) - sum(structure, _ZERO)

        values: list[Any] = [
            offset + 1,
            employee.employee_code,
            employee.full_name,
            departments.get(employee.department_id) if employee.department_id else None,
            project_of(payslip, employee),
            employee.designation,
            locations.get(employee.location_id) if employee.location_id else None,
            employee.date_of_joining,
            employee.pan,
            employee.uan,
            employee.esi_number,
            employee.bank_account_number,
            employee.bank_ifsc,
            _dec(payslip.working_days),
            _dec(payslip.paid_days),
            _dec(payslip.lop_days),
            _dec(payslip.not_employed_days),
            *[earning_amounts.get(code, _ZERO) for code, _ in earnings],
            gross,
            other_earnings,
            gross + other_earnings,
            *statutory,
            *structure,
            other_deductions,
            total_deductions,
            _dec(payslip.net_pay),
            *[_dec(getattr(payslip, field)) for field, _ in _STATUTORY_EMPLOYER],
            *[employer_amounts.get(code, _ZERO) for code, _ in employer_lines],
            _dec(payslip.employer_cost),
        ]
        row_number = first_data + offset
        for index, ((_, _, kind), value) in enumerate(zip(columns, values, strict=True), start=1):
            cell = sheet.cell(row=row_number, column=index)
            cell.border = _BOX
            if kind == "text":
                _text(cell, None if value is None else str(value))
            elif kind == "date":
                cell.value = value
                cell.number_format = _DATE
            elif kind == "money":
                cell.value = value
                cell.number_format = _MONEY
                totals[index - 1] = (totals[index - 1] or _ZERO) + value
            else:
                cell.value = value

    total_row = first_data + len(rows)
    sheet.cell(row=total_row, column=3, value="Total").font = Font(bold=True)
    for index, total in enumerate(totals, start=1):
        cell = sheet.cell(row=total_row, column=index)
        cell.fill, cell.border = _TOTAL_FILL, _BOX
        if total is not None:
            cell.value = total
            cell.number_format = _MONEY
            cell.font = Font(bold=True)

    _widths(
        sheet,
        [6, 13, 26, 18, 20, 16, 13, 13, 14, 13, 18, 13, 9, 9, 9, 10] + [14] * (len(columns) - 16),
    )
    sheet.freeze_panes = sheet.cell(row=first_data, column=4)
    sheet.auto_filter.ref = f"A{head_row}:{get_column_letter(len(columns))}{total_row - 1}"

    # ── Other pay and deductions, item by item ──────────────────────────────
    items = book.create_sheet("Other pay & deductions")
    items["A1"] = f"Other pay and deductions — {label}"
    items["A1"].font = _TITLE_FONT
    items["A2"] = (
        f"What the {imported_from(run)} register paid between gross, deductions and the "
        "bank transfer: net-salary differences, adjustments and rounding. Their sums are "
        "the Other earnings and Other deductions columns on the register."
        if imported
        else "What other modules put on the payslips: claims, overtime, adjustments, bonus, "
        "loan and advance recoveries. Their sums are the Other earnings and Other "
        "deductions columns on the register."
    )
    headings = ["Employee code", "Employee name", "Item", "Type", "Amount", "Note"]
    for index, heading in enumerate(headings, start=1):
        cell = items.cell(row=4, column=index, value=heading)
        cell.font, cell.fill, cell.border = _HEADER_FONT, _HEADER_FILL, _BOX
    line_row = 5
    for payslip, employee in rows:
        for line in payslip.input_lines or []:
            kind = str(line.get("kind") or "")
            skipped = bool(line.get("skipped"))
            note = (
                "Held this month: not recovered (net pay would fall below the floor)."
                if skipped
                else str(line.get("basis") or "")
            )
            cells = [
                employee.employee_code,
                employee.full_name,
                str(line.get("name") or line.get("code") or ""),
                _INPUT_KIND_LABEL.get(kind, kind.replace("_", " ").capitalize()),
            ]
            for index, value in enumerate(cells, start=1):
                cell = items.cell(row=line_row, column=index)
                _text(cell, value)
                cell.border = _BOX
            amount = items.cell(row=line_row, column=5, value=_dec(line.get("amount")))
            amount.number_format, amount.border = _MONEY, _BOX
            _text(items.cell(row=line_row, column=6), note)
            items.cell(row=line_row, column=6).border = _BOX
            line_row += 1
    if line_row == 5:
        items.cell(row=5, column=1, value="Nothing from other modules on this run.")
    _widths(items, [14, 26, 36, 22, 14, 60])
    items.freeze_panes = "A5"

    buffer = io.BytesIO()
    book.save(buffer)

    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.EXPORT,
        entity_type="payroll_run",
        entity_id=run.id,
        actor=user,
        changes={"export": "register_xlsx", "rows": len(rows)},
        summary=(
            f"Downloaded the salary register (Excel) for {run.period_month}/{run.period_year}"
        ),
        request=request,
    )
    await session.commit()

    code = f"{entity.code}-" if entity is not None and entity.code else ""
    kind = "" if run.run_type.value == "regular" else f"-{run.run_type.value.replace('_', '-')}"
    filename = f"salary-register-{code}{run.period_year}-{run.period_month:02d}{kind}.xlsx"
    return filename, buffer.getvalue()
