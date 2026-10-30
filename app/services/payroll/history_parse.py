"""Read factoHR salary registers (Excel) into rows, exactly as typed. Pure: no database.

The registers are hand-built workbooks, one per month, with live formulas whose
cached values are what was paid. They come in several layouts over the years
(column letters shift, headers are renamed, an "actual" rate block is missing in
the first month, F&F settlements sit on a sheet of their own), so columns are
found by header name, never by letter.

Every figure is kept as the cell had it: a number becomes `Decimal(repr(value))`
(never binary arithmetic) and its raw text is kept beside it, so the importer and
an independent reconciler can both prove a stored figure equals the cell. PAN,
Aadhaar, UAN, ESIC, bank and e-mail columns are recognised only to be left out:
nothing here ever carries them (DPDP).

`parse_registers` groups files by payroll month. What the parser could not take
as given — a file whose title names another month, text in a money column, rows
after the data block — comes back as issues for the import report.
"""

from __future__ import annotations

import calendar
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

__all__ = [
    "DAY_KEYS",
    "DEDUCTION_KEYS",
    "EARNING_PARTS",
    "MONEY_KEYS",
    "Issue",
    "RegisterFile",
    "RegisterMonth",
    "RegisterRow",
    "RegisterSheet",
    "month_from_text",
    "parse_register",
    "parse_registers",
]

#: A person row's employee code. Anything else in column A ends the data block.
CODE_RE = re.compile(r"^C[A-Z]{2}\d{2}-\d{3}$")

#: The four earning parts and the gross, in both blocks (actual rate, earned).
EARNING_PARTS = ("basic", "hra", "special", "bonus")
_EARNING_KEYS = (*EARNING_PARTS, "gross")

#: Deduction cells that make up the register's total deduction (TDS included).
DEDUCTION_KEYS = (
    "pf_employee",
    "esi_employee",
    "medical_insurance",
    "professional_tax",
    "tds",
    "other_deduction",
)

DAY_KEYS = ("workable_days", "days_worked", "lop_days")

#: Every scalar figure a row can carry, besides the two earning blocks.
_SCALAR_MONEY = (
    "ctc_monthly",
    "ctc_earned",
    "pf_wage",
    "esi_wage",
    *DEDUCTION_KEYS,
    "total_deduction",
    "net_salary",
    "adjustment",
    "bank_transfer",
    "bank_transfer_rounded",
    "pf_employer",
    "esi_employer",
    "gratuity",
)
MONEY_KEYS = (
    *_SCALAR_MONEY,
    *(f"actual_{k}" for k in _EARNING_KEYS),
    *(f"earned_{k}" for k in _EARNING_KEYS),
)
_NUMERIC_KEYS = frozenset((*MONEY_KEYS, *DAY_KEYS))

_TEXT_KEYS = (
    "name",
    "status",
    "client",
    "project",
    "domain",
    "remarks",
    "pt_status",
    "gratuity_status",
)
_DATE_KEYS = ("doj", "dol", "payment_status")


def _norm(value: object) -> str:
    """A header as compared: spaces collapsed, case folded."""
    return re.sub(r"\s+", " ", str(value)).strip().casefold()


#: Header (normalised) → the key it holds. First occurrence wins; a repeat is a
#: helper column (`client`, `project` copies at the far right) and stays a cell.
_HEADERS: dict[str, str] = {
    "emp code": "code",
    "empid": "code",
    "emp id": "code",
    "name": "name",
    "employees name": "name",
    "employee name": "name",
    "status": "status",
    "date of joining": "doj",
    "doj": "doj",
    "client": "client",
    "project": "project",
    "domain": "domain",
    "dol": "dol",
    "lwd": "dol",
    "master ctc": "ctc_monthly",
    "ctc": "ctc_monthly",
    "monthly ctc": "ctc_monthly",
    "total ctc": "ctc_monthly",
    "ctc earned": "ctc_earned",
    "earned ctc": "ctc_earned",
    "workable days": "workable_days",
    "days worked": "days_worked",
    "paid days": "days_worked",
    "lop": "lop_days",
    "pf basic": "pf_wage",
    "esi gross": "esi_wage",
    "pf": "pf_employee",
    "esi": "esi_employee",
    "med ins": "medical_insurance",
    "medical insurance": "medical_insurance",
    "insurance": "medical_insurance",
    "pt": "professional_tax",
    "tds": "tds",
    "other deduction": "other_deduction",
    "total deduction": "total_deduction",
    "net salary": "net_salary",
    "payable amount": "payable",
    "adjustement": "adjustment",
    "adjustment": "adjustment",
    "bank tfr": "bank_transfer",
    "bank trs": "bank_transfer",
    "bank transfer": "bank_transfer",
    "bank transfer amount": "bank_transfer",
    "rounded": "bank_transfer_rounded",
    "round": "bank_transfer_rounded",
    "employer pf": "pf_employer",
    "pf emr": "pf_employer",
    "employer esi": "esi_employer",
    "esi emr": "esi_employer",
    "gratuity": "gratuity",
    "remarks": "remarks",
    "payment status": "payment_status",
    "salary status": "payment_status",
    "all pt status": "pt_status",
    "gratuity status": "gratuity_status",
}

#: Earning headers: the 1st occurrence is the actual (full-month) rate, the 2nd
#: what was earned. A register with a single block (Sep-2023) has earned only.
_EARNING_HEADERS = {
    "basic salary": "basic",
    "basic": "basic",
    "house rent allowance": "hra",
    "hra": "hra",
    "special allowance": "special",
    "special all": "special",
    "bonus": "bonus",
    "gross salary": "gross",
}

#: Identifier and bank columns: recognised only so that they are never kept.
_IDENTIFIER_HEADERS = frozenset(
    {
        "uan",
        "uan no",
        "uan no.",
        "esic no",
        "esic no.",
        "esi no",
        "esi no.",
        "pan",
        "pan no",
        "pan no.",
        "aadhar no",
        "adhar no",
        "aadhaar",
        "aadhaar no",
        "aadhar",
        "bank name",
        "account",
        "account no",
        "account no.",
        "account number",
        "ifsc",
        "ifsc code",
        "email id",
        "email",
        "personal email id",
        "personal mail id",
        "mail id",
    }
)

_MONTH_NAMES = {
    "january": 1,
    "february": 2,
    "march": 3,
    "april": 4,
    "may": 5,
    "june": 6,
    "july": 7,
    "august": 8,
    "september": 9,
    "october": 10,
    "november": 11,
    "december": 12,
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "sept": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}
_MONTH_RE = re.compile(
    r"(?<![a-z])(" + "|".join(sorted(_MONTH_NAMES, key=len, reverse=True)) + r")"
    r"[\s._]*[-" + chr(0x2013) + r"]?[\s._]*(\d{2,})",  # a hyphen or an en dash
    re.IGNORECASE,
)
_MONTH_CTC_RE = re.compile(r"^([a-z]+) ctc$")
_DIVISOR_RE = re.compile(
    r"^=\s*\$?[A-Z]{1,3}\$?\d+\s*/\s*(\d+(?:\.\d+)?)\s*\*\s*\$?[A-Z]{1,3}\$?\d+\s*$"
)
_NUMBER_RE = re.compile(r"^-?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?$")
_PAN_RE = re.compile(r"^[A-Z]{5}\d{4}[A-Z]$")
_IFSC_RE = re.compile(r"^[A-Z]{4}0[A-Z0-9]{6}$")
_HEADER_SCAN_ROWS = 8
_MAX_COLUMNS = 200


# ── Results ─────────────────────────────────────────────────────────────────


@dataclass(slots=True)
class Issue:
    """Something the report must show. Errors block the import."""

    level: str  # "error" | "warning" | "note"
    category: str
    message: str
    month: str | None = None
    file: str | None = None
    sheet: str | None = None
    row: int | None = None
    code: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "level": self.level,
            "category": self.category,
            "message": self.message,
            "month": self.month,
            "file": self.file,
            "sheet": self.sheet,
            "row": self.row,
            "code": self.code,
        }


@dataclass(slots=True)
class RegisterRow:
    """One person row of a register sheet, every figure as the cell had it."""

    file: str
    sheet: str
    row: int
    kind: str  # "main" | "fnf"
    hidden: bool
    code: str
    name: str | None = None
    status: str | None = None
    doj: str | None = None
    dol: str | None = None
    client: str | None = None
    project: str | None = None
    domain: str | None = None
    remarks: str | None = None
    payment_status: str | None = None
    pt_status: str | None = None
    gratuity_status: str | None = None
    ctc_monthly: Decimal | None = None
    ctc_earned: Decimal | None = None
    workable_days: Decimal | None = None
    days_worked: Decimal | None = None
    lop_days: Decimal | None = None
    #: The full-month rate block; None where the register has none (Sep-2023).
    actual: dict[str, Decimal | None] | None = None
    earned: dict[str, Decimal | None] = field(default_factory=dict)
    pf_wage: Decimal | None = None
    esi_wage: Decimal | None = None
    pf_employee: Decimal | None = None
    esi_employee: Decimal | None = None
    medical_insurance: Decimal | None = None
    professional_tax: Decimal | None = None
    tds: Decimal | None = None
    other_deduction: Decimal | None = None
    total_deduction: Decimal | None = None
    net_salary: Decimal | None = None
    adjustment: Decimal | None = None
    bank_transfer: Decimal | None = None
    bank_transfer_rounded: Decimal | None = None
    pf_employer: Decimal | None = None
    esi_employer: Decimal | None = None
    gratuity: Decimal | None = None
    #: {"Header (letter)": raw text} for every non-identifier, non-empty cell.
    cells: dict[str, str] = field(default_factory=dict)
    #: key → raw text of the cell the key was read from.
    raw: dict[str, str] = field(default_factory=dict)
    #: key → (column letter, header), shared by every row of the sheet.
    columns: dict[str, tuple[str, str]] = field(default_factory=dict)
    #: The sheet has a ROUNDED / Round column (the paid, rounded bank figure).
    has_rounded: bool = False
    #: The sheet has a Domain column, so Project (not Client) names the account.
    has_domain: bool = False

    def money(self, key: str) -> Decimal | None:
        """A figure by its flat key, including `actual_basic` / `earned_gross`."""
        if key.startswith("actual_"):
            return (self.actual or {}).get(key[7:])
        if key.startswith("earned_"):
            return self.earned.get(key[7:])
        return getattr(self, key)

    @property
    def account(self) -> str | None:
        """The client account the row was paid under, as typed."""
        return self.project if self.has_domain else self.client


@dataclass(slots=True)
class RegisterSheet:
    name: str
    role: str  # "main" | "fnf"
    header_row: int
    columns: dict[str, tuple[str, str]]
    rows: list[RegisterRow]
    first_row: int | None
    last_row: int | None
    has_actual: bool
    has_rounded: bool
    has_domain: bool
    #: The totals row straight after the data: key → raw text.
    register_totals: dict[str, str]
    totals_row: int | None
    notes: list[str] = field(default_factory=list)

    def summary(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "role": self.role,
            "header_row": self.header_row,
            "rows_read": len(self.rows),
            "first_row": self.first_row,
            "last_row": self.last_row,
            "totals_row": self.totals_row,
        }


@dataclass(slots=True)
class RegisterFile:
    name: str
    sha256: str
    size: int
    year: int | None
    month: int | None
    title: str | None
    sheets: list[RegisterSheet] = field(default_factory=list)
    #: (sheet name, row count) of sheets that are not imported.
    other_sheets: list[tuple[str, int]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)

    @property
    def rows(self) -> list[RegisterRow]:
        return [row for sheet in self.sheets for row in sheet.rows]

    @property
    def main(self) -> RegisterSheet | None:
        return next((s for s in self.sheets if s.role == "main"), None)

    def summary(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "sha256": self.sha256,
            "size": self.size,
            "sheets": [s.summary() for s in self.sheets],
            "other_sheets": [{"name": n, "rows": r} for n, r in self.other_sheets],
            "notes": list(self.notes),
        }


@dataclass(slots=True)
class RegisterMonth:
    year: int
    month: int
    files: list[RegisterFile]
    issues: list[Issue] = field(default_factory=list)

    @property
    def label(self) -> str:
        return f"{calendar.month_abbr[self.month]} {self.year}"

    @property
    def key(self) -> str:
        return f"{self.year}-{self.month:02d}"

    @property
    def rows(self) -> list[RegisterRow]:
        return [row for f in self.files for row in f.rows]


# ── Values ──────────────────────────────────────────────────────────────────


def _raw(value: object) -> str | None:
    """The cell as text: numbers as Python wrote them, dates as ISO."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _number(value: object) -> tuple[Decimal | None, bool]:
    """(the figure, whether the cell held text that is not a number)."""
    if value is None or isinstance(value, bool):
        return None, isinstance(value, bool)
    if isinstance(value, int):
        return Decimal(value), False
    if isinstance(value, float):
        return Decimal(repr(value)), False
    if isinstance(value, Decimal):
        return value, False
    if isinstance(value, datetime | date):
        return None, True
    text = str(value).strip().replace(",", "")
    if not text:
        return None, False
    if _NUMBER_RE.match(text):
        try:
            return Decimal(text), False
        except InvalidOperation:  # pragma: no cover - the pattern admits only numbers
            return None, True
    return None, True


_IDENTIFIER_TOKENS = (
    re.compile(r"\S+@\S+"),  # an e-mail address
    re.compile(r"\b[A-Za-z]{5}\d{4}[A-Za-z]\b"),  # PAN
    re.compile(r"\b[A-Za-z]{4}0[A-Za-z0-9]{6}\b"),  # IFSC
    re.compile(r"(?<![\d.])\d[\d ]{7,}\d(?![\d.])"),  # account, Aadhaar, UAN: 9+ digits
)
REMOVED = "[removed]"


def _scrub(text: str) -> tuple[str, bool]:
    """Free text with any account number, IFSC, PAN or e-mail in it taken out.

    A remark such as "KOTAK MAHINDRA BANK KKBK0005297 6346262973 (15 days F&F
    paid)" keeps its words and loses the bank details.
    """
    cleaned = text
    for pattern in _IDENTIFIER_TOKENS:
        cleaned = pattern.sub(REMOVED, cleaned)
    return cleaned, cleaned != text


def _looks_like_identifier(raw: str) -> bool:
    """A cell under no recognised header that still looks like an ID or bank detail."""
    text = raw.strip()
    if "@" in text:
        return True
    upper = text.upper()
    if _PAN_RE.match(upper) or _IFSC_RE.match(upper):
        return True
    digits = text.removesuffix(".0")
    return digits.isdigit() and len(digits) >= 9


def month_from_text(text: str) -> tuple[int, int] | None:
    """(year, month) named in a file name or a title, e.g. "May-241" → (2024, 5)."""
    for match in _MONTH_RE.finditer(text):
        month = _MONTH_NAMES[match.group(1).casefold()]
        digits = match.group(2)
        if len(digits) >= 4 and 1990 <= int(digits[:4]) <= 2099:
            return int(digits[:4]), month
        # Two digits, or more that make no possible year ("May-241"): the first two.
        return 2000 + int(digits[:2]), month
    return None


Grid = list[tuple[Any, ...]]


def _header_row(grid: Grid, wanted: tuple[str, ...]) -> int | None:
    """The first row among the first eight whose column A is `Emp Code` / `EmpId`."""
    for r, values in enumerate(grid[:_HEADER_SCAN_ROWS], start=1):
        value = values[0] if values else None
        if isinstance(value, str) and re.sub(r"\s+", "", value).casefold() in wanted:
            return r
    return None


def _cell(grid: Grid, r: int, c: int) -> Any:
    """Grid value at 1-based (row, column); None outside it."""
    if r < 1 or r > len(grid):
        return None
    values = grid[r - 1]
    return values[c - 1] if 0 < c <= len(values) else None


_ROW_TAG = re.compile(rb"<(?:\w+:)?row(\s[^>]*)?>")
_ROW_NUM = re.compile(rb'(?:^|\s)r="(\d+)"')
_HIDDEN = re.compile(rb'(?:^|\s)hidden="(?:1|true)"')


def _hidden_rows(data: bytes, ws) -> set[int]:
    """Rows hidden in the sheet. Read from the sheet's XML: the fast read-only
    loader does not keep row properties."""
    path = getattr(ws, "_worksheet_path", None)
    if not path:
        return set()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read(path.lstrip("/"))
    except (KeyError, zipfile.BadZipFile, OSError):
        return set()
    hidden: set[int] = set()
    for match in _ROW_TAG.finditer(xml):
        attrs = match.group(1) or b""
        if _HIDDEN.search(attrs):
            number = _ROW_NUM.search(attrs)
            if number:
                hidden.add(int(number.group(1)))
    return hidden


@dataclass(slots=True)
class _Layout:
    #: key → column index (1-based)
    index: dict[str, int]
    columns: dict[str, tuple[str, str]]
    #: column index → (letter, header) for every column kept as a cell
    cell_columns: dict[int, tuple[str, str]]
    has_actual: bool
    cells_only: list[str]


def _layout(
    headers: list[object], month: int | None, role: str, issues: list[Issue], where: dict
) -> _Layout:
    index: dict[str, int] = {}
    columns: dict[str, tuple[str, str]] = {}
    cell_columns: dict[int, tuple[str, str]] = {}
    cells_only: list[str] = []
    earning_seen: dict[str, int] = {}
    month_ctc: int | None = None

    gross_count = sum(1 for h in headers if h is not None and _norm(h) == "gross salary")
    single_block = gross_count == 1

    for i, header in enumerate(headers, start=1):
        letter = get_column_letter(i)
        text = _raw(header) if header is not None else None
        norm = _norm(header) if text else ""
        if norm in _IDENTIFIER_HEADERS:
            continue  # never kept, not even as a cell
        label = f"{text} ({letter})" if text else f"({letter})"
        cell_columns[i] = (letter, label)
        if not norm:
            continue
        key: str | None = None
        if norm in _EARNING_HEADERS:
            part = _EARNING_HEADERS[norm]
            seen = earning_seen.get(part, 0)
            earning_seen[part] = seen + 1
            if single_block:
                key = f"earned_{part}" if seen == 0 else None
            elif seen == 0:
                key = f"actual_{part}"
            elif seen == 1:
                key = f"earned_{part}"
        elif norm in _HEADERS:
            key = _HEADERS[norm]
        else:
            ctc = _MONTH_CTC_RE.match(norm)
            if ctc and ctc.group(1) in _MONTH_NAMES:
                if month is not None and _MONTH_NAMES[ctc.group(1)] == month and month_ctc is None:
                    month_ctc = i
                    continue
                cells_only.append(label)
                continue
        if key is None:
            cells_only.append(label)
            continue
        if key == "payable":
            if "payable" in index:
                cells_only.append(label)
                continue
            index["payable"] = i
            columns["payable"] = (letter, text or "")
            continue
        if key in index:
            cells_only.append(label)  # a helper repeating a header further right
            continue
        index[key] = i
        columns[key] = (letter, text or "")

    if month_ctc is not None:
        if "ctc_monthly" in index:
            cells_only.append(cell_columns[index["ctc_monthly"]][1])
        index["ctc_monthly"] = month_ctc
        letter, label = cell_columns[month_ctc]
        columns["ctc_monthly"] = (letter, label.rsplit(" (", 1)[0])

    # The F&F sheet's Payable Amount is the bank transfer, and the net where the
    # sheet has no net of its own.
    if "payable" in index:
        if "bank_transfer" not in index:
            index["bank_transfer"] = index["payable"]
            columns["bank_transfer"] = columns["payable"]
        if "net_salary" not in index:
            index["net_salary"] = index["payable"]
            columns["net_salary"] = columns["payable"]
        index.pop("payable")
        columns.pop("payable")

    for needed in ("code", "earned_gross"):
        if needed not in index:
            issues.append(
                Issue(
                    "error",
                    "missing_column",
                    f"The {'F&F ' if role == 'fnf' else ''}sheet has no "
                    f"{'Emp Code' if needed == 'code' else 'earned Gross Salary'} column.",
                    **where,
                )
            )
    return _Layout(
        index=index,
        columns=columns,
        cell_columns=cell_columns,
        has_actual=not single_block and "actual_gross" in index,
        cells_only=cells_only,
    )


def _divisors(formula_grid: Grid | None, rows: list[int], column: int | None) -> dict[int, Decimal]:
    """The days divisor of each F&F row, from its earned-basic formula (=H3/31*E3)."""
    found: dict[int, Decimal] = {}
    if formula_grid is None or column is None:
        return found
    for r in rows:
        value = _cell(formula_grid, r, column)
        if isinstance(value, str):
            match = _DIVISOR_RE.match(value.replace(" ", ""))
            if match:
                found[r] = Decimal(match.group(1))
    return found


def _read_sheet(
    grid: Grid,
    *,
    title: str,
    file: str,
    role: str,
    header_row: int,
    month: int | None,
    month_key: str | None,
    issues: list[Issue],
    hidden: set[int],
    formula_grid: Grid | None = None,
) -> RegisterSheet:
    where = {"month": month_key, "file": file, "sheet": title}
    headers = list(grid[header_row - 1]) if header_row <= len(grid) else []
    layout = _layout(headers, month, role, issues, where)
    index = layout.index
    mapped_columns = set(index.values())
    has_rounded = "bank_transfer_rounded" in index
    has_domain = "domain" in index
    notes: list[str] = []
    if layout.cells_only:
        notes.append(
            f"{title}: columns kept only as the row's cells (not imported as figures): "
            + ", ".join(layout.cells_only)
        )

    rows: list[RegisterRow] = []
    first = last = None
    code_col = index.get("code", 1)
    stop_at: int | None = None
    dropped_cells = 0
    scrubbed = 0
    for r in range(header_row + 1, len(grid) + 1):
        values = grid[r - 1]
        cell = _cell(grid, r, code_col)
        code = re.sub(r"\s+", "", str(cell)) if cell is not None else ""
        if not CODE_RE.match(code):
            stop_at = r
            break
        first = first or r
        last = r
        row = RegisterRow(
            file=file,
            sheet=title,
            row=r,
            kind=role,
            hidden=r in hidden,
            code=code,
            columns=layout.columns,
            has_rounded=has_rounded,
            has_domain=has_domain,
        )
        earned: dict[str, Decimal | None] = {}
        actual: dict[str, Decimal | None] | None = {} if layout.has_actual else None
        texts: dict[int, str] = {}
        for col, value in enumerate(values, start=1):
            if isinstance(value, str) and value.strip():
                text, changed = _scrub(value.strip())
                texts[col] = text
                scrubbed += changed
        for key, col in index.items():
            if key == "code":
                continue
            value = values[col - 1] if col <= len(values) else None
            raw = texts.get(col) or _raw(value)
            if raw is not None:
                row.raw[key] = raw
            if key in _NUMERIC_KEYS:
                number, is_text = _number(value)
                if is_text and raw is not None:
                    level = "warning" if key in MONEY_KEYS else "note"
                    letter, header = layout.columns[key]
                    issues.append(
                        Issue(
                            level,
                            "text_in_figure",
                            f"Row {r}, {header} ({letter}) holds text '{raw}'; read as empty.",
                            row=r,
                            code=code,
                            **where,
                        )
                    )
                if key.startswith("actual_"):
                    if actual is not None:
                        actual[key[7:]] = number
                elif key.startswith("earned_"):
                    earned[key[7:]] = number
                else:
                    setattr(row, key, number)
            elif key in _DATE_KEYS or key in _TEXT_KEYS:
                setattr(row, key, raw)
        row.earned = {k: earned.get(k) for k in _EARNING_KEYS}
        if actual is not None:
            row.actual = {k: actual.get(k) for k in _EARNING_KEYS}
        for col, (_letter, label) in layout.cell_columns.items():
            value = values[col - 1] if col <= len(values) else None
            raw = texts.get(col) or _raw(value)
            if raw is None:
                continue
            if col not in mapped_columns and _looks_like_identifier(raw):
                dropped_cells += 1
                continue
            row.cells[label] = raw
        rows.append(row)

    # F&F: the days the month was divided by, from the earned-basic formula.
    if role == "fnf":
        divisors = _divisors(formula_grid, [row.row for row in rows], index.get("earned_basic"))
        for row in rows:
            divisor = divisors.get(row.row)
            row.workable_days = divisor
            if divisor is None:
                issues.append(
                    Issue(
                        "note",
                        "fnf_days",
                        f"Row {row.row}: the days divisor could not be read from the earned "
                        "Basic formula; working days are taken as the paid days.",
                        row=row.row,
                        code=row.code,
                        **where,
                    )
                )
            else:
                row.raw["workable_days"] = str(divisor)
            # The sheet has no adjustment and no total deduction of its own.
            if "adjustment" not in index:
                row.adjustment = Decimal(0)
            if "total_deduction" not in index:
                parts = [row.money(k) for k in DEDUCTION_KEYS]
                total = sum((p for p in parts if p is not None), Decimal(0))
                row.total_deduction = total
                row.raw["total_deduction"] = str(total)

    if dropped_cells:
        notes.append(
            f"{title}: {dropped_cells} cell(s) under unnamed or helper columns looked like an "
            "ID, account or e-mail and were left out."
        )
    if scrubbed:
        notes.append(
            f"{title}: {scrubbed} text cell(s) (remarks and the like) held an account number, "
            f"IFSC, PAN or e-mail; those parts were replaced by {REMOVED}."
        )
    hidden_read = sorted(row.row for row in rows if row.hidden)
    if hidden_read:
        notes.append(
            f"{title}: hidden row(s) {', '.join(map(str, hidden_read))} read like any other "
            "(a SUBTOTAL in the sheet leaves them out)."
        )

    # The totals row: straight after the data, if it has figures in money columns.
    register_totals: dict[str, str] = {}
    totals_row: int | None = None
    if stop_at is not None and rows:
        for key, col in index.items():
            if key not in _NUMERIC_KEYS:
                continue
            value = _cell(grid, stop_at, col)
            number, _ = _number(value)
            if number is not None and not isinstance(value, str):
                register_totals[key] = _raw(value) or ""
        if any(key in MONEY_KEYS for key in register_totals):
            totals_row = stop_at
        else:
            register_totals = {}

    # Codes further down that the data block did not reach (a stray duplicate
    # after a blank row, an embedded copy of the F&F sheet).
    if stop_at is not None:
        stray: list[str] = []
        for r in range(stop_at, len(grid) + 1):
            value = _cell(grid, r, code_col)
            if value is None:
                continue
            text = re.sub(r"\s+", "", str(value))
            if CODE_RE.match(text) or text.casefold() in ("empid", "empcode"):
                stray.append(f"row {r} ({text})")
        if stray:
            more = len(stray) - 12
            shown = ", ".join(stray[:12]) + (f" and {more} more" if more > 0 else "")
            notes.append(
                f"{title}: not read, below the data block (which ends at the first row "
                f"without an employee code): {shown}."
            )

    return RegisterSheet(
        name=title,
        role=role,
        header_row=header_row,
        columns=layout.columns,
        rows=rows,
        first_row=first,
        last_row=last,
        has_actual=layout.has_actual,
        has_rounded=has_rounded,
        has_domain=has_domain,
        register_totals=register_totals,
        totals_row=totals_row,
        notes=notes,
    )


# ── One file ────────────────────────────────────────────────────────────────


def _title(grid: Grid, header_row: int) -> str | None:
    """A title in the first rows naming the register's month ("Salary Register Apr - 24")."""
    for r in range(1, min(header_row, 4)):
        for value in grid[r - 1][:30]:
            if isinstance(value, str) and "register" in value.casefold() and month_from_text(value):
                return value.strip()
    return None


def _grid(ws) -> Grid:
    """Every row of a sheet as values, column A first, up to `_MAX_COLUMNS` wide."""
    width = min(ws.max_column or _MAX_COLUMNS, _MAX_COLUMNS)
    return [tuple(values) for values in ws.iter_rows(max_col=width, values_only=True)]


def _open(data: bytes, *, data_only: bool):
    # Read-only: a register's styling (conditional formats, thousands of styles)
    # takes minutes to load in full mode and adds nothing to the figures.
    return load_workbook(io.BytesIO(data), data_only=data_only, read_only=True)


def parse_register(filename: str, data: bytes) -> RegisterFile:
    """Read one register workbook. Problems come back as issues, never as exceptions."""
    found = month_from_text(filename)
    result = RegisterFile(
        name=filename,
        sha256=hashlib.sha256(data).hexdigest(),
        size=len(data),
        year=found[0] if found else None,
        month=found[1] if found else None,
        title=None,
    )
    month_key = f"{found[0]}-{found[1]:02d}" if found else None
    where = {"month": month_key, "file": filename}
    if found is None:
        result.issues.append(
            Issue(
                "error",
                "month",
                "The file name names no month and year (e.g. 'Salary Register Apr-24').",
                **where,
            )
        )
    try:
        book = _open(data, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        result.issues.append(
            Issue("error", "unreadable", f"Not an Excel workbook that can be read: {exc}", **where)
        )
        return result

    main: tuple[Any, Grid, int] | None = None
    fnf: list[tuple[Any, Grid, int]] = []
    try:
        for ws in book.worksheets:
            head = [
                tuple(values)
                for values in ws.iter_rows(
                    min_row=1, max_row=_HEADER_SCAN_ROWS, max_col=1, values_only=True
                )
            ]
            if main is None and _header_row(head, ("empcode",)) is not None:
                grid = _grid(ws)
                main = (ws, grid, _header_row(grid, ("empcode",)) or 1)
                continue
            if _header_row(head, ("empid",)) is not None:
                grid = _grid(ws)
                fnf.append((ws, grid, _header_row(grid, ("empid",)) or 1))
                continue
            result.other_sheets.append((ws.title, ws.max_row or 0))
    finally:
        book.close()

    if main is None:
        result.issues.append(
            Issue(
                "error",
                "no_register",
                "No sheet has 'Emp Code' in column A of its first rows; nothing can be read.",
                **where,
            )
        )
        return result

    main_ws, main_grid, main_header = main
    title = _title(main_grid, main_header)
    result.title = title
    if title and found:
        named = month_from_text(title)
        if named and named != found:
            result.issues.append(
                Issue(
                    "error",
                    "month",
                    f"The file name says {calendar.month_abbr[found[1]]} {found[0]} but the title "
                    f"'{title}' says {calendar.month_abbr[named[1]]} {named[0]}.",
                    **where,
                )
            )

    formula_grids: dict[str, Grid] = {}
    if fnf:
        try:
            formulas = _open(data, data_only=False)
            try:
                for ws, _values, _header in fnf:
                    formula_grids[ws.title] = _grid(formulas[ws.title])
            finally:
                formulas.close()
        except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError):
            formula_grids = {}

    month = found[1] if found else None
    result.sheets.append(
        _read_sheet(
            main_grid,
            title=main_ws.title,
            file=filename,
            role="main",
            header_row=main_header,
            month=month,
            month_key=month_key,
            issues=result.issues,
            hidden=_hidden_rows(data, main_ws),
        )
    )
    for ws, grid, header in fnf:
        result.sheets.append(
            _read_sheet(
                grid,
                title=ws.title,
                file=filename,
                role="fnf",
                header_row=header,
                month=month,
                month_key=month_key,
                issues=result.issues,
                hidden=_hidden_rows(data, ws),
                formula_grid=formula_grids.get(ws.title),
            )
        )
    for sheet in result.sheets:
        result.notes.extend(sheet.notes)
    for name, count in result.other_sheets:
        result.notes.append(f"Sheet '{name}' ({count} rows) is not a register sheet; not imported.")
    if not any(sheet.rows for sheet in result.sheets):
        result.issues.append(
            Issue("error", "no_rows", "The register has no person rows under its header.", **where)
        )
    return result


def parse_registers(files: list[tuple[str, bytes]]) -> list[RegisterMonth]:
    """Read every file and group them by payroll month, oldest first."""
    by_month: dict[tuple[int, int], list[RegisterFile]] = {}
    unplaced: list[RegisterFile] = []
    for name, data in files:
        parsed = parse_register(name, data)
        if parsed.year is None or parsed.month is None:
            unplaced.append(parsed)
            continue
        by_month.setdefault((parsed.year, parsed.month), []).append(parsed)
    months: list[RegisterMonth] = []
    for (year, month), group in sorted(by_month.items()):
        result = RegisterMonth(year=year, month=month, files=group)
        if len(group) > 1:
            result.issues.append(
                Issue(
                    "error",
                    "duplicate_month",
                    f"{len(group)} files are for {result.label}: "
                    + ", ".join(f.name for f in group)
                    + ". Keep one.",
                    month=result.key,
                )
            )
        months.append(result)
    if unplaced:
        # Kept so the report can say why these files were left out.
        months.append(RegisterMonth(year=0, month=0, files=unplaced))
    return months
