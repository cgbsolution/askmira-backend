"""The employee import template: an Employees sheet to fill and an Instructions sheet.

Its headers are factoHR's own wording wherever the two systems share a field
(Emp Code, Joining Date, Pan No, UA No…), so a column copied out of a factoHR
export lands in the right place. The workspace's custom fields are added at
the end. No example row sits in the Employees sheet — a template uploaded
untouched must not create a person; the examples are on Instructions.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.models.employee import CustomFieldDefinition
from app.services.employee_import.fields import FIELDS

#: Rows the drop-downs cover.
_ROWS = 2000

_HEADER_FILL = PatternFill("solid", fgColor="EB061D")  # AskMira Red
_REQUIRED_FILL = PatternFill("solid", fgColor="F8A01A")  # Marigold
_SECTION_FILL = PatternFill("solid", fgColor="F9EFE8")  # Cream


def build_template(custom: list[CustomFieldDefinition]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Employees"

    columns = [field for field in FIELDS if field.template]
    headers = [field.label for field in columns] + [d.label for d in custom]
    sheet.append(headers)
    for index, field in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="100404" if field.required else "FFFFFF")
        cell.fill = _REQUIRED_FILL if field.required else _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(field.label) + 4)
        if field.choices:
            validation = DataValidation(
                type="list",
                formula1='"' + ",".join(field.choices) + '"',
                allow_blank=True,
                showErrorMessage=False,
            )
            sheet.add_data_validation(validation)
            letter = get_column_letter(index)
            validation.add(f"{letter}2:{letter}{_ROWS}")
    for offset, definition in enumerate(custom, start=len(columns) + 1):
        cell = sheet.cell(row=1, column=offset)
        cell.font = Font(bold=True, color="100404" if definition.is_required else "FFFFFF")
        cell.fill = _REQUIRED_FILL if definition.is_required else _HEADER_FILL
        width = max(14, len(definition.label) + 4)
        sheet.column_dimensions[get_column_letter(offset)].width = width
        if definition.field_type == "select" and definition.options:
            joined = ",".join(definition.options)
            if len(joined) <= 250:
                validation = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True)
                sheet.add_data_validation(validation)
                letter = get_column_letter(offset)
                validation.add(f"{letter}2:{letter}{_ROWS}")
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "C2"

    guide = book.create_sheet("Instructions")
    guide.append(["How to fill the Employees sheet"])
    guide["A1"].font = Font(bold=True, size=14)
    for line in (
        "One row per person. Marigold headers must be filled; the rest are optional.",
        "Dates as DD-MM-YYYY. Reporting Manager is the manager's Emp Code.",
        "Somebody who has left: Status Inactive and their Leaving Date.",
        "An export from factoHR (Employee Detail report) can be uploaded as it is.",
        "Employee codes already in AskMira are skipped, so a corrected file can be "
        "imported again.",
    ):
        guide.append([line])
    guide.append([])
    guide.append(["Column", "Required", "Section", "What to enter", "Example"])
    header_row = guide.max_row
    for column in range(1, 6):
        cell = guide.cell(row=header_row, column=column)
        cell.font = Font(bold=True)
        cell.fill = _SECTION_FILL
    for field in columns:
        hint = field.hint
        if field.choices and not hint:
            hint = "One of: " + ", ".join(field.choices) + "."
        guide.append([field.label, "Yes" if field.required else "", field.section, hint,
                      field.example])
    for definition in custom:
        options = ", ".join(definition.options or [])
        guide.append([
            definition.label,
            "Yes" if definition.is_required else "",
            "Custom fields",
            f"One of: {options}." if options else f"A {definition.field_type} value.",
            "",
        ])
    for letter, width in zip("ABCDE", (34, 10, 16, 70, 28), strict=True):
        guide.column_dimensions[letter].width = width
    for row in guide.iter_rows(min_row=header_row + 1):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
