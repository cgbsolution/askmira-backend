"""The employee export: every field of the employee record, in Excel or CSV.

HR chooses the columns; nothing chosen means everything. The catalogue below
is the whole record as the profile shows it to HR — basic and job details,
personal, addresses, statutory numbers, the salary account, emergency
contacts, family, nominees, education, previous jobs, what the offer carried,
and the workspace's own custom fields and tables.

Lists (family, education…) become numbered column sets: "Family member 1 -
Name", "Family member 2 - Name", as many as the most anybody exported has, and
at least one so the columns are there to fill.

Three things are left out on purpose:

- Salary. It has its own export under Payroll → Salaries, and a sub-admin may
  export People without being allowed to see pay.
- The candidate's fit score from hiring: an AI score is never shown without
  its signals (FS-ANL-02), and those live on the candidate record.
- Inactive custom fields and tables: the profile does not show them either.

PAN, Aadhaar and account numbers are written in full (HR sees them in full on
the profile). One that cannot be decrypted — a copy of the data under another
key — says so in its cell instead of failing the whole file.
"""

from __future__ import annotations

import csv
import io
import uuid
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Literal

from fastapi import HTTPException, Request, status
from sqlalchemy import Text, select, type_coerce
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import load_only

from app.core import permissions
from app.core.crypto import EncryptionError, field_reader, secret_reader
from app.core.deps import tenant_scoped
from app.models.audit import AuditAction
from app.models.custom_tables import CustomTableDefinition, CustomTableRow
from app.models.disbursement import EmployeeBankAccount
from app.models.employee import CustomFieldDefinition, Employee, EmployeeProfile
from app.models.org import Department, LegalEntity, Location
from app.models.timesheets import Project
from app.models.user import User, UserRole
from app.services import audit
from app.services.employee import filter_directory

__all__ = [
    "EXPORT_LIMIT",
    "ExportFile",
    "ExportFilters",
    "catalogue",
    "export_employees",
]

#: The most people one export carries; the screen pages, a file does not.
EXPORT_LIMIT = 20_000

FileType = Literal["xlsx", "csv"]

#: Written in place of a number that cannot be decrypted with this server's key.
UNREADABLE = "Could not be read"

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

#: Columns that are encrypted on `employees`; read raw and decrypted one by one.
_ENCRYPTED = ("pan", "aadhaar_number", "bank_account_number")


# ═══════════════════════════════════════════════════════════════════════════
# How values are shown
# ═══════════════════════════════════════════════════════════════════════════

_STATUS = {
    "active": "Active",
    "on_notice": "On notice",
    "on_leave": "On leave",
    "exited": "Exited",
    "pre_joining": "Pre-joining",
}
_EMPLOYMENT_TYPE = {
    "full_time": "Full time",
    "part_time": "Part time",
    "contract": "Contract",
    "intern": "Intern",
    "consultant": "Consultant",
}
_GENDER = {
    "female": "Female",
    "male": "Male",
    "non_binary": "Non-binary",
    "undisclosed": "Prefer not to say",
}
_MARITAL = {
    "single": "Single",
    "married": "Married",
    "separated": "Separated",
    "divorced": "Divorced",
    "widowed": "Widowed",
    "undisclosed": "Prefer not to say",
}
_RELATIONSHIP = {
    "spouse": "Spouse",
    "child": "Child",
    "father": "Father",
    "mother": "Mother",
    "sibling": "Sibling",
    "father_in_law": "Father-in-law",
    "mother_in_law": "Mother-in-law",
    "grandparent": "Grandparent",
    "friend": "Friend",
    "other": "Other",
}
_PURPOSE = {
    "pf": "Provident Fund",
    "eps": "Pension (EPS)",
    "gratuity": "Gratuity",
    "insurance": "Insurance",
    "superannuation": "Superannuation",
}
_EDUCATION_LEVEL = {
    "secondary": "Secondary (10th)",
    "higher_secondary": "Higher secondary (12th)",
    "diploma": "Diploma",
    "graduate": "Graduate",
    "postgraduate": "Postgraduate",
    "doctorate": "Doctorate",
    "certification": "Certification",
    "other": "Other",
}
_ROLE = {
    UserRole.SUPER_ADMIN: "Platform admin",
    UserRole.ADMIN: "Admin",
    UserRole.SUB_ADMIN: "Sub-admin",
    UserRole.EMPLOYEE: "Employee",
}
#: State and union-territory names by the code AskMira stores.
_STATES = {
    "AN": "Andaman & Nicobar Islands", "AP": "Andhra Pradesh", "AR": "Arunachal Pradesh",
    "AS": "Assam", "BR": "Bihar", "CH": "Chandigarh", "CG": "Chhattisgarh",
    "CT": "Chhattisgarh", "DN": "Dadra & Nagar Haveli", "DD": "Daman & Diu",
    "DH": "Dadra & Nagar Haveli and Daman & Diu", "DL": "Delhi", "GA": "Goa",
    "GJ": "Gujarat", "HR": "Haryana", "HP": "Himachal Pradesh", "JK": "Jammu & Kashmir",
    "JH": "Jharkhand", "KA": "Karnataka", "KL": "Kerala", "LA": "Ladakh", "LD": "Lakshadweep",
    "MP": "Madhya Pradesh", "MH": "Maharashtra", "MN": "Manipur", "ML": "Meghalaya",
    "MZ": "Mizoram", "NL": "Nagaland", "OD": "Odisha", "OR": "Odisha", "PY": "Puducherry",
    "PB": "Punjab", "RJ": "Rajasthan", "SK": "Sikkim", "TN": "Tamil Nadu", "TG": "Telangana",
    "TS": "Telangana", "TR": "Tripura", "UP": "Uttar Pradesh", "UK": "Uttarakhand",
    "UT": "Uttarakhand", "WB": "West Bengal",
}  # fmt: skip


def _yes_no(value: Any) -> str | None:
    if value is None or value == "":
        return None
    return "Yes" if value is True or str(value).lower() in {"true", "yes", "1"} else "No"


def _day(value: Any) -> date | str | None:
    """A stored date (a date, or ISO text from a JSON section) as a date."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return str(value)


def _number(value: Any) -> int | Decimal | str | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return _yes_no(value)
    if isinstance(value, int | Decimal):
        return value
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        return str(value)
    return int(number) if number == number.to_integral_value() else number


def _label(table: dict[str, str], value: Any, own_words: Any = None) -> str | None:
    """An enum value in words; the source system's own words when it gave them."""
    if own_words:
        return str(own_words)
    if value is None or value == "":
        return None
    raw = getattr(value, "value", value)
    return table.get(str(raw), str(raw))


def _state(code: Any) -> str | None:
    if not code:
        return None
    return _STATES.get(str(code).upper(), str(code))


def _custom_value(kind: str, value: Any) -> Any:
    if value is None or value == "" or value == []:
        return None
    if kind == "boolean":
        return _yes_no(value)
    if kind == "date":
        return _day(value)
    if kind == "number":
        return _number(value)
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return value


# ═══════════════════════════════════════════════════════════════════════════
# One person, as the export reads them
# ═══════════════════════════════════════════════════════════════════════════


@dataclass
class _Person:
    employee: Employee
    profile: dict[str, Any] = field(default_factory=dict)
    pan: str | None = None
    aadhaar: str | None = None
    account: dict[str, Any] = field(default_factory=dict)
    login: tuple[UserRole, bool] | None = None
    tables: dict[str, list[dict[str, Any]]] = field(default_factory=dict)


@dataclass
class _Context:
    companies: dict[uuid.UUID, str]
    projects: dict[uuid.UUID, str]
    departments: dict[uuid.UUID, str]
    locations: dict[uuid.UUID, str]
    people: dict[uuid.UUID, tuple[str, str]]
    reports: dict[uuid.UUID, int]


Getter = Callable[[_Context, _Person], Any]


@dataclass(frozen=True)
class _Column:
    key: str
    label: str
    get: Getter


@dataclass(frozen=True)
class _ListField:
    key: str
    label: str
    get: Callable[[dict[str, Any]], Any]


@dataclass(frozen=True)
class _Section:
    key: str
    label: str
    columns: tuple[_Column, ...] = ()
    #: For a list section: what one entry is called in a header ("Family member").
    entry: str | None = None
    fields: tuple[_ListField, ...] = ()
    items: Callable[[_Person], list[dict[str, Any]]] | None = None

    @property
    def repeats(self) -> bool:
        return self.items is not None


def _manager(index: int) -> Getter:
    def get(ctx: _Context, p: _Person) -> Any:
        found = ctx.people.get(p.employee.manager_id) if p.employee.manager_id else None
        return found[index] if found else None

    return get


def _dotted(index: int) -> Getter:
    def get(ctx: _Context, p: _Person) -> Any:
        target = p.employee.dotted_line_manager_id
        found = ctx.people.get(target) if target else None
        return found[index] if found else None

    return get


def _named(names: str, attribute: str) -> Getter:
    """A company, project, department or location by name."""

    def get(ctx: _Context, p: _Person) -> Any:
        value = getattr(p.employee, attribute)
        return getattr(ctx, names).get(value) if value else None

    return get


def _personal(key: str) -> Getter:
    return lambda ctx, p: (p.profile.get("personal") or {}).get(key) or None


def _address(which: str, key: str) -> Getter:
    def get(ctx: _Context, p: _Person) -> Any:
        addresses = p.profile.get("addresses") or {}
        place = addresses.get(which) or {}
        value = place.get(key)
        return _state(value) if key == "state_code" else (value or None)

    return get


def _login_status(ctx: _Context, p: _Person) -> str:
    if p.login is None:
        return "No login"
    return "Active" if p.login[1] else "Disabled"


def _aadhaar(ctx: _Context, p: _Person) -> str | None:
    if p.aadhaar:
        return p.aadhaar
    last4 = p.employee.aadhaar_last4
    return f"Last 4 only: {last4}" if last4 else None


def _hiring(key: str, *, number: bool = False) -> Getter:
    def get(ctx: _Context, p: _Person) -> Any:
        value = (p.profile.get("hiring_snapshot") or {}).get(key)
        return _number(value) if number else (value or None)

    return get


def _list(key: str) -> Callable[[_Person], list[dict[str, Any]]]:
    return lambda p: [item for item in (p.profile.get(key) or []) if isinstance(item, dict)]


def _text(key: str) -> Callable[[dict[str, Any]], Any]:
    return lambda item: item.get(key) or None


def _relationship(item: dict[str, Any]) -> Any:
    return _label(_RELATIONSHIP, item.get("relationship"), item.get("relationship_text"))


_STANDARD: tuple[_Section, ...] = (
    _Section(
        "basic",
        "Basic details",
        (
            _Column("code", "Employee code", lambda c, p: p.employee.employee_code),
            _Column("name", "Name", lambda c, p: p.employee.full_name),
            _Column("first_name", "First name", lambda c, p: p.employee.first_name),
            _Column("last_name", "Last name", lambda c, p: p.employee.last_name),
            _Column("display_name", "Display name", lambda c, p: p.employee.display_name),
            _Column("gender", "Gender", lambda c, p: _label(_GENDER, p.employee.gender)),
            _Column("date_of_birth", "Date of birth", lambda c, p: p.employee.date_of_birth),
            _Column("work_email", "Work email", lambda c, p: p.employee.work_email),
            _Column("personal_email", "Personal email", lambda c, p: p.employee.personal_email),
            _Column("phone", "Mobile", lambda c, p: p.employee.phone),
        ),
    ),
    _Section(
        "job",
        "Job details",
        (
            _Column("designation", "Designation", lambda c, p: p.employee.designation),
            _Column("company", "Company", _named("companies", "legal_entity_id")),
            _Column("project", "Project", _named("projects", "project_id")),
            _Column("department", "Department", _named("departments", "department_id")),
            _Column("location", "Location", _named("locations", "location_id")),
            _Column("grade", "Grade", lambda c, p: p.employee.grade),
            _Column(
                "employment_type",
                "Employment type",
                lambda c, p: _label(_EMPLOYMENT_TYPE, p.employee.employment_type),
            ),
            _Column("status", "Status", lambda c, p: _label(_STATUS, p.employee.status)),
            _Column("date_of_joining", "Date of joining", lambda c, p: p.employee.date_of_joining),
            _Column(
                "probation_end_date", "Probation ends", lambda c, p: p.employee.probation_end_date
            ),
            _Column(
                "notice_period_days",
                "Notice period (days)",
                lambda c, p: p.employee.notice_period_days,
            ),
            _Column("date_of_exit", "Date of exit", lambda c, p: p.employee.date_of_exit),
            _Column("manager", "Manager", _manager(0)),
            _Column("manager_code", "Manager code", _manager(1)),
            _Column("dotted_manager", "Dotted-line manager", _dotted(0)),
            _Column("dotted_manager_code", "Dotted-line manager code", _dotted(1)),
            _Column(
                "direct_reports", "Direct reports", lambda c, p: c.reports.get(p.employee.id, 0)
            ),
            _Column(
                "portal_role",
                "Portal role",
                lambda c, p: _ROLE.get(p.login[0]) if p.login else None,
            ),
            _Column("login_status", "Login status", _login_status),
            _Column(
                "added_on",
                "Added on",
                lambda c, p: p.employee.created_at.date() if p.employee.created_at else None,
            ),
        ),
    ),
    _Section(
        "personal",
        "Personal",
        (
            _Column(
                "marital_status",
                "Marital status",
                lambda c, p: _label(_MARITAL, _personal("marital_status")(c, p)),
            ),
            _Column("blood_group", "Blood group", _personal("blood_group")),
            _Column("nationality", "Nationality", _personal("nationality")),
            _Column("father_name", "Father's name", _personal("father_name")),
            _Column("spouse_name", "Spouse's name", _personal("spouse_name")),
            _Column("alternate_phone", "Alternate mobile", _personal("alternate_phone")),
        ),
    ),
    _Section(
        "address",
        "Address",
        (
            _Column("cur_line1", "Current address line 1", _address("current", "line1")),
            _Column("cur_line2", "Current address line 2", _address("current", "line2")),
            _Column("cur_city", "Current city", _address("current", "city")),
            _Column("cur_district", "Current district", _address("current", "district")),
            _Column("cur_state", "Current state", _address("current", "state_code")),
            _Column("cur_pincode", "Current PIN code", _address("current", "pincode")),
            _Column(
                "perm_same",
                "Permanent same as current",
                lambda c, p: _yes_no(
                    (p.profile.get("addresses") or {}).get("permanent_same_as_current")
                ),
            ),
            _Column("perm_line1", "Permanent address line 1", _address("permanent", "line1")),
            _Column("perm_line2", "Permanent address line 2", _address("permanent", "line2")),
            _Column("perm_city", "Permanent city", _address("permanent", "city")),
            _Column("perm_district", "Permanent district", _address("permanent", "district")),
            _Column("perm_state", "Permanent state", _address("permanent", "state_code")),
            _Column("perm_pincode", "Permanent PIN code", _address("permanent", "pincode")),
        ),
    ),
    _Section(
        "statutory",
        "Statutory",
        (
            _Column("pan", "PAN", lambda c, p: p.pan),
            _Column("aadhaar", "Aadhaar", _aadhaar),
            _Column("uan", "UAN", lambda c, p: p.employee.uan),
            _Column("pf_number", "PF number", lambda c, p: p.employee.pf_number),
            _Column("esi_number", "ESI number", lambda c, p: p.employee.esi_number),
        ),
    ),
    _Section(
        "bank",
        "Bank",
        (
            _Column("bank_name", "Bank name", lambda c, p: p.account.get("bank_name")),
            _Column("account_number", "Account number", lambda c, p: p.account.get("number")),
            _Column("ifsc", "IFSC", lambda c, p: p.account.get("ifsc")),
            _Column("account_holder", "Name as per bank", lambda c, p: p.account.get("holder")),
        ),
    ),
    _Section(
        "emergency",
        "Emergency contacts",
        entry="Emergency contact",
        items=_list("emergency_contacts"),
        fields=(
            _ListField("name", "Name", _text("name")),
            _ListField("relationship", "Relationship", _relationship),
            _ListField("phone", "Phone", _text("phone")),
            _ListField("alternate_phone", "Alternate phone", _text("alternate_phone")),
            _ListField("landline", "Landline", _text("landline")),
            _ListField("email", "Email", _text("email")),
            _ListField("address", "Address", _text("address")),
            _ListField("is_primary", "Primary", lambda i: _yes_no(i.get("is_primary"))),
        ),
    ),
    _Section(
        "family",
        "Family members",
        entry="Family member",
        items=_list("family_members"),
        fields=(
            _ListField("name", "Name", _text("name")),
            _ListField("relationship", "Relationship", _relationship),
            _ListField("title", "Title", _text("title")),
            _ListField("date_of_birth", "Date of birth", lambda i: _day(i.get("date_of_birth"))),
            _ListField("gender", "Gender", lambda i: _label(_GENDER, i.get("gender"))),
            _ListField("occupation", "Occupation", _text("occupation")),
            _ListField("phone", "Phone", _text("phone")),
            _ListField("blood_group", "Blood group", _text("blood_group")),
            _ListField("nationality", "Nationality", _text("nationality")),
            _ListField("marital_status", "Marital status", _text("marital_status")),
            _ListField("marriage_date", "Marriage date", lambda i: _day(i.get("marriage_date"))),
            _ListField("is_dependent", "Dependant", lambda i: _yes_no(i.get("is_dependent"))),
            _ListField(
                "health_insurance", "Health insurance", lambda i: _yes_no(i.get("health_insurance"))
            ),
            _ListField("insurance_number", "Insurance number", _text("insurance_number")),
            _ListField("remarks", "Remarks", _text("remarks")),
        ),
    ),
    _Section(
        "nominees",
        "Nominees",
        entry="Nominee",
        items=_list("nominees"),
        fields=(
            _ListField("name", "Name", _text("name")),
            _ListField("relationship", "Relationship", _relationship),
            _ListField(
                "purpose",
                "Nominated for",
                lambda i: _label(_PURPOSE, i.get("purpose"), i.get("purpose_text")),
            ),
            _ListField("share_percent", "Share (%)", lambda i: _number(i.get("share_percent"))),
            _ListField("date_of_birth", "Date of birth", lambda i: _day(i.get("date_of_birth"))),
            _ListField("guardian_name", "Guardian's name", _text("guardian_name")),
            _ListField("phone", "Phone", _text("phone")),
            _ListField("address", "Address", _text("address")),
            _ListField("city", "City", _text("city")),
            _ListField("country", "Country", _text("country")),
            _ListField("remarks", "Remarks", _text("remarks")),
        ),
    ),
    _Section(
        "education",
        "Education",
        entry="Education",
        items=_list("education"),
        fields=(
            _ListField("level", "Level", lambda i: _label(_EDUCATION_LEVEL, i.get("level"))),
            _ListField("qualification", "Qualification", _text("qualification")),
            _ListField("course", "Course", _text("course")),
            _ListField("stream", "Stream", _text("stream")),
            _ListField("specialisation", "Specialisation", _text("specialisation")),
            _ListField("institution", "Institution", _text("institution")),
            _ListField("board_or_university", "Board or university", _text("board_or_university")),
            _ListField("nature_of_course", "Full or part time", _text("nature_of_course")),
            _ListField("status", "Status", _text("status")),
            _ListField("start_year", "Start year", lambda i: _number(i.get("start_year"))),
            _ListField("end_year", "End year", lambda i: _number(i.get("end_year"))),
            _ListField("from_date", "From", lambda i: _day(i.get("from_date"))),
            _ListField("to_date", "To", lambda i: _day(i.get("to_date"))),
            _ListField("year", "Year of passing", _text("year")),
            _ListField(
                "date_of_passing", "Date of passing", lambda i: _day(i.get("date_of_passing"))
            ),
            _ListField("percentage", "Percentage", _text("percentage")),
            _ListField("grade", "Grade or score", _text("grade")),
            _ListField("duration", "Duration", _text("duration")),
            _ListField(
                "is_highest", "Highest qualification", lambda i: _yes_no(i.get("is_highest"))
            ),
            _ListField("remarks", "Remarks", _text("remarks")),
        ),
    ),
    _Section(
        "experience",
        "Previous employment",
        entry="Previous job",
        items=_list("previous_employment"),
        fields=(
            _ListField("employer", "Employer", _text("employer")),
            _ListField("designation", "Designation", _text("designation")),
            _ListField("from_date", "From", lambda i: _day(i.get("from_date"))),
            _ListField("to_date", "To", lambda i: _day(i.get("to_date"))),
            _ListField("location", "Location", _text("location")),
            _ListField("reason_for_leaving", "Reason for leaving", _text("reason_for_leaving")),
        ),
    ),
    _Section(
        "hiring",
        "From the offer",
        (
            _Column("offer_code", "Offer", _hiring("offer_code")),
            _Column("hiring_source", "Hiring source", _hiring("source")),
            _Column(
                "notice_at_offer",
                "Notice period at offer (days)",
                _hiring("notice_period_days_at_offer", number=True),
            ),
            _Column("expected_ctc", "Expected CTC", _hiring("expected_ctc", number=True)),
            _Column(
                "ctc_at_offer",
                "CTC at offer (previous employer)",
                _hiring("current_ctc_at_offer", number=True),
            ),
        ),
    ),
)

#: Columns whose values are personal identifiers, named in the audit row.
_IDENTIFIER_KEYS = frozenset({"pan", "aadhaar", "account_number"})
_PROFILE_SECTIONS = frozenset(
    {"personal", "address", "emergency", "family", "nominees", "education", "experience", "hiring"}
)


# ═══════════════════════════════════════════════════════════════════════════
# The workspace's own fields and tables
# ═══════════════════════════════════════════════════════════════════════════


def _custom_field_sections(definitions: list[CustomFieldDefinition]) -> list[_Section]:
    """One section per profile tab the fields sit on ("factoHR", or Custom fields)."""
    by_tab: dict[str, list[_Column]] = {}
    for definition in definitions:
        key, kind = definition.key, definition.field_type

        def get(ctx: _Context, p: _Person, key: str = key, kind: str = kind) -> Any:
            return _custom_value(kind, (p.employee.custom_fields or {}).get(key))

        by_tab.setdefault(definition.tab or "", []).append(
            _Column(f"cf.{key}", definition.label, get)
        )
    return [
        _Section(f"cf:{tab}" if tab else "cf", tab or "Custom fields", tuple(columns))
        for tab, columns in sorted(by_tab.items(), key=lambda item: (item[0] != "", item[0]))
    ]


def _custom_table_sections(tables: list[CustomTableDefinition]) -> list[_Section]:
    sections: list[_Section] = []
    for table in tables:
        fields: list[_ListField] = []
        for column in table.columns or []:
            col_key, kind = column.get("key"), column.get("field_type", "text")
            if not col_key:
                continue

            def get(item: dict[str, Any], col_key: str = col_key, kind: str = kind) -> Any:
                return _custom_value(kind, item.get(col_key))

            fields.append(_ListField(col_key, column.get("label") or col_key, get))
        table_key = table.key

        def items(p: _Person, table_key: str = table_key) -> list[dict[str, Any]]:
            return p.tables.get(table_key, [])

        sections.append(
            _Section(
                f"ct.{table.key}", table.label, entry=table.label, fields=tuple(fields), items=items
            )
        )
    return sections


async def _sections(session: AsyncSession, user: User) -> list[_Section]:
    definitions = (
        await session.scalars(
            tenant_scoped(
                select(CustomFieldDefinition)
                .where(
                    CustomFieldDefinition.entity == "employee",
                    CustomFieldDefinition.is_active.is_(True),
                )
                .order_by(CustomFieldDefinition.display_order, CustomFieldDefinition.label),
                CustomFieldDefinition,
                user,
            )
        )
    ).all()
    tables = (
        await session.scalars(
            tenant_scoped(
                select(CustomTableDefinition)
                .where(CustomTableDefinition.active.is_(True))
                .order_by(CustomTableDefinition.display_order, CustomTableDefinition.label),
                CustomTableDefinition,
                user,
            )
        )
    ).all()
    return [
        *_STANDARD,
        *_custom_field_sections(list(definitions)),
        *_custom_table_sections(list(tables)),
    ]


def _column_keys(section: _Section) -> list[str]:
    if section.repeats:
        return [f"{section.key}.{f.key}" for f in section.fields]
    return [c.key for c in section.columns]


def _require_hr(user: User) -> None:
    if not permissions.is_hr(user):
        raise permissions.forbidden("Only HR can export employee data.")


async def catalogue(session: AsyncSession, user: User) -> dict[str, Any]:
    """Every column the export can write, by section, for the column picker."""
    _require_hr(user)
    sections = await _sections(session, user)
    return {
        "sections": [
            {
                "key": section.key,
                "label": section.label,
                "repeats": section.repeats,
                "entry": section.entry,
                "columns": [
                    {"key": key, "label": label}
                    for key, label in zip(
                        _column_keys(section),
                        [f.label for f in section.fields]
                        if section.repeats
                        else [c.label for c in section.columns],
                        strict=True,
                    )
                ],
            }
            for section in sections
        ]
    }


# ═══════════════════════════════════════════════════════════════════════════
# The export
# ═══════════════════════════════════════════════════════════════════════════


@dataclass(frozen=True)
class ExportFilters:
    """The directory's filters, as on screen."""

    search: str | None = None
    department_id: uuid.UUID | None = None
    location_id: uuid.UUID | None = None
    manager_id: uuid.UUID | None = None
    legal_entity_id: uuid.UUID | None = None
    without_legal_entity: bool = False
    project_id: uuid.UUID | None = None
    without_project: bool = False
    status_filter: str | None = None
    current_only: bool = False

    def audit(self) -> dict[str, Any]:
        def text(value: Any) -> Any:
            return str(value) if isinstance(value, uuid.UUID) else value

        return {
            "search": self.search,
            "legal_entity_id": text(self.legal_entity_id),
            "without_legal_entity": self.without_legal_entity,
            "project_id": text(self.project_id),
            "without_project": self.without_project,
            "department_id": text(self.department_id),
            "location_id": text(self.location_id),
            "manager_id": text(self.manager_id),
            "status": self.status_filter,
            "current": self.current_only,
        }


def _decrypt(value: str | None, read: Callable[[str], str | None]) -> str | None:
    if not value:
        return None
    try:
        return read(value) or None
    except EncryptionError:
        return UNREADABLE


async def _names(session: AsyncSession, user: User, model: Any) -> dict[uuid.UUID, str]:
    rows = await session.execute(tenant_scoped(select(model.id, model.name), model, user))
    return dict(rows.all())  # type: ignore[arg-type]


async def _context(session: AsyncSession, user: User) -> _Context:
    everyone = (
        await session.scalars(
            tenant_scoped(select(Employee), Employee, user).options(
                load_only(
                    Employee.id,
                    Employee.employee_code,
                    Employee.first_name,
                    Employee.last_name,
                    Employee.display_name,
                    Employee.manager_id,
                    Employee.deleted_at,
                )
            )
        )
    ).all()
    reports: dict[uuid.UUID, int] = {}
    for person in everyone:
        if person.manager_id and person.deleted_at is None:
            reports[person.manager_id] = reports.get(person.manager_id, 0) + 1
    return _Context(
        companies=await _names(session, user, LegalEntity),
        projects=await _names(session, user, Project),
        departments=await _names(session, user, Department),
        locations=await _names(session, user, Location),
        people={p.id: (p.full_name, p.employee_code) for p in everyone},
        reports=reports,
    )


async def _people(
    session: AsyncSession,
    user: User,
    filters: ExportFilters,
    employee_ids: list[uuid.UUID] | None,
    wanted: set[str],
    needed: set[str],
) -> tuple[list[_Person], int]:
    """The people to export, each with what the chosen columns need.

    `needed` is the keys of the sections with a chosen column: a section
    nobody chose is not read at all.
    """
    plain = [attr for attr in Employee.__mapper__.column_attrs if attr.key not in _ENCRYPTED]
    stmt = tenant_scoped(select(Employee), Employee, user).options(
        load_only(*(getattr(Employee, attr.key) for attr in plain))
    )
    # The encrypted columns raw, so one unreadable value does not sink the file.
    stmt = stmt.add_columns(
        type_coerce(Employee.pan, Text()).label("pan_raw"),
        type_coerce(Employee.aadhaar_number, Text()).label("aadhaar_raw"),
        type_coerce(Employee.bank_account_number, Text()).label("account_raw"),
    )
    stmt = filter_directory(
        stmt,
        search=filters.search,
        department_id=filters.department_id,
        location_id=filters.location_id,
        manager_id=filters.manager_id,
        legal_entity_id=filters.legal_entity_id,
        without_legal_entity=filters.without_legal_entity,
        project_id=filters.project_id,
        without_project=filters.without_project,
        status_filter=filters.status_filter,
        current_only=filters.current_only,
    )
    if employee_ids is not None:
        stmt = stmt.where(Employee.id.in_(employee_ids))
    rows = (
        await session.execute(
            stmt.order_by(Employee.first_name, Employee.last_name, Employee.id).limit(
                EXPORT_LIMIT + 1
            )
        )
    ).all()
    total = len(rows)
    rows = rows[:EXPORT_LIMIT]

    decrypt_field = field_reader()
    people: list[_Person] = []
    for employee, pan_raw, aadhaar_raw, account_raw in rows:
        person = _Person(employee=employee)
        if "pan" in wanted:
            person.pan = _decrypt(pan_raw, decrypt_field)
        if "aadhaar" in wanted:
            person.aadhaar = _decrypt(aadhaar_raw, decrypt_field)
        # An account kept on the record itself, from before salary accounts had
        # their own table; the salary account below replaces it when there is one.
        if employee.bank_ifsc:
            person.account["ifsc"] = employee.bank_ifsc
        if account_raw and "account_number" in wanted:
            person.account["number"] = _decrypt(account_raw, decrypt_field)
        people.append(person)
    by_id = {p.employee.id: p for p in people}
    ids = list(by_id)
    if not ids:
        return people, total

    if needed & _PROFILE_SECTIONS:
        profiles = (
            await session.scalars(
                tenant_scoped(
                    select(EmployeeProfile).where(EmployeeProfile.employee_id.in_(ids)),
                    EmployeeProfile,
                    user,
                )
            )
        ).all()
        for profile in profiles:
            by_id[profile.employee_id].profile = {
                "personal": profile.personal or {},
                "addresses": profile.addresses or {},
                "emergency_contacts": profile.emergency_contacts or [],
                "family_members": profile.family_members or [],
                "nominees": profile.nominees or [],
                "education": profile.education or [],
                "previous_employment": profile.previous_employment or [],
                "hiring_snapshot": profile.hiring_snapshot or {},
            }

    if "bank" in needed:
        accounts = (
            await session.scalars(
                tenant_scoped(
                    select(EmployeeBankAccount)
                    .where(
                        EmployeeBankAccount.employee_id.in_(ids),
                        EmployeeBankAccount.is_primary.is_(True),
                        EmployeeBankAccount.active.is_(True),
                    )
                    .order_by(EmployeeBankAccount.created_at.desc()),
                    EmployeeBankAccount,
                    user,
                )
            )
        ).all()
        decrypt_secret = secret_reader(user.tenant_id)
        seen: set[uuid.UUID] = set()
        for account in accounts:
            if account.employee_id in seen:
                continue
            seen.add(account.employee_id)
            by_id[account.employee_id].account = {
                "number": _decrypt(account.account_number, decrypt_secret)
                if "account_number" in wanted
                else None,
                "ifsc": account.ifsc,
                "holder": account.account_holder_name,
                "bank_name": account.bank_name,
            }

    if {"portal_role", "login_status"} & wanted:
        logins = (
            await session.execute(
                tenant_scoped(
                    select(User.employee_id, User.role, User.is_active)
                    .where(User.employee_id.in_(ids))
                    .order_by(User.is_active.desc(), User.created_at),
                    User,
                    user,
                )
            )
        ).all()
        for employee_id, role, active in logins:
            person = by_id.get(employee_id)
            if person is not None and person.login is None:
                person.login = (role, bool(active))

    table_keys = {key.split(".")[1] for key in wanted if key.startswith("ct.")}
    if table_keys:
        tables = {
            t.id: t.key
            for t in (
                await session.scalars(
                    tenant_scoped(
                        select(CustomTableDefinition).where(
                            CustomTableDefinition.key.in_(table_keys)
                        ),
                        CustomTableDefinition,
                        user,
                    )
                )
            ).all()
        }
        if tables:
            rows_ = (
                await session.scalars(
                    tenant_scoped(
                        select(CustomTableRow)
                        .where(
                            CustomTableRow.table_id.in_(list(tables)),
                            CustomTableRow.employee_id.in_(ids),
                        )
                        .order_by(CustomTableRow.created_at),
                        CustomTableRow,
                        user,
                    )
                )
            ).all()
            for row in rows_:
                by_id[row.employee_id].tables.setdefault(tables[row.table_id], []).append(
                    dict(row.values or {})
                )
    return people, total


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class ExportFile:
    content: bytes
    filename: str
    media_type: str
    rows: int


@dataclass
class _Plan:
    """The header and the way to fill each cell, chosen columns only."""

    headers: list[str] = field(default_factory=list)
    cells: list[Callable[[_Context, _Person], Any]] = field(default_factory=list)
    #: Columns written as text even when they look like numbers.
    text_columns: set[int] = field(default_factory=set)


#: Kept as text in Excel, so a long number keeps its digits and its leading zeros.
_TEXT_KEYS = frozenset(
    {
        "code",
        "phone",
        "alternate_phone",
        "pan",
        "aadhaar",
        "uan",
        "pf_number",
        "esi_number",
        "account_number",
        "ifsc",
        "cur_pincode",
        "perm_pincode",
        "manager_code",
        "dotted_manager_code",
    }
)


def _plan(sections: list[_Section], wanted: set[str], people: list[_Person]) -> _Plan:
    plan = _Plan()
    for section in sections:
        if not section.repeats:
            for column in section.columns:
                if column.key in wanted:
                    if column.key in _TEXT_KEYS:
                        plan.text_columns.add(len(plan.headers))
                    plan.headers.append(column.label)
                    plan.cells.append(column.get)
            continue
        chosen = [f for f in section.fields if f"{section.key}.{f.key}" in wanted]
        if not chosen:
            continue
        assert section.items is not None
        most = max((len(section.items(p)) for p in people), default=0)
        for index in range(max(1, most)):
            for list_field in chosen:
                if list_field.key in {"phone", "alternate_phone", "landline"}:
                    plan.text_columns.add(len(plan.headers))
                plan.headers.append(f"{section.entry} {index + 1} - {list_field.label}")

                def cell(
                    ctx: _Context,
                    p: _Person,
                    items: Callable[[_Person], list[dict[str, Any]]] = section.items,
                    index: int = index,
                    get: Callable[[dict[str, Any]], Any] = list_field.get,
                ) -> Any:
                    entries = items(p)
                    return get(entries[index]) if index < len(entries) else None

                plan.cells.append(cell)
    return plan


def _csv(plan: _Plan, ctx: _Context, people: list[_Person]) -> bytes:
    def show(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, date):
            return value.strftime("%d-%m-%Y")
        text = str(value)
        return "'" + text if text.startswith(_FORMULA_PREFIXES) else text

    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([show(header) for header in plan.headers])
    for person in people:
        writer.writerow([show(cell(ctx, person)) for cell in plan.cells])
    # The byte-order mark makes Excel read the file as UTF-8.
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _xlsx(plan: _Plan, ctx: _Context, people: list[_Person]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook(write_only=True)
    sheet = book.create_sheet("Employees")
    widths = [min(max(len(h) + 2, 12), 40) for h in plan.headers]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C2"
    if plan.headers:
        last = get_column_letter(len(plan.headers))
        sheet.auto_filter.ref = f"A1:{last}{len(people) + 1}"

    fill = PatternFill("solid", fgColor="7D0712")  # AskMira maroon
    font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(vertical="center", wrap_text=True)
    header_row = []
    for header in plan.headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.fill, cell.font, cell.alignment = fill, font, wrap
        header_row.append(cell)
    sheet.append(header_row)

    # Most cells go in as plain values; only those that need a format (dates,
    # numbers kept as text) or must not become a formula get a cell object.
    for person in people:
        row: list[Any] = []
        for index, get in enumerate(plan.cells):
            value = get(ctx, person)
            if value is None:
                row.append(None)
                continue
            if isinstance(value, datetime):
                value = value.date()
            as_text = index in plan.text_columns
            if as_text:
                value = str(value)
            if isinstance(value, str):
                if not as_text and not value.startswith("="):
                    row.append(value)
                    continue
                cell = WriteOnlyCell(sheet, value=value)
                # Never a formula: a value starting "=" is shown, not run.
                cell.data_type = "s"
                if as_text:
                    cell.number_format = "@"
            elif isinstance(value, date):
                cell = WriteOnlyCell(sheet, value=value)
                cell.number_format = "DD-MM-YYYY"
            else:
                row.append(value)
                continue
            row.append(cell)
        sheet.append(row)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def export_employees(
    session: AsyncSession,
    user: User,
    filters: ExportFilters,
    *,
    columns: list[str] | None = None,
    employee_ids: list[uuid.UUID] | None = None,
    file_type: FileType = "xlsx",
    request: Request | None = None,
) -> ExportFile:
    """The chosen columns for the chosen people, as a file.

    `columns` None means every column. `employee_ids` narrows to people ticked
    on screen (within the filters). HR only, and recorded as an export.
    """
    _require_hr(user)
    sections = await _sections(session, user)
    every = [key for section in sections for key in _column_keys(section)]
    if columns is None:
        wanted = set(every)
    else:
        wanted = set(columns) & set(every)
        if not wanted:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Choose at least one column to export.",
            )

    needed = {
        section.key for section in sections if any(k in wanted for k in _column_keys(section))
    }
    people, total = await _people(session, user, filters, employee_ids, wanted, needed)
    ctx = await _context(session, user)
    plan = _plan(sections, wanted, people)
    content = _xlsx(plan, ctx, people) if file_type == "xlsx" else _csv(plan, ctx, people)

    chosen_sections = [
        section.label for section in sections if any(key in wanted for key in _column_keys(section))
    ]
    kind = "Excel" if file_type == "xlsx" else "CSV"
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        actor=user,
        action=AuditAction.EXPORT,
        entity_type="employee_directory",
        summary=(
            f"Exported employee data ({len(people)} people, {len(plan.headers)} columns) as {kind}"
        ),
        changes={
            "rows": len(people),
            "more_than_limit": total > EXPORT_LIMIT,
            "format": file_type,
            "columns": len(wanted),
            "all_columns": len(wanted) == len(every),
            "sections": chosen_sections,
            "identifiers": sorted(_IDENTIFIER_KEYS & wanted),
            "ticked_people": len(employee_ids) if employee_ids is not None else None,
            **filters.audit(),
        },
        request=request,
    )
    from app.services.attendance.service import local_today

    stamp = (await local_today(session, user)).isoformat()
    return ExportFile(
        content=content,
        filename=f"employees-{stamp}.{file_type}",
        media_type=_XLSX if file_type == "xlsx" else "text/csv; charset=utf-8",
        rows=len(people),
    )
