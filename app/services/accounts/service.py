"""Client accounts: the accounts, what was billed and spent, who works on each, the P&L.

Everything here is HR's or finance's, behind the payroll area: an account's
P&L carries what people are paid. A sub-admin needs the payroll rights the
request path asks for (view to read, create/edit/delete to change), and one
narrowed to some departments sees only their people's cost — the P&L says so.

Every change writes its audit row in the same transaction.
"""

from __future__ import annotations

import csv
import io
import re
import uuid
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any

from fastapi import HTTPException, Request, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import and_, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import hr_scope, permissions
from app.core.deps import tenant_scoped
from app.models.accounts import AccountAllocation, AccountEntry, AccountEntryKind, ClientAccount
from app.models.audit import AuditAction
from app.models.employee import Employee
from app.models.org import Department, LegalEntity
from app.models.payroll import PayrollRun, PayrollRunStatus, Payslip
from app.models.user import User
from app.schemas.accounts import (
    AccountEntryCreate,
    AccountEntryRead,
    AccountEntryUpdate,
    AccountImportRead,
    AccountImportRowRead,
    AccountPnlMonth,
    AccountPnlPerson,
    AccountPnlRead,
    AccountPnlRow,
    AllocationCreate,
    AllocationEnd,
    AllocationRead,
    ClientAccountCreate,
    ClientAccountRead,
    ClientAccountUpdate,
    PnlDetailRead,
)
from app.services import audit
from app.services.accounts import pnl as math
from app.services.attendance.service import local_today
from app.services.employee_import.fields import read_date, text

__all__ = [
    "account_pnl",
    "create_account",
    "create_allocation",
    "create_entry",
    "delete_account",
    "delete_allocation",
    "delete_entry",
    "end_allocation",
    "find_accounts",
    "import_entries",
    "list_accounts",
    "list_allocations",
    "list_entries",
    "pnl",
    "template",
    "update_account",
    "update_entry",
]

#: Runs whose payslips are not a cost: cancelled, or failed before anything was computed.
_NOT_A_COST = (PayrollRunStatus.CANCELLED, PayrollRunStatus.FAILED)
_FINAL = (PayrollRunStatus.APPROVED, PayrollRunStatus.PAID)
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UPLOAD_ROWS = 5000
ZERO = Decimal(0)


def _not_found(detail: str) -> HTTPException:
    return HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=detail)


def _unprocessable(detail: str) -> HTTPException:
    return HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=detail)


def _conflict(detail: str) -> HTTPException:
    return HTTPException(status_code=status.HTTP_409_CONFLICT, detail=detail)


def _require(user: User) -> None:
    permissions.require_hr(
        user,
        detail="Client accounts and their P&L are for admins and sub-admins with payroll rights.",
    )


def _month(value: str) -> tuple[int, int]:
    return int(value[:4]), int(value[5:7])


def _name(employee: Employee) -> str:
    return employee.full_name


def _percent(value: Decimal) -> str:
    """A share as people write it: 110, 33.5 (never 1.1E+2)."""
    text = f"{Decimal(value):f}"
    return text.rstrip("0").rstrip(".") if "." in text else text


# ── Accounts ─────────────────────────────────────────────────────────────────


async def _account(session: AsyncSession, user: User, account_id: uuid.UUID) -> ClientAccount:
    row = await session.scalar(
        tenant_scoped(
            select(ClientAccount).where(ClientAccount.id == account_id), ClientAccount, user
        )
    )
    if row is None:
        raise _not_found("No such account.")
    return row


async def _check_links(
    session: AsyncSession,
    user: User,
    *,
    legal_entity_id: uuid.UUID | None,
    owner_employee_id: uuid.UUID | None,
) -> None:
    if legal_entity_id is not None:
        found = await session.scalar(
            tenant_scoped(
                select(LegalEntity.id).where(LegalEntity.id == legal_entity_id), LegalEntity, user
            )
        )
        if found is None:
            raise _unprocessable("That company is not in this workspace.")
    if owner_employee_id is not None:
        found = await session.scalar(
            tenant_scoped(
                select(Employee.id).where(
                    Employee.id == owner_employee_id, Employee.deleted_at.is_(None)
                ),
                Employee,
                user,
            )
        )
        if found is None:
            raise _unprocessable("That person is not in this workspace.")


async def _reads(
    session: AsyncSession, user: User, accounts: list[ClientAccount], today: date
) -> list[ClientAccountRead]:
    if not accounts:
        return []
    ids = [a.id for a in accounts]
    entity_ids = {a.legal_entity_id for a in accounts if a.legal_entity_id}
    owner_ids = {a.owner_employee_id for a in accounts if a.owner_employee_id}
    entities: dict[uuid.UUID, str] = {}
    if entity_ids:
        entities = {
            row.id: row.name
            for row in await session.execute(
                tenant_scoped(
                    select(LegalEntity.id, LegalEntity.name).where(LegalEntity.id.in_(entity_ids)),
                    LegalEntity,
                    user,
                )
            )
        }
    owners: dict[uuid.UUID, str] = {}
    if owner_ids:
        owners = {
            row.id: _name(row)
            for row in await session.scalars(
                tenant_scoped(select(Employee).where(Employee.id.in_(owner_ids)), Employee, user)
            )
        }
    mapped: dict[uuid.UUID, tuple[int, Decimal]] = {}
    for row in await session.execute(
        tenant_scoped(
            select(
                AccountAllocation.account_id,
                func.count(func.distinct(AccountAllocation.employee_id)),
                func.coalesce(func.sum(AccountAllocation.percent), 0),
            )
            .where(
                AccountAllocation.account_id.in_(ids),
                AccountAllocation.effective_from <= today,
                or_(
                    AccountAllocation.effective_to.is_(None),
                    AccountAllocation.effective_to >= today,
                ),
            )
            .group_by(AccountAllocation.account_id),
            AccountAllocation,
            user,
        )
    ):
        mapped[row[0]] = (int(row[1]), (Decimal(row[2]) / 100).quantize(math.PAISA))
    return [
        ClientAccountRead(
            id=a.id,
            code=a.code,
            name=a.name,
            client_name=a.client_name,
            legal_entity_id=a.legal_entity_id,
            legal_entity_name=entities.get(a.legal_entity_id) if a.legal_entity_id else None,
            owner_employee_id=a.owner_employee_id,
            owner_name=owners.get(a.owner_employee_id) if a.owner_employee_id else None,
            is_active=a.is_active,
            notes=a.notes,
            people_now=mapped.get(a.id, (0, ZERO))[0],
            mapped_fte=mapped.get(a.id, (0, ZERO))[1],
            created_at=a.created_at,
        )
        for a in accounts
    ]


async def list_accounts(
    session: AsyncSession, user: User, *, include_inactive: bool = True
) -> list[ClientAccountRead]:
    _require(user)
    stmt = select(ClientAccount).order_by(ClientAccount.is_active.desc(), ClientAccount.name)
    if not include_inactive:
        stmt = stmt.where(ClientAccount.is_active.is_(True))
    rows = list((await session.scalars(tenant_scoped(stmt, ClientAccount, user))).all())
    return await _reads(session, user, rows, await local_today(session, user))


async def find_accounts(session: AsyncSession, user: User, query: str) -> list[ClientAccount]:
    """Accounts whose code or name holds the words typed (for Ask Mira)."""
    stmt = select(ClientAccount).order_by(ClientAccount.name)
    words = [w for w in re.split(r"\s+", query.strip()) if w]
    for word in words:
        like = "%" + word.replace("\\", "").replace("%", "").replace("_", "") + "%"
        stmt = stmt.where(
            or_(
                ClientAccount.code.ilike(like),
                ClientAccount.name.ilike(like),
                ClientAccount.client_name.ilike(like),
            )
        )
    return list((await session.scalars(tenant_scoped(stmt, ClientAccount, user))).all())


async def _code_taken(
    session: AsyncSession, user: User, code: str, *, other_than: uuid.UUID | None = None
) -> bool:
    stmt = select(ClientAccount.id).where(func.upper(ClientAccount.code) == code.upper())
    if other_than is not None:
        stmt = stmt.where(ClientAccount.id != other_than)
    return await session.scalar(tenant_scoped(stmt, ClientAccount, user)) is not None


async def create_account(
    session: AsyncSession, user: User, payload: ClientAccountCreate, *, request: Request | None
) -> ClientAccountRead:
    _require(user)
    if await _code_taken(session, user, payload.code):
        raise _conflict(f"An account with the code {payload.code} already exists.")
    await _check_links(
        session,
        user,
        legal_entity_id=payload.legal_entity_id,
        owner_employee_id=payload.owner_employee_id,
    )
    row = ClientAccount(tenant_id=user.tenant_id, **payload.model_dump())
    session.add(row)
    await session.flush()
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.CREATE,
        entity_type="client_account",
        entity_id=row.id,
        actor=user,
        changes=payload.model_dump(mode="json"),
        summary=f"Client account {row.code} ({row.name}) added",
        request=request,
    )
    await session.commit()
    await session.refresh(row)
    return (await _reads(session, user, [row], await local_today(session, user)))[0]


async def update_account(
    session: AsyncSession,
    user: User,
    account_id: uuid.UUID,
    payload: ClientAccountUpdate,
    *,
    request: Request | None,
) -> ClientAccountRead:
    _require(user)
    row = await _account(session, user, account_id)
    changes = payload.model_dump(exclude_unset=True)
    await _check_links(
        session,
        user,
        legal_entity_id=changes.get("legal_entity_id"),
        owner_employee_id=changes.get("owner_employee_id"),
    )
    before = {key: getattr(row, key) for key in changes}
    for key, value in changes.items():
        if key == "name" and value is None:
            continue
        setattr(row, key, value)
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.UPDATE,
        entity_type="client_account",
        entity_id=row.id,
        actor=user,
        changes={
            "before": {k: str(v) if v is not None else None for k, v in before.items()},
            "after": {k: str(v) if v is not None else None for k, v in changes.items()},
        },
        summary=f"Client account {row.code} changed",
        request=request,
    )
    await session.commit()
    await session.refresh(row)
    return (await _reads(session, user, [row], await local_today(session, user)))[0]


async def delete_account(
    session: AsyncSession, user: User, account_id: uuid.UUID, *, request: Request | None
) -> None:
    _require(user)
    row = await _account(session, user, account_id)
    used = await session.scalar(
        tenant_scoped(
            select(func.count(AccountEntry.id)).where(AccountEntry.account_id == row.id),
            AccountEntry,
            user,
        )
    )
    mapped = await session.scalar(
        tenant_scoped(
            select(func.count(AccountAllocation.id)).where(AccountAllocation.account_id == row.id),
            AccountAllocation,
            user,
        )
    )
    if used or mapped:
        raise _conflict(
            "This account has billing, costs or people mapped to it, so its P&L would be lost. "
            "Mark it inactive instead."
        )
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.DELETE,
        entity_type="client_account",
        entity_id=row.id,
        actor=user,
        changes={"code": row.code, "name": row.name},
        summary=f"Client account {row.code} deleted",
        request=request,
    )
    await session.delete(row)
    await session.commit()


# ── Billing and costs ───────────────────────────────────────────────────────


def _entry_read(row: AccountEntry, account: ClientAccount) -> AccountEntryRead:
    return AccountEntryRead(
        id=row.id,
        account_id=row.account_id,
        account_code=account.code,
        account_name=account.name,
        kind=row.kind,
        month=math.month_key(row.period_year, row.period_month),
        amount=Decimal(row.amount),
        reference=row.reference,
        entry_date=row.entry_date,
        notes=row.notes,
        created_at=row.created_at,
    )


async def _not_after_this_month(session: AsyncSession, user: User, month: str) -> None:
    today = await local_today(session, user)
    if _month(month) > (today.year, today.month):
        raise _unprocessable(
            "Billing and costs are recorded for this month or earlier, once they have happened."
        )


async def list_entries(
    session: AsyncSession,
    user: User,
    *,
    account_id: uuid.UUID | None = None,
    kind: str | None = None,
    start: str | None = None,
    end: str | None = None,
) -> list[AccountEntryRead]:
    _require(user)
    stmt = (
        select(AccountEntry, ClientAccount)
        .join(ClientAccount, ClientAccount.id == AccountEntry.account_id)
        .order_by(
            AccountEntry.period_year.desc(),
            AccountEntry.period_month.desc(),
            ClientAccount.name,
            AccountEntry.created_at.desc(),
        )
    )
    if account_id is not None:
        stmt = stmt.where(AccountEntry.account_id == account_id)
    if kind in AccountEntryKind.ALL:
        stmt = stmt.where(AccountEntry.kind == kind)
    if start:
        year, month = _month(start)
        stmt = stmt.where(
            or_(
                AccountEntry.period_year > year,
                and_(AccountEntry.period_year == year, AccountEntry.period_month >= month),
            )
        )
    if end:
        year, month = _month(end)
        stmt = stmt.where(
            or_(
                AccountEntry.period_year < year,
                and_(AccountEntry.period_year == year, AccountEntry.period_month <= month),
            )
        )
    rows = (await session.execute(tenant_scoped(stmt, AccountEntry, user))).all()
    return [_entry_read(entry, account) for entry, account in rows]


async def _entry(session: AsyncSession, user: User, entry_id: uuid.UUID) -> AccountEntry:
    row = await session.scalar(
        tenant_scoped(select(AccountEntry).where(AccountEntry.id == entry_id), AccountEntry, user)
    )
    if row is None:
        raise _not_found("No such entry.")
    return row


async def create_entry(
    session: AsyncSession, user: User, payload: AccountEntryCreate, *, request: Request | None
) -> AccountEntryRead:
    _require(user)
    account = await _account(session, user, payload.account_id)
    await _not_after_this_month(session, user, payload.month)
    year, month = _month(payload.month)
    row = AccountEntry(
        tenant_id=user.tenant_id,
        account_id=account.id,
        kind=payload.kind,
        period_year=year,
        period_month=month,
        amount=payload.amount,
        reference=payload.reference,
        entry_date=payload.entry_date,
        notes=payload.notes,
    )
    session.add(row)
    await session.flush()
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.CREATE,
        entity_type="account_entry",
        entity_id=row.id,
        actor=user,
        changes=payload.model_dump(mode="json"),
        summary=(
            f"{'Billing' if payload.kind == 'billing' else 'Cost'} of {payload.amount} for "
            f"{account.code}, {payload.month}"
        ),
        request=request,
    )
    await session.commit()
    await session.refresh(row)
    return _entry_read(row, account)


async def update_entry(
    session: AsyncSession,
    user: User,
    entry_id: uuid.UUID,
    payload: AccountEntryUpdate,
    *,
    request: Request | None,
) -> AccountEntryRead:
    _require(user)
    row = await _entry(session, user, entry_id)
    account = await _account(session, user, row.account_id)
    changes = payload.model_dump(exclude_unset=True)
    before = {
        "month": math.month_key(row.period_year, row.period_month),
        "amount": str(row.amount),
        "reference": row.reference,
        "entry_date": row.entry_date.isoformat() if row.entry_date else None,
        "notes": row.notes,
    }
    if changes.get("month"):
        await _not_after_this_month(session, user, changes["month"])
        row.period_year, row.period_month = _month(changes["month"])
    if "amount" in changes and changes["amount"] is not None:
        try:
            # The same rules as a new entry: two decimals, not zero, a cost never negative.
            AccountEntryCreate(
                account_id=row.account_id,
                kind=row.kind,
                month=math.month_key(row.period_year, row.period_month),
                amount=changes["amount"],
            )
        except ValidationError as exc:
            raise _unprocessable(exc.errors()[0]["msg"].removeprefix("Value error, ")) from exc
        row.amount = changes["amount"]
    for key in ("reference", "entry_date", "notes"):
        if key in changes:
            setattr(row, key, changes[key])
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.UPDATE,
        entity_type="account_entry",
        entity_id=row.id,
        actor=user,
        changes={"before": before, "after": payload.model_dump(mode="json", exclude_unset=True)},
        summary=f"Entry for {account.code} changed",
        request=request,
    )
    await session.commit()
    await session.refresh(row)
    return _entry_read(row, account)


async def delete_entry(
    session: AsyncSession, user: User, entry_id: uuid.UUID, *, request: Request | None
) -> None:
    _require(user)
    row = await _entry(session, user, entry_id)
    account = await _account(session, user, row.account_id)
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.DELETE,
        entity_type="account_entry",
        entity_id=row.id,
        actor=user,
        changes={
            "account": account.code,
            "kind": row.kind,
            "month": math.month_key(row.period_year, row.period_month),
            "amount": str(row.amount),
            "reference": row.reference,
        },
        summary=f"Entry of {row.amount} for {account.code} deleted",
        request=request,
    )
    await session.delete(row)
    await session.commit()


# ── Upload ──────────────────────────────────────────────────────────────────

_HEADERS = ("Account code", "Type", "Month", "Amount", "Invoice no. / cost head", "Date", "Notes")
_HEADER_KEYS = {
    "account code": "code",
    "account": "code",
    "code": "code",
    "type": "kind",
    "month": "month",
    "period": "month",
    "amount": "amount",
    "invoice no. / cost head": "reference",
    "invoice no": "reference",
    "invoice number": "reference",
    "reference": "reference",
    "cost head": "reference",
    "date": "date",
    "invoice date": "date",
    "notes": "notes",
    "remarks": "notes",
}
_KINDS = {
    "billing": "billing",
    "billed": "billing",
    "invoice": "billing",
    "revenue": "billing",
    "credit note": "billing",
    "cost": "cost",
    "expense": "cost",
    "direct cost": "cost",
}
_MONTHS = {
    name: number
    for number, names in enumerate(
        (
            ("jan", "january"),
            ("feb", "february"),
            ("mar", "march"),
            ("apr", "april"),
            ("may",),
            ("jun", "june"),
            ("jul", "july"),
            ("aug", "august"),
            ("sep", "sept", "september"),
            ("oct", "october"),
            ("nov", "november"),
            ("dec", "december"),
        ),
        start=1,
    )
    for name in names
}


def _key(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


async def template(session: AsyncSession, user: User) -> bytes:
    """An upload sheet: one row per invoice or cost, with the accounts listed to pick from."""
    _require(user)
    accounts = (
        await session.scalars(
            tenant_scoped(
                select(ClientAccount)
                .where(ClientAccount.is_active.is_(True))
                .order_by(ClientAccount.code),
                ClientAccount,
                user,
            )
        )
    ).all()
    book = Workbook()
    sheet = book.active
    sheet.title = "Billing and costs"
    sheet.append(list(_HEADERS))
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="7D0712")
    for cell in sheet[1]:
        cell.font = bold
        cell.fill = fill
    for column, width in zip("ABCDEFG", (16, 12, 12, 16, 26, 14, 30), strict=True):
        sheet.column_dimensions[column].width = width
    kinds = DataValidation(type="list", formula1='"Billing,Cost"', allow_blank=True)
    sheet.add_data_validation(kinds)
    kinds.add("B2:B5000")

    guide = book.create_sheet("Instructions")
    for line in (
        ["How to fill this sheet"],
        ["Account code: the code of an account in AskMira (see the list below)."],
        ["Type: Billing for an invoice raised (negative for a credit note)."],
        ["        Cost for a direct cost (travel, software, a subcontractor)."],
        ["Month: the month it belongs to, as 2026-09 or Sep 2026."],
        ["Amount: in rupees, at most two decimals. A cost cannot be negative."],
        ["Invoice no. / cost head: the invoice number, or what the cost was (Travel, Software)."],
        ["Date and Notes are optional."],
        [],
        ["Accounts"],
        *[[a.code, a.name] for a in accounts],
    ):
        guide.append(line)
    guide["A1"].font = Font(bold=True)
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 60
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _sheet_rows(data: bytes, filename: str) -> list[tuple[int, list[Any]]]:
    if not data:
        raise _unprocessable("The file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise _unprocessable("The file is larger than 5 MB.")
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            decoded = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = data.decode("latin-1")
        return list(enumerate(csv.reader(io.StringIO(decoded)), start=1))
    if not name.endswith((".xlsx", ".xlsm")):
        raise _unprocessable("Upload the Excel (.xlsx) template, or a CSV file.")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds for a bad file
        raise _unprocessable("This is not an Excel (.xlsx) file AskMira can read.") from exc
    for sheet in book.worksheets:
        if _key(sheet.title) in {"instructions", "sample"}:
            continue
        rows = [
            (number, list(values))
            for number, values in enumerate(sheet.iter_rows(values_only=True), start=1)
        ]
        if any(_HEADER_KEYS.get(_key(v)) == "code" for _, values in rows[:10] for v in values):
            return rows
    raise _unprocessable("No sheet has an 'Account code' column. Start from the template.")


def _read_month(raw: Any) -> str:
    if isinstance(raw, date):
        return math.month_key(raw.year, raw.month)
    value = _key(raw)
    if not value:
        raise ValueError("Month is empty")
    match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})(?:[-/.]\d{1,2})?", value)
    if match:
        year, month = int(match.group(1)), int(match.group(2))
    else:
        match = re.fullmatch(r"(\d{1,2})[-/.](\d{4})", value)
        if match:
            month, year = int(match.group(1)), int(match.group(2))
        else:
            match = re.fullmatch(r"([a-z]+)[\s\-'/,]*(\d{2}|\d{4})", value)
            if not match or match.group(1) not in _MONTHS:
                raise ValueError(f"Month '{raw}': write it as 2026-09 or Sep 2026")
            month = _MONTHS[match.group(1)]
            year = int(match.group(2))
            year = 2000 + year if year < 100 else year
    if not 1 <= month <= 12 or year < 2000:
        raise ValueError(f"Month '{raw}' is not a month")
    return math.month_key(year, month)


def _read_amount(raw: Any) -> Decimal:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        raise ValueError("Amount is empty")
    if isinstance(raw, bool):
        raise ValueError(f"Amount '{raw}' is not a number")
    cleaned = str(raw).strip().replace(",", "").replace("₹", "").replace(" ", "")
    if re.fullmatch(r"\(\d+(\.\d+)?\)", cleaned):
        cleaned = "-" + cleaned[1:-1]
    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"Amount '{raw}' is not a number") from exc


async def import_entries(
    session: AsyncSession,
    user: User,
    data: bytes,
    filename: str,
    *,
    apply: bool,
    request: Request | None,
) -> AccountImportRead:
    """Check an upload. With `apply`, save every row at once, only if none has an error."""
    _require(user)
    rows = _sheet_rows(data, filename)
    header_at = next(
        (
            index
            for index, (_, values) in enumerate(rows[:10])
            if any(_HEADER_KEYS.get(_key(v)) == "code" for v in values)
        ),
        None,
    )
    if header_at is None:
        raise _unprocessable("No 'Account code' column was found. Start from the template.")
    columns = {
        _HEADER_KEYS[_key(v)]: position
        for position, v in enumerate(rows[header_at][1])
        if _key(v) in _HEADER_KEYS
    }
    missing = [
        label
        for label, key in (("Type", "kind"), ("Month", "month"), ("Amount", "amount"))
        if key not in columns
    ]
    if missing:
        raise _unprocessable(
            f"The sheet has no {', '.join(missing)} column. Start from the template."
        )
    body = rows[header_at + 1 :]
    if len(body) > MAX_UPLOAD_ROWS:
        raise _unprocessable(f"Upload at most {MAX_UPLOAD_ROWS} rows at a time.")

    accounts = {
        a.code.upper(): a
        for a in (
            await session.scalars(tenant_scoped(select(ClientAccount), ClientAccount, user))
        ).all()
    }
    today = await local_today(session, user)
    out: list[AccountImportRowRead] = []
    plans: list[tuple[ClientAccount, AccountEntryCreate]] = []

    def cell(values: list[Any], key: str) -> Any:
        position = columns.get(key)
        return values[position] if position is not None and position < len(values) else None

    for number, values in body:
        if not any(text(v) for v in values):
            out.append(
                AccountImportRowRead(
                    row=number,
                    status="skipped",
                    account_code=None,
                    kind=None,
                    month=None,
                    amount=None,
                    reference=None,
                    message=None,
                )
            )
            continue
        code = (text(cell(values, "code")) or "").upper()
        reference = text(cell(values, "reference"))
        try:
            if not code:
                raise ValueError("Account code is empty")
            account = accounts.get(code)
            if account is None:
                raise ValueError(f"No account has the code {code}")
            kind = _KINDS.get(_key(cell(values, "kind")))
            if kind is None:
                raise ValueError("Type must be Billing or Cost")
            month = _read_month(cell(values, "month"))
            if _month(month) > (today.year, today.month):
                raise ValueError(f"{month} has not happened yet")
            amount = _read_amount(cell(values, "amount"))
            entry_date = read_date(cell(values, "date"))
            payload = AccountEntryCreate(
                account_id=account.id,
                kind=kind,
                month=month,
                amount=amount,
                reference=reference,
                entry_date=entry_date,
                notes=text(cell(values, "notes")),
            )
        except ValidationError as exc:
            message = "; ".join(e["msg"].removeprefix("Value error, ") for e in exc.errors())
            out.append(
                AccountImportRowRead(
                    row=number,
                    status="error",
                    account_code=code or None,
                    kind=None,
                    month=None,
                    amount=None,
                    reference=reference,
                    message=message,
                )
            )
            continue
        except ValueError as exc:
            out.append(
                AccountImportRowRead(
                    row=number,
                    status="error",
                    account_code=code or None,
                    kind=None,
                    month=None,
                    amount=None,
                    reference=reference,
                    message=str(exc),
                )
            )
            continue
        plans.append((account, payload))
        out.append(
            AccountImportRowRead(
                row=number,
                status="ready",
                account_code=account.code,
                kind=payload.kind,
                month=payload.month,
                amount=payload.amount,
                reference=payload.reference,
                message=None,
            )
        )

    errors = sum(1 for r in out if r.status == "error")
    result = AccountImportRead(
        filename=filename,
        rows=out,
        ready=len(plans),
        errors=errors,
        skipped=sum(1 for r in out if r.status == "skipped"),
    )
    if not apply:
        return result
    if errors:
        raise _unprocessable(f"Fix the {errors} row(s) marked as errors first; nothing was saved.")
    if not plans:
        raise _unprocessable("The sheet has no rows to save.")
    for account, payload in plans:
        year, month = _month(payload.month)
        session.add(
            AccountEntry(
                tenant_id=user.tenant_id,
                account_id=account.id,
                kind=payload.kind,
                period_year=year,
                period_month=month,
                amount=payload.amount,
                reference=payload.reference,
                entry_date=payload.entry_date,
                notes=payload.notes,
            )
        )
    await session.flush()
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.CREATE,
        entity_type="account_entry_import",
        actor=user,
        changes={
            "filename": filename,
            "rows": len(plans),
            "accounts": sorted({a.code for a, _ in plans}),
            "billed": str(sum((p.amount for _, p in plans if p.kind == "billing"), ZERO)),
            "costs": str(sum((p.amount for _, p in plans if p.kind == "cost"), ZERO)),
        },
        summary=f"{len(plans)} billing and cost row(s) uploaded from {filename}",
        request=request,
    )
    await session.commit()
    result.saved = len(plans)
    return result


# ── People on accounts ──────────────────────────────────────────────────────


async def _allocation(
    session: AsyncSession, user: User, allocation_id: uuid.UUID
) -> AccountAllocation:
    row = await session.scalar(
        tenant_scoped(
            select(AccountAllocation).where(AccountAllocation.id == allocation_id),
            AccountAllocation,
            user,
        )
    )
    if row is None:
        raise _not_found("No such mapping.")
    return row


async def list_allocations(
    session: AsyncSession,
    user: User,
    *,
    account_id: uuid.UUID | None = None,
    employee_id: uuid.UUID | None = None,
    current_only: bool = False,
) -> list[AllocationRead]:
    _require(user)
    today = await local_today(session, user)
    stmt = (
        select(AccountAllocation, ClientAccount, Employee, Department.name)
        .join(ClientAccount, ClientAccount.id == AccountAllocation.account_id)
        .join(Employee, Employee.id == AccountAllocation.employee_id)
        .outerjoin(Department, Department.id == Employee.department_id)
        .order_by(Employee.first_name, Employee.last_name, AccountAllocation.effective_from.desc())
    )
    if account_id is not None:
        stmt = stmt.where(AccountAllocation.account_id == account_id)
    if employee_id is not None:
        stmt = stmt.where(AccountAllocation.employee_id == employee_id)
    if current_only:
        stmt = stmt.where(
            or_(AccountAllocation.effective_to.is_(None), AccountAllocation.effective_to >= today)
        )
    rows = (await session.execute(tenant_scoped(stmt, AccountAllocation, user))).all()
    return [
        AllocationRead(
            id=allocation.id,
            employee_id=employee.id,
            employee_code=employee.employee_code,
            employee_name=_name(employee),
            department_name=department,
            account_id=account.id,
            account_code=account.code,
            account_name=account.name,
            percent=Decimal(allocation.percent),
            effective_from=allocation.effective_from,
            effective_to=allocation.effective_to,
            active=allocation.effective_from <= today
            and (allocation.effective_to is None or allocation.effective_to >= today),
        )
        for allocation, account, employee, department in rows
    ]


async def _overbooked(
    session: AsyncSession,
    user: User,
    employee_id: uuid.UUID,
    start: date,
    end: date | None,
    percent: Decimal,
    *,
    other_than: uuid.UUID | None = None,
) -> tuple[date, Decimal] | None:
    """The first day the person's shares would add up to more than 100%, and the total."""
    stmt = select(AccountAllocation).where(
        AccountAllocation.employee_id == employee_id,
        or_(AccountAllocation.effective_to.is_(None), AccountAllocation.effective_to >= start),
    )
    if end is not None:
        stmt = stmt.where(AccountAllocation.effective_from <= end)
    if other_than is not None:
        stmt = stmt.where(AccountAllocation.id != other_than)
    others = (await session.scalars(tenant_scoped(stmt, AccountAllocation, user))).all()
    # The total only changes where a mapping starts, so those days are enough to check.
    days = sorted({start, *(o.effective_from for o in others if o.effective_from >= start)})
    for day in days:
        if end is not None and day > end:
            break
        total = percent + sum(
            (
                Decimal(o.percent)
                for o in others
                if o.effective_from <= day and (o.effective_to is None or o.effective_to >= day)
            ),
            ZERO,
        )
        if total > 100:
            return day, total
    return None


async def create_allocation(
    session: AsyncSession, user: User, payload: AllocationCreate, *, request: Request | None
) -> AllocationRead:
    _require(user)
    account = await _account(session, user, payload.account_id)
    if not account.is_active:
        raise _unprocessable("This account is inactive. Mark it active to map people to it.")
    employee = await session.scalar(
        tenant_scoped(
            select(Employee).where(
                Employee.id == payload.employee_id, Employee.deleted_at.is_(None)
            ),
            Employee,
            user,
        )
    )
    if employee is None:
        raise _not_found("That person is not in this workspace.")
    if employee.date_of_joining and payload.effective_from < employee.date_of_joining:
        raise _unprocessable(
            f"{_name(employee)} joined on {employee.date_of_joining:%d %b %Y}; the mapping "
            "cannot start before that."
        )
    if employee.date_of_exit and payload.effective_from > employee.date_of_exit:
        raise _unprocessable(
            f"{_name(employee)} left on {employee.date_of_exit:%d %b %Y}; the mapping cannot "
            "start after that."
        )
    end = payload.effective_to
    if employee.date_of_exit and (end is None or end > employee.date_of_exit):
        end = employee.date_of_exit
    clash = await _overbooked(
        session, user, employee.id, payload.effective_from, end, payload.percent
    )
    if clash is not None:
        day, total = clash
        raise _unprocessable(
            f"From {day:%d %b %Y} {_name(employee)} would be mapped {_percent(total)}% in all. "
            "A person's shares add up to at most 100%; end or lower another mapping first."
        )
    row = AccountAllocation(
        tenant_id=user.tenant_id,
        employee_id=employee.id,
        account_id=account.id,
        percent=payload.percent,
        effective_from=payload.effective_from,
        effective_to=end,
    )
    session.add(row)
    await session.flush()
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.CREATE,
        entity_type="account_allocation",
        entity_id=row.id,
        actor=user,
        changes={
            "employee": employee.employee_code,
            "account": account.code,
            "percent": str(payload.percent),
            "effective_from": payload.effective_from.isoformat(),
            "effective_to": end.isoformat() if end else None,
        },
        summary=f"{_name(employee)} mapped {_percent(payload.percent)}% to {account.code}",
        request=request,
    )
    await session.commit()
    found = await list_allocations(session, user, employee_id=employee.id)
    return next(a for a in found if a.id == row.id)


async def end_allocation(
    session: AsyncSession,
    user: User,
    allocation_id: uuid.UUID,
    payload: AllocationEnd,
    *,
    request: Request | None,
) -> AllocationRead:
    """Close a mapping on its last day. A new share is a new mapping from the next day."""
    _require(user)
    row = await _allocation(session, user, allocation_id)
    if payload.last_day < row.effective_from:
        raise _unprocessable(
            f"The mapping starts on {row.effective_from:%d %b %Y}; it cannot end before that. "
            "Delete it if it was a mistake."
        )
    if row.effective_to is not None and payload.last_day > row.effective_to:
        raise _unprocessable(
            f"The mapping already ends on {row.effective_to:%d %b %Y}. To extend it, add a new "
            "mapping from the next day."
        )
    before = row.effective_to
    row.effective_to = payload.last_day
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.UPDATE,
        entity_type="account_allocation",
        entity_id=row.id,
        actor=user,
        changes={
            "effective_to": {
                "before": before.isoformat() if before else None,
                "after": payload.last_day.isoformat(),
            }
        },
        summary=f"Account mapping ended on {payload.last_day:%d %b %Y}",
        request=request,
    )
    await session.commit()
    found = await list_allocations(session, user, employee_id=row.employee_id)
    return next(a for a in found if a.id == row.id)


async def delete_allocation(
    session: AsyncSession, user: User, allocation_id: uuid.UUID, *, request: Request | None
) -> None:
    """Remove a mapping entered by mistake.

    Past P&Ls change with it, so the audit row keeps what it was.
    """
    _require(user)
    row = await _allocation(session, user, allocation_id)
    await audit.record(
        session,
        tenant_id=user.tenant_id,
        action=AuditAction.DELETE,
        entity_type="account_allocation",
        entity_id=row.id,
        actor=user,
        changes={
            "employee_id": str(row.employee_id),
            "account_id": str(row.account_id),
            "percent": str(row.percent),
            "effective_from": row.effective_from.isoformat(),
            "effective_to": row.effective_to.isoformat() if row.effective_to else None,
        },
        summary="Account mapping deleted",
        request=request,
    )
    await session.delete(row)
    await session.commit()


# ── P&L ─────────────────────────────────────────────────────────────────────


async def _payroll_costs(
    session: AsyncSession, user: User, months: list[tuple[int, int]]
) -> tuple[dict[tuple[uuid.UUID, int, int], Decimal], set[str]]:
    """Each person's cost to company per month from payroll, and the months not final yet."""
    period = or_(
        *[and_(PayrollRun.period_year == y, PayrollRun.period_month == m) for y, m in months]
    )
    costs: dict[tuple[uuid.UUID, int, int], Decimal] = defaultdict(lambda: ZERO)
    pending: set[str] = set()
    rows = await session.execute(
        tenant_scoped(
            select(
                Payslip.employee_id,
                PayrollRun.period_year,
                PayrollRun.period_month,
                PayrollRun.status,
                func.coalesce(func.sum(Payslip.employer_cost), 0),
            )
            .join(PayrollRun, PayrollRun.id == Payslip.run_id)
            .where(period, PayrollRun.status.notin_(_NOT_A_COST))
            .group_by(
                Payslip.employee_id,
                PayrollRun.period_year,
                PayrollRun.period_month,
                PayrollRun.status,
            ),
            Payslip,
            user,
        )
    )
    for employee_id, year, month, run_status, amount in rows:
        costs[(employee_id, year, month)] += Decimal(amount)
        if run_status not in _FINAL:
            pending.add(math.month_key(year, month))
    return costs, pending


async def _allocations_in(
    session: AsyncSession, user: User, first: date, last: date, account_ids: list[uuid.UUID] | None
) -> list[AccountAllocation]:
    stmt = select(AccountAllocation).where(
        AccountAllocation.effective_from <= last,
        or_(AccountAllocation.effective_to.is_(None), AccountAllocation.effective_to >= first),
    )
    if account_ids is not None:
        stmt = stmt.where(AccountAllocation.account_id.in_(account_ids))
    return list((await session.scalars(tenant_scoped(stmt, AccountAllocation, user))).all())


async def _entries_in(
    session: AsyncSession, user: User, months: list[tuple[int, int]], account_ids: list[uuid.UUID]
) -> dict[tuple[uuid.UUID, str, int, int], Decimal]:
    period = or_(
        *[and_(AccountEntry.period_year == y, AccountEntry.period_month == m) for y, m in months]
    )
    out: dict[tuple[uuid.UUID, str, int, int], Decimal] = defaultdict(lambda: ZERO)
    for account_id, kind, year, month, amount in await session.execute(
        tenant_scoped(
            select(
                AccountEntry.account_id,
                AccountEntry.kind,
                AccountEntry.period_year,
                AccountEntry.period_month,
                func.coalesce(func.sum(AccountEntry.amount), 0),
            )
            .where(period, AccountEntry.account_id.in_(account_ids))
            .group_by(
                AccountEntry.account_id,
                AccountEntry.kind,
                AccountEntry.period_year,
                AccountEntry.period_month,
            ),
            AccountEntry,
            user,
        )
    ):
        out[(account_id, kind, year, month)] += Decimal(amount)
    return out


def _row(
    account: ClientAccount, billed: Decimal, people: Decimal, other: Decimal, heads: int
) -> AccountPnlRow:
    billed, people, other = math.rupees(billed), math.rupees(people), math.rupees(other)
    total = people + other
    margin = billed - total
    return AccountPnlRow(
        account_id=account.id,
        code=account.code,
        name=account.name,
        billed=billed,
        people_cost=people,
        other_cost=other,
        total_cost=total,
        margin=margin,
        margin_pct=math.margin_pct(billed, margin),
        people=heads,
    )


class _Computed:
    def __init__(self) -> None:
        self.billed: dict[tuple[uuid.UUID, str], Decimal] = defaultdict(lambda: ZERO)
        self.other: dict[tuple[uuid.UUID, str], Decimal] = defaultdict(lambda: ZERO)
        self.people: dict[tuple[uuid.UUID, str], Decimal] = defaultdict(lambda: ZERO)
        #: account → employee → (cost, share-months)
        self.by_person: dict[uuid.UUID, dict[uuid.UUID, list[Decimal]]] = defaultdict(dict)
        self.payroll_cost = ZERO
        self.mapped_cost = ZERO
        self.pending: set[str] = set()


async def _compute(
    session: AsyncSession,
    user: User,
    accounts: list[ClientAccount],
    months: list[tuple[int, int]],
) -> _Computed:
    ids = [a.id for a in accounts]
    first = math.month_bounds(*months[0])[0]
    last = math.month_bounds(*months[-1])[1]
    out = _Computed()
    costs, out.pending = await _payroll_costs(session, user, months)
    # Every mapping in the period, not only these accounts', so the part of a
    # person's cost that is on no account can be told apart.
    allocations = await _allocations_in(session, user, first, last, None)
    by_employee: dict[uuid.UUID, list[AccountAllocation]] = defaultdict(list)
    for allocation in allocations:
        by_employee[allocation.employee_id].append(allocation)
    wanted = set(ids)
    for (employee_id, year, month), cost in costs.items():
        key = math.month_key(year, month)
        out.payroll_cost += cost
        mine = by_employee.get(employee_id, [])
        shares: dict[uuid.UUID, Decimal] = defaultdict(lambda: ZERO)
        for allocation in mine:
            shares[allocation.account_id] += math.share_of_month(
                [(Decimal(allocation.percent), allocation.effective_from, allocation.effective_to)],
                year,
                month,
            )
        total_share = min(sum(shares.values(), ZERO), Decimal(1))
        out.mapped_cost += cost * total_share
        for account_id, share in shares.items():
            if account_id not in wanted or share == 0:
                continue
            amount = cost * share
            out.people[(account_id, key)] += amount
            entry = out.by_person[account_id].setdefault(employee_id, [ZERO, ZERO])
            entry[0] += amount
            entry[1] += share
    for (account_id, kind, year, month), amount in (
        await _entries_in(session, user, months, ids)
    ).items():
        target = out.billed if kind == AccountEntryKind.BILLING else out.other
        target[(account_id, math.month_key(year, month))] += amount
    return out


def _notes(out: _Computed, months: list[tuple[int, int]]) -> list[str]:
    notes = [
        "People cost is what payroll paid (cost to company on each payslip, all runs of the "
        "month that were not cancelled), shared out by each person's account mapping for the "
        "days of the month it covers."
    ]
    if out.pending:
        notes.append(
            "Payroll for "
            + ", ".join(sorted(out.pending))
            + " is not approved yet, so those months' cost may still change."
        )
    if hr_scope.current_scope() is not None:
        notes.append("You see only the people in your scope; cost of anyone else is left out.")
    keys = {math.month_key(y, m) for y, m in months}
    if not any(k[1] in keys for k in out.billed):
        notes.append(
            "Nothing has been billed in this period yet. Add billing under Billing & costs."
        )
    return notes


async def _period(session: AsyncSession, user: User, start: str | None, end: str | None):
    today = await local_today(session, user)
    if not end:
        end = math.month_key(today.year, today.month)
    if not start:
        # The financial year so far (April to the end month).
        year, month = _month(end)
        start = math.month_key(year if month >= 4 else year - 1, 4)
    try:
        months = math.months_between(_month(start), _month(end))
    except ValueError as exc:
        raise _unprocessable(str(exc)) from exc
    return start, end, months


async def pnl(
    session: AsyncSession,
    user: User,
    *,
    start: str | None = None,
    end: str | None = None,
    include_inactive: bool = False,
) -> AccountPnlRead:
    """Every account's P&L for the period, and the payroll cost on no account."""
    _require(user)
    return await account_pnl(session, user, start=start, end=end, include_inactive=include_inactive)


async def account_pnl(
    session: AsyncSession,
    user: User,
    *,
    start: str | None = None,
    end: str | None = None,
    include_inactive: bool = False,
    accounts: list[ClientAccount] | None = None,
) -> AccountPnlRead:
    """The P&L without the role check: callers check who may see it (Ask Mira does its own)."""
    start, end, months = await _period(session, user, start, end)
    if accounts is None:
        stmt = select(ClientAccount).order_by(ClientAccount.name)
        if not include_inactive:
            stmt = stmt.where(ClientAccount.is_active.is_(True))
        accounts = list((await session.scalars(tenant_scoped(stmt, ClientAccount, user))).all())
    keys = [math.month_key(y, m) for y, m in months]
    if not accounts:
        return AccountPnlRead(
            start=start,
            end=end,
            months=keys,
            accounts=[],
            totals=None,
            payroll_cost=ZERO,
            unmapped_cost=ZERO,
            notes=["No client accounts are set up yet. Add one under Accounts."],
        )
    out = await _compute(session, user, accounts, months)
    rows = [
        _row(
            account,
            sum((out.billed[(account.id, k)] for k in keys), ZERO),
            sum((out.people[(account.id, k)] for k in keys), ZERO),
            sum((out.other[(account.id, k)] for k in keys), ZERO),
            len(out.by_person.get(account.id, {})),
        )
        for account in accounts
    ]
    rows.sort(key=lambda r: (-r.billed, r.name))
    totals = AccountPnlRow(
        account_id=uuid.UUID(int=0),
        code="TOTAL",
        name="All accounts",
        billed=sum((r.billed for r in rows), ZERO),
        people_cost=sum((r.people_cost for r in rows), ZERO),
        other_cost=sum((r.other_cost for r in rows), ZERO),
        total_cost=sum((r.total_cost for r in rows), ZERO),
        margin=sum((r.margin for r in rows), ZERO),
        margin_pct=None,
        people=len({p for persons in out.by_person.values() for p in persons}),
    )
    totals.margin_pct = math.margin_pct(totals.billed, totals.margin)
    return AccountPnlRead(
        start=start,
        end=end,
        months=keys,
        accounts=rows,
        totals=totals,
        payroll_cost=math.rupees(out.payroll_cost),
        unmapped_cost=math.rupees(out.payroll_cost - out.mapped_cost),
        notes=_notes(out, months),
    )


async def pnl_detail(
    session: AsyncSession,
    user: User,
    account_id: uuid.UUID,
    *,
    start: str | None = None,
    end: str | None = None,
    checked: bool = False,
) -> PnlDetailRead:
    """One account month by month, and the people whose cost is on it."""
    if not checked:
        _require(user)
    account = await _account(session, user, account_id)
    start, end, months = await _period(session, user, start, end)
    out = await _compute(session, user, [account], months)
    keys = [math.month_key(y, m) for y, m in months]
    monthly = []
    for key in keys:
        row = _row(
            account,
            out.billed[(account.id, key)],
            out.people[(account.id, key)],
            out.other[(account.id, key)],
            0,
        )
        monthly.append(
            AccountPnlMonth(
                month=key,
                billed=row.billed,
                people_cost=row.people_cost,
                other_cost=row.other_cost,
                total_cost=row.total_cost,
                margin=row.margin,
                margin_pct=row.margin_pct,
            )
        )
    persons = out.by_person.get(account.id, {})
    employees: dict[uuid.UUID, Employee] = {}
    if persons:
        employees = {
            e.id: e
            for e in await session.scalars(
                tenant_scoped(
                    select(Employee).where(Employee.id.in_(list(persons))), Employee, user
                )
            )
        }
    people = [
        AccountPnlPerson(
            employee_id=employee_id,
            employee_code=employees[employee_id].employee_code if employee_id in employees else "",
            name=_name(employees[employee_id]) if employee_id in employees else "—",
            share_pct=(share_months / Decimal(len(months)) * 100).quantize(Decimal("0.1")),
            cost=math.rupees(cost),
        )
        for employee_id, (cost, share_months) in persons.items()
    ]
    people.sort(key=lambda p: -p.cost)
    summary = _row(
        account,
        sum((m.billed for m in monthly), ZERO),
        sum((m.people_cost for m in monthly), ZERO),
        sum((m.other_cost for m in monthly), ZERO),
        len(people),
    )
    read = (await _reads(session, user, [account], await local_today(session, user)))[0]
    return PnlDetailRead(
        account=read,
        start=start,
        end=end,
        summary=summary,
        months=monthly,
        people=people,
        notes=_notes(out, months),
    )
