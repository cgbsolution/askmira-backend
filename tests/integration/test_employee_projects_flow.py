"""Employees on projects: a client's, or the company's own (Internal).

CGB employs everybody through one legal entity; about half work on clients'
projects (DXC INET, JK Paper, UIDAI…), the rest are Internal. The project is
set per person, moved in bulk, made from the departments they used to stand in
for, or read from the old HRMS's salary register (Emp Code + PROJECT). Payroll
keeps the project on each payslip, and the salary register shows it.
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import httpx
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.audit import AuditLog
from app.models.employee import Employee, EmploymentStatus
from app.models.lifecycle import LifecycleEvent, LifecycleEventType
from app.models.org import Department, LegalEntity
from app.models.payroll import Payslip
from app.models.tenant import Tenant
from app.models.timesheets import Project
from app.models.user import UserRole
from app.services.payroll import run as run_service
from tests.integration.conftest import requires_db
from tests.integration.platform_support import bearer, client, make_person  # noqa: F401
from tests.integration.test_payroll_run_flow import (  # noqa: F401  (fixture)
    _approved_rules,
    _lock,
    org,
)

pytestmark = requires_db

API = "/api/v1"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _project(session: AsyncSession, tenant: Tenant, code: str, name: str, **extra):
    row = Project(tenant_id=tenant.id, code=code, name=name, **extra)
    session.add(row)
    await session.commit()
    await session.refresh(row)
    return row


async def _employee(
    session: AsyncSession, tenant: Tenant, code: str, *, department=None, project=None, **extra
):
    row = Employee(
        tenant_id=tenant.id,
        employee_code=code,
        first_name=code.title(),
        work_email=f"{code.lower()}-{uuid.uuid4().hex[:6]}@cgb.test",
        status=EmploymentStatus.ACTIVE,
        date_of_joining=date(2024, 1, 1),
        department_id=department.id if department else None,
        project_id=project.id if project else None,
        **extra,
    )
    session.add(row)
    await session.commit()
    await session.refresh(row)
    return row


def _register_sheet(rows: list[tuple[str, str | None]]) -> bytes:
    """Like the old HRMS's salary register: title rows, then S#, Emp Code, Name, PROJECT."""
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "CGBiz Solutions Private Limited"
    sheet["A2"] = "SALARY DETAIL FOR THE MONTH OF SEP-26"
    for col, head in enumerate(["S#", "Emp Code", "Name", "Designation", "PROJECT"], start=1):
        sheet.cell(row=6, column=col, value=head)
    for i, (code, project) in enumerate(rows, start=1):
        sheet.cell(row=6 + i, column=1, value=i)
        sheet.cell(row=6 + i, column=2, value=code)
        sheet.cell(row=6 + i, column=3, value=f"Person {code}")
        sheet.cell(row=6 + i, column=4, value="Consultant")
        sheet.cell(row=6 + i, column=5, value=project)
    sheet.cell(row=7 + len(rows), column=1, value="Grand Total:-")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ── The list of projects ────────────────────────────────────────────────────


async def test_a_project_names_its_company_and_counts_its_people(
    session: AsyncSession,
    client: httpx.AsyncClient,  # noqa: F811
    tenant: Tenant,
    other_tenant: Tenant,
    legal_entity: LegalEntity,
) -> None:
    hr = await make_person(session, tenant, role=UserRole.ADMIN, name="Hari HR")
    staff = await make_person(session, tenant, role=UserRole.EMPLOYEE, name="Sunil Staff")
    foreign = LegalEntity(tenant_id=other_tenant.id, name="Globex", code="GLOBEX")
    session.add(foreign)
    await session.commit()

    body = {
        "code": "DXC-INET",
        "name": "DXC INET",
        "client": "DXC Technology",
        "billable": True,
        "legal_entity_id": str(legal_entity.id),
    }
    created = await client.post(f"{API}/timesheets/projects", json=body, headers=bearer(hr))
    assert created.status_code == 201, created.text
    project = created.json()
    assert project["legal_entity_id"] == str(legal_entity.id)

    # Another workspace's company cannot be named.
    bad = {**body, "code": "X1", "legal_entity_id": str(foreign.id)}
    assert (
        await client.post(f"{API}/timesheets/projects", json=bad, headers=bearer(hr))
    ).status_code == 422

    await _employee(
        session, tenant, "DXC001", project=await session.get(Project, uuid.UUID(project["id"]))
    )
    leaver = await _employee(
        session, tenant, "DXC002", project=await session.get(Project, uuid.UUID(project["id"]))
    )
    leaver.status = EmploymentStatus.EXITED
    await session.commit()

    listed = (
        await client.get(f"{API}/timesheets/projects?include_inactive=true", headers=bearer(hr))
    ).json()
    assert [(p["code"], p["headcount"]) for p in listed] == [("DXC-INET", 1)]
    # Everybody can read the list (timesheets need it); only HR sees the counts.
    seen = (await client.get(f"{API}/timesheets/projects", headers=bearer(staff))).json()
    assert seen[0]["headcount"] is None

    # The company can be cleared, and other fields left alone.
    cleared = await client.patch(
        f"{API}/timesheets/projects/{project['id']}",
        json={"legal_entity_id": None},
        headers=bearer(hr),
    )
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()["legal_entity_id"] is None
    assert cleared.json()["client"] == "DXC Technology"


# ── Moving people onto projects ─────────────────────────────────────────────


async def test_hr_puts_people_on_a_project_in_one_go(
    session: AsyncSession,
    client: httpx.AsyncClient,  # noqa: F811
    tenant: Tenant,
    legal_entity: LegalEntity,
) -> None:
    internal = await _project(session, tenant, "INTERNAL", "Internal", billable=False)
    dxc = await _project(session, tenant, "DXC-INET", "DXC INET", client="DXC Technology")
    closed = await _project(session, tenant, "OLD", "Old client", active=False)
    hr = await make_person(session, tenant, role=UserRole.SUB_ADMIN, name="Hema HR")
    staff = await make_person(session, tenant, role=UserRole.EMPLOYEE, name="Sunil Staff")
    one = await _employee(session, tenant, "P001", project=internal)
    two = await _employee(session, tenant, "P002")
    there = await _employee(session, tenant, "P003", project=dxc)
    gone = await _employee(session, tenant, "P004", project=internal)
    gone.status = EmploymentStatus.EXITED
    await session.commit()

    url = f"{API}/employees/bulk/project"
    body = {
        "employee_ids": [
            str(one.id),
            str(two.id),
            str(there.id),
            str(gone.id),
            str(hr.employee_id),
            str(uuid.uuid4()),
        ],
        "project_id": str(dxc.id),
    }
    assert (await client.patch(url, json=body, headers=bearer(staff))).status_code == 403
    response = await client.patch(url, json=body, headers=bearer(hr))
    assert response.status_code == 200, response.text
    result = response.json()
    assert (result["moved"], result["already_there"], len(result["skipped"])) == (2, 1, 3)
    assert result["project_name"] == "DXC INET"

    await session.refresh(one)
    await session.refresh(gone)
    assert one.project_id == dxc.id
    assert gone.project_id == internal.id  # a leaver keeps the project they left on
    # Their company does not change: payroll follows the company, not the project.
    assert one.legal_entity_id is None
    event = await session.scalar(select(LifecycleEvent).where(LifecycleEvent.employee_id == one.id))
    assert event.event_type is LifecycleEventType.PROJECT_CHANGE
    assert event.previous_values == {"project_id": str(internal.id)}
    assert event.new_values == {"project_id": str(dxc.id)}
    audited = await session.scalar(
        select(func.count())
        .select_from(AuditLog)
        .where(AuditLog.entity_type == "employee", AuditLog.entity_id.in_([one.id, two.id]))
    )
    assert audited == 2

    # The directory filters by project and says each person's project.
    listed = (
        await client.get(f"{API}/employees?project_id={dxc.id}", headers=bearer(staff))
    ).json()
    assert {row["employee_code"] for row in listed["items"]} == {"P001", "P002", "P003"}
    assert all(row["project_id"] == str(dxc.id) for row in listed["items"])
    loose = (await client.get(f"{API}/employees?without_project=true", headers=bearer(hr))).json()
    assert str(two.id) not in {row["id"] for row in loose["items"]}
    csv = await client.get(f"{API}/employees/export?project_id={dxc.id}", headers=bearer(hr))
    assert "Project" in csv.text.splitlines()[0]
    assert "DXC INET" in csv.text

    # A closed project takes nobody new; null takes people off every project.
    body = {"employee_ids": [str(two.id)], "project_id": str(closed.id)}
    response = await client.patch(url, json=body, headers=bearer(hr))
    assert response.status_code == 422
    assert "closed" in response.json()["detail"]
    body = {"employee_ids": [str(two.id)], "project_id": None}
    response = await client.patch(url, json=body, headers=bearer(hr))
    assert response.json()["moved"] == 1
    await session.refresh(two)
    assert two.project_id is None


async def test_a_project_on_the_record_is_checked_and_kept(
    session: AsyncSession,
    client: httpx.AsyncClient,  # noqa: F811
    tenant: Tenant,
    other_tenant: Tenant,
) -> None:
    hr = await make_person(session, tenant, role=UserRole.ADMIN, name="Hari HR")
    dxc = await _project(session, tenant, "DXC-INET", "DXC INET")
    foreign = await _project(session, other_tenant, "GLB", "Globex work")
    person = await _employee(session, tenant, "R001", project=dxc)

    url = f"{API}/employees/{person.id}"
    response = await client.patch(url, json={"project_id": str(foreign.id)}, headers=bearer(hr))
    assert response.status_code == 422
    # The project closes; saving their phone (the form sends the project too) still works.
    dxc.active = False
    await session.commit()
    response = await client.patch(
        url, json={"project_id": str(dxc.id), "phone": "9800000000"}, headers=bearer(hr)
    )
    assert response.status_code == 200, response.text
    # A change made on the record is a project change in their history.
    internal = await _project(session, tenant, "INTERNAL", "Internal", billable=False)
    response = await client.patch(url, json={"project_id": str(internal.id)}, headers=bearer(hr))
    assert response.status_code == 200, response.text
    event = await session.scalar(
        select(LifecycleEvent).where(
            LifecycleEvent.employee_id == person.id,
            LifecycleEvent.event_type == LifecycleEventType.PROJECT_CHANGE,
        )
    )
    assert event.new_values == {"project_id": str(internal.id)}


# ── From departments, and from the old register ─────────────────────────────


async def test_departments_that_stood_in_for_projects_become_projects(
    session: AsyncSession,
    client: httpx.AsyncClient,  # noqa: F811
    tenant: Tenant,
    legal_entity: LegalEntity,
) -> None:
    admin = await make_person(session, tenant, role=UserRole.ADMIN, name="Deepak Admin")
    sub = await make_person(session, tenant, role=UserRole.SUB_ADMIN, name="Sana Sub")
    internal = Department(tenant_id=tenant.id, name="Internal", code="INTERNAL")
    external = Department(tenant_id=tenant.id, name="External", code="EXTERNAL")
    engineering = Department(tenant_id=tenant.id, name="Engineering", code="ENG")
    session.add_all([internal, external, engineering])
    await session.commit()
    a = await _employee(session, tenant, "I001", department=internal)
    b = await _employee(session, tenant, "E001", department=external)
    c = await _employee(session, tenant, "G001", department=engineering)

    url = f"{API}/org/projects/from-departments"
    body = {"department_ids": [str(internal.id), str(external.id)], "clear_department": True}
    assert (await client.post(url, json=body, headers=bearer(sub))).status_code == 403
    response = await client.post(url, json=body, headers=bearer(admin))
    assert response.status_code == 200, response.text
    result = response.json()
    assert sorted(result["projects_created"]) == ["External", "Internal"]
    assert result["moved"] == 2

    projects = {p.code: p for p in (await session.scalars(select(Project))).all()}
    assert projects["INTERNAL"].billable is False  # the company's own staff
    assert projects["EXTERNAL"].billable is True
    assert projects["INTERNAL"].legal_entity_id == legal_entity.id  # the only company
    for person, code in ((a, "INTERNAL"), (b, "EXTERNAL")):
        await session.refresh(person)
        assert person.project_id == projects[code].id
        assert person.department_id is None
    await session.refresh(c)
    assert c.project_id is None and c.department_id == engineering.id

    # Again: nothing new, nothing moved.
    again = (await client.post(url, json=body, headers=bearer(admin))).json()
    assert again["projects_created"] == [] and again["moved"] == 0
    assert sorted(again["projects_reused"]) == ["External", "Internal"]


async def test_the_old_salary_register_puts_everybody_on_their_project(
    session: AsyncSession,
    client: httpx.AsyncClient,  # noqa: F811
    tenant: Tenant,
    legal_entity: LegalEntity,
) -> None:
    admin = await make_person(session, tenant, role=UserRole.ADMIN, name="Deepak Admin")
    internal = await _project(session, tenant, "INTERNAL", "Internal", billable=False)
    first = await _employee(session, tenant, "CDL23-232")
    second = await _employee(session, tenant, "DIC-CNST1")
    third = await _employee(session, tenant, "INT-CNST54", project=internal)
    await _employee(session, tenant, "AB0031")
    sheet = _register_sheet(
        [
            ("CDL23-232", "JK Paper Limited"),
            ("DIC-CNST1", "Data  Intensity"),  # spaces as typed
            ("INT-CNST54", "Internal"),
            ("AB0031", None),
            ("ZZZ-999", "UIDAI"),
            ("CDL23-232", "Powergrid"),
        ]
    )
    files = {"file": ("Salary Register - Sep-26.xlsx", sheet, XLSX)}
    url = f"{API}/org/projects/assign-from-sheet"

    preview = await client.post(url, files=files, data={"apply": "false"}, headers=bearer(admin))
    assert preview.status_code == 200, preview.text
    checked = preview.json()
    outcomes = {(row["employee_code"], row["outcome"]) for row in checked["rows"]}
    assert ("CDL23-232", "assign") in outcomes
    assert ("CDL23-232", "duplicate") in outcomes
    assert ("INT-CNST54", "unchanged") in outcomes
    assert ("AB0031", "blank") in outcomes
    assert ("ZZZ-999", "unknown") in outcomes
    assert checked["new_projects"] == ["Data Intensity", "JK Paper Limited"]
    # The preview changed nothing.
    assert await session.scalar(select(func.count()).select_from(Project)) == 1

    applied = await client.post(url, files=files, data={"apply": "true"}, headers=bearer(admin))
    assert applied.status_code == 200, applied.text
    assert applied.json()["moved"] == 2
    projects = {p.name: p for p in (await session.scalars(select(Project))).all()}
    assert set(projects) == {"Internal", "JK Paper Limited", "Data Intensity"}
    assert projects["JK Paper Limited"].legal_entity_id == legal_entity.id
    await session.refresh(first)
    await session.refresh(second)
    await session.refresh(third)
    assert first.project_id == projects["JK Paper Limited"].id
    assert second.project_id == projects["Data Intensity"].id
    assert third.project_id == internal.id

    # The same file again changes nothing.
    again = (
        await client.post(url, files=files, data={"apply": "true"}, headers=bearer(admin))
    ).json()
    assert again["moved"] == 0 and again["projects_created"] == []

    # A sheet without a project column is refused, saying what it needs.
    book = Workbook()
    book.active.append(["Emp Code", "Name"])
    book.active.append(["CDL23-232", "Someone"])
    out = io.BytesIO()
    book.save(out)
    bad = {"file": ("people.xlsx", out.getvalue(), XLSX)}
    response = await client.post(url, files=bad, data={"apply": "false"}, headers=bearer(admin))
    assert response.status_code == 422
    assert "Project" in response.json()["detail"]


async def test_the_employee_import_reads_a_project_column(
    session: AsyncSession,
    client: httpx.AsyncClient,  # noqa: F811
    tenant: Tenant,
    legal_entity: LegalEntity,
) -> None:
    admin = await make_person(session, tenant, role=UserRole.ADMIN, name="Deepak Admin")
    internal = await _project(session, tenant, "INTERNAL", "Internal", billable=False)
    book = Workbook()
    sheet = book.active
    sheet.append(["Emp Code", "First Name", "Last Name", "Email", "Joining Date", "Project"])
    sheet.append(["NEW001", "Asha", "Verma", "asha.verma@example.com", "01-04-2025", "Internal"])
    sheet.append(["NEW002", "Ravi", "Kumar", "ravi.kumar@example.com", "01-04-2025", "UIDAI"])
    out = io.BytesIO()
    book.save(out)
    files = {"file": ("joiners.xlsx", out.getvalue(), XLSX)}

    preview = (
        await client.post(f"{API}/employees/import/preview", files=files, headers=bearer(admin))
    ).json()
    assert preview["new_projects"] == ["UIDAI"]
    imported = await client.post(f"{API}/employees/import", files=files, headers=bearer(admin))
    assert imported.status_code == 201, imported.text
    assert imported.json()["projects_created"] == ["UIDAI"]
    asha = await session.scalar(select(Employee).where(Employee.employee_code == "NEW001"))
    ravi = await session.scalar(select(Employee).where(Employee.employee_code == "NEW002"))
    uidai = await session.scalar(select(Project).where(Project.name == "UIDAI"))
    assert asha.project_id == internal.id
    assert ravi.project_id == uidai.id


# ── Payroll keeps the project on each payslip ───────────────────────────────


async def test_the_salary_register_says_which_project_each_person_was_on(
    session: AsyncSession,
    client: httpx.AsyncClient,  # noqa: F811
    org,  # noqa: F811
) -> None:
    await _approved_rules(session, org)
    session.add(_lock(org["tenant"], org["maker"], 2026, 4))
    await session.commit()
    dxc = await _project(session, org["tenant"], "DXC-INET", "DXC INET")
    internal = await _project(session, org["tenant"], "INTERNAL", "Internal", billable=False)
    meera = await session.get(Employee, org["meera_emp"].id)
    meera.project_id = dxc.id
    await session.commit()

    april = await run_service.create_run(
        session, org["maker"], legal_entity_id=org["entity"].id, period_year=2026, period_month=4
    )
    await run_service.compute_run(session, org["maker"], april.id)
    slip = await session.scalar(
        select(Payslip).where(Payslip.run_id == april.id, Payslip.employee_id == meera.id)
    )
    assert slip.project_id == dxc.id

    # She moves to Internal afterwards: April's register still says DXC INET.
    meera.project_id = internal.id
    await session.commit()
    register = await client.get(
        f"{API}/payroll/runs/{april.id}/register", headers=bearer(org["maker"])
    )
    assert register.status_code == 200, register.text
    rows = {row["employee_code"]: row for row in register.json()["rows"]}
    assert rows["MEERA"]["project_name"] == "DXC INET"

    download = await client.get(
        f"{API}/payroll/runs/{april.id}/register.xlsx", headers=bearer(org["maker"])
    )
    assert download.status_code == 200, download.text
    sheet = load_workbook(io.BytesIO(download.content)).active
    headings = [cell.value for cell in sheet[7]]
    column = headings.index("Project")
    codes = [cell.value for cell in sheet["B"]]
    assert sheet.cell(row=codes.index("MEERA") + 1, column=column + 1).value == "DXC INET"
