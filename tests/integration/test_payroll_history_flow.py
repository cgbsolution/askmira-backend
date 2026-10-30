"""factoHR payroll history import, end to end against Postgres.

Three synthetic registers (Jan-24 header row 1, Feb-24 with an F&F sheet, Mar-24
with ROUNDED) for nine people chosen so that every rule runs once:

  CHD23-101 Asha    live salary = the register rate, wrongly dated → re-dated
  CHD23-102 Bala    increment in Feb; live = new rate → old rate inserted, live re-dated
  CHD23-103 Chitra  active, live differs from every register rate → later increment, warned
  CHD23-104 Dev     left in Feb, live differs → register rates inserted, live deleted
  CHD23-105 Esha    joined Jan-10, no salary in AskMira → inserted from her joining date
  CHD23-106 Farah   live starts after her last register while still employed → kept
  CHD23-107 Gopal   mid-month change in Feb, new rate's row listed first → old rate first
  CHD23-108 Hari    Feb: main "Hold" zero-pay row + F&F row → one payslip, adjustment 439
  CDL24-565 Kushal  not an AskMira code; AB0024 has his name and joining date → mapped

Gates: one locked, published, imported run per month whose totals are its
payslips'; every payslip adds up as the PDF checks it; TDS is never a line;
register differences are lines of their own; salary series have no gaps or
overlaps and each payslip links to the revision in force; months already run
are refused; nothing is written while any error remains; undo removes only
imported runs; audit rows share the transaction.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncGenerator
from datetime import date
from decimal import Decimal

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import create_access_token
from app.db.session import get_session
from app.main import app
from app.models.audit import AuditAction, AuditLog
from app.models.employee import Employee, EmploymentStatus
from app.models.org import LegalEntity
from app.models.payroll import (
    EmployeeSalary,
    PayrollRun,
    PayrollRunStatus,
    PayrollRunType,
    Payslip,
    SalaryStructure,
)
from app.models.tenant import Tenant
from app.models.timesheets import Project
from app.models.user import User, UserRole
from app.schemas.payroll import EmployeeSalaryCreate, SalaryComponentAmountInput
from app.services.payroll import history_import, payslip_pdf
from app.services.payroll import service as salary_service
from app.services.payroll.history_parse import parse_registers
from app.services.payroll.run import other_earnings_of
from app.services.payroll.salary_master import revision_locks
from app.services.payroll.standard_setup import install_standard
from tests.integration.conftest import requires_db
from tests.payroll.payroll_history_books import DEC23, FNF, MAY24, Sheet, book, person

pytestmark = requires_db

D = Decimal
JAN, FEB, MAR = 31, 29, 31

RATES = {
    "asha": dict(ctc=30000, basic=12000, hra=4800, special=10000),
    "bala_old": dict(ctc=40000, basic=16000, hra=6400, special=14000),
    "bala_new": dict(ctc=44000, basic=17600, hra=7040, special=15400),
    "chitra": dict(ctc=50000, basic=20000, hra=8000, special=18000),
    "dev": dict(ctc=25000, basic=10000, hra=4000, special=9000),
    "esha": dict(ctc=20000, basic=8000, hra=3200, special=7000),
    "farah": dict(ctc=35000, basic=14000, hra=5600, special=12500),
    "gopal_old": dict(ctc=24000, basic=9600, hra=3840, special=8500),
    "gopal_new": dict(ctc=30000, basic=12000, hra=4800, special=10600),
    "hari": dict(ctc=60000, basic=24000, hra=9600, special=22000),
    "kushal": dict(ctc=18000, basic=7200, hra=2880, special=6500),
}
JOINED = {
    "CHD23-101": date(2022, 1, 1),
    "CHD23-102": date(2022, 1, 1),
    "CHD23-103": date(2022, 1, 1),
    "CHD23-104": date(2022, 1, 1),
    "CHD23-105": date(2024, 1, 10),
    "CHD23-106": date(2022, 1, 1),
    "CHD23-107": date(2022, 1, 1),
    "CHD23-108": date(2022, 1, 1),
    "CDL24-565": date(2024, 1, 5),
}


def _row(code: str, name: str, rate: str, **kwargs):
    return person(code, name, JOINED[code], **RATES[rate], **kwargs)


def _zero(code: str, name: str, rate: str, workable: int, **kwargs):
    row = _row(code, name, rate, workable=workable, worked=0, lop=workable, **kwargs)
    for key in ("basic", "hra", "special", "bonus", "gross", "pf", "esi", "ins", "pt", "tds",
                "total_ded", "net", "adj", "bank", "rounded", "er_pf", "er_esi", "gratuity",
                "ctc_earned", "a_basic", "a_hra", "a_special", "a_bonus", "a_gross"):
        row[key] = 0
    return row


def registers() -> list[tuple[str, bytes]]:
    jan_rows = [
        _row("CHD23-101", "Asha Rao", "asha", workable=JAN, worked=30, lop=1,
             client="CGB Internal"),
        _row("CHD23-102", "Bala Iyer", "bala_old", workable=JAN, client="BOB"),
        _row("CHD23-103", "Chitra N", "chitra", workable=JAN, tds=1250.5),
        _row("CHD23-104", "Dev K", "dev", workable=JAN, client="Transition"),
        _row("CHD23-105", "Esha M", "esha", workable=JAN, worked=22, lop=0),
        _row("CHD23-106", "Farah S", "farah", workable=JAN),
        _row("CHD23-107", "Gopal R", "gopal_old", workable=JAN),
        _row("CHD23-108", "Hari P", "hari", workable=JAN),
        _row("CDL24-565", "Kushal Sunaria", "kushal", workable=JAN, worked=27, lop=0),
    ]
    jan = book([Sheet("Sheet1", DEC23, jan_rows)])

    hold = _row("CHD23-108", "Hari P", "hari", workable=FEB, worked=0, lop=FEB, pf=0,
                status="Hold", remarks="23 feb 24 LEFT")
    for key in ("basic", "hra", "special", "bonus", "gross", "er_pf", "gratuity", "ctc_earned"):
        hold[key] = 0
    hold.update(net=-439, total_ded=439, bank=0)
    feb_rows = [
        _row("CHD23-101", "Asha Rao", "asha", workable=FEB, adj=-200, client="CGB Internal"),
        _row("CHD23-102", "Bala Iyer", "bala_new", workable=FEB, adj=500, client="BOB"),
        _row("CHD23-103", "Chitra N", "chitra", workable=FEB),
        _row("CHD23-104", "Dev K", "dev", workable=FEB, worked=10, lop=0, status="F&F Paid",
             client="Transition"),
        _row("CHD23-105", "Esha M", "esha", workable=FEB),
        _row("CHD23-106", "Farah S", "farah", workable=FEB),
        _row("CHD23-107", "Gopal R", "gopal_new", workable=FEB, worked=16, lop=0),
        _row("CHD23-107", "Gopal R", "gopal_old", workable=FEB, worked=13, lop=0),
        hold,
        _row("CDL24-565", "Kushal Sunaria", "kushal", workable=FEB),
    ]
    main = Sheet("Salary Register", DEC23, feb_rows, above=[None, None])
    settlement = _row("CHD23-108", "Hari P", "hari", workable=FEB, worked=23, pt=0,
                      dol=date(2024, 2, 23))
    fnf = Sheet("F&F Payment", FNF, [settlement], above=[None],
                formulas={"basic": "={a_basic_col}{r}/29*{worked_col}{r}"})
    feb = book([main, fnf])

    hold_mar = _row("CHD23-105", "Esha M", "esha", workable=MAR, whole=True, status="HOLD")
    hold_mar.update(net=0, bank=0, rounded=0)
    kushal = _row("CDL24-565", "Kushal Sunaria", "kushal", workable=MAR, whole=True)
    kushal["gross"] = kushal["gross"] + 10  # a typed gross, not the sum of its parts
    kushal["total_ded"] = kushal["total_ded"] + 61  # a typed total deduction
    kushal["net"] = kushal["gross"] - kushal["total_ded"]
    kushal["bank"] = kushal["rounded"] = kushal["net"]
    mar_rows = [
        _row("CHD23-101", "Asha Rao", "asha", workable=MAR, whole=True, client="CGB Internal"),
        _row("CHD23-102", "Bala Iyer", "bala_new", workable=MAR, whole=True, client="BOB"),
        _row("CHD23-103", "Chitra N", "chitra", workable=MAR, whole=True, tds=100.4),
        _zero("CHD23-104", "Dev K", "dev", MAR, status="Inactive"),
        hold_mar,
        _row("CHD23-107", "Gopal R", "gopal_new", workable=MAR, whole=True),
        kushal,
    ]
    mar = book([Sheet("Sheet1", MAY24, mar_rows, above=[["Salary Register Mar - 24"], None])])
    return [
        ("Salary Register Jan-24.xlsx", jan),
        ("Salary Register Feb-24.xlsx", feb),
        ("Salary Register Mar-24.xlsx", mar),
    ]


# ── Fixtures ────────────────────────────────────────────────────────────────


async def _employee(session, tenant, entity, code, first, last, *, status=EmploymentStatus.ACTIVE,
                    exited=None, joined=None) -> Employee:
    employee = Employee(
        tenant_id=tenant.id,
        employee_code=code,
        first_name=first,
        last_name=last,
        work_email=f"{code.lower()}@acme.test",
        legal_entity_id=entity.id,
        status=status,
        date_of_joining=joined or JOINED.get(code),
        date_of_exit=exited,
    )
    session.add(employee)
    await session.flush()
    return employee


async def _live(session, user, employee, structure, start: date, *, ctc, basic, hra, special,
                bonus=0) -> EmployeeSalary:
    """A current revision as the Employee Detail import left it: yearly figures."""
    payload = EmployeeSalaryCreate(
        entry="manual",
        ctc_annual=D(ctc * 12),
        gross_amount=D((basic + hra + special + bonus) * 12),
        gross_basis="annual",
        components=[
            SalaryComponentAmountInput(component_code=code, amount=D(amount * 12), basis="annual")
            for code, amount in (("BASIC", basic), ("HRA", hra), ("SPECIAL", special),
                                 ("BONUS", bonus))
            if amount
        ],
        structure_id=structure.id,
        effective_from=start,
        revision_reason="From factoHR Employee Detail; to be corrected from payroll history.",
    )
    return await salary_service.assign_salary(session, user, employee.id, payload, commit=False)


@pytest.fixture
async def world(session: AsyncSession, tenant: Tenant, hr_user: User, legal_entity: LegalEntity):
    await install_standard(session, hr_user)
    structure = await session.scalar(select(SalaryStructure))
    session.add_all(
        [
            Project(tenant_id=tenant.id, code="INTERNAL", name="Internal"),
            Project(tenant_id=tenant.id, code="BOB", name="BOB"),
        ]
    )
    people = {}
    people["asha"] = await _employee(session, tenant, legal_entity, "CHD23-101", "Asha", "Rao")
    people["bala"] = await _employee(session, tenant, legal_entity, "CHD23-102", "Bala", "Iyer")
    people["chitra"] = await _employee(session, tenant, legal_entity, "CHD23-103", "Chitra", "N")
    people["dev"] = await _employee(session, tenant, legal_entity, "CHD23-104", "Dev", "K",
                                    status=EmploymentStatus.EXITED, exited=date(2024, 2, 10))
    people["esha"] = await _employee(session, tenant, legal_entity, "CHD23-105", "Esha", "M")
    people["farah"] = await _employee(session, tenant, legal_entity, "CHD23-106", "Farah", "S",
                                      status=EmploymentStatus.EXITED, exited=date(2024, 6, 30))
    people["gopal"] = await _employee(session, tenant, legal_entity, "CHD23-107", "Gopal", "R")
    people["hari"] = await _employee(session, tenant, legal_entity, "CHD23-108", "Hari", "P",
                                     status=EmploymentStatus.EXITED, exited=date(2024, 2, 23))
    people["kushal"] = await _employee(session, tenant, legal_entity, "AB0024", "Kushal",
                                       "Sunaria", joined=date(2024, 1, 5))
    live = {
        "asha": await _live(session, hr_user, people["asha"], structure, date(2024, 9, 1),
                            **RATES["asha"]),
        "bala": await _live(session, hr_user, people["bala"], structure, date(2023, 3, 15),
                            **RATES["bala_new"]),
        "chitra": await _live(session, hr_user, people["chitra"], structure, date(2024, 2, 1),
                              ctc=55000, basic=22000, hra=8800, special=20000),
        "dev": await _live(session, hr_user, people["dev"], structure, date(2024, 1, 15),
                           ctc=26000, basic=10400, hra=4160, special=9400),
        "farah": await _live(session, hr_user, people["farah"], structure, date(2024, 5, 1),
                             ctc=38000, basic=15200, hra=6080, special=13500),
        "hari": await _live(session, hr_user, people["hari"], structure, date(2024, 3, 1),
                            **RATES["hari"]),
    }
    await session.commit()
    return {"people": people, "live": live, "structure": structure, "entity": legal_entity}


async def _plan(session, user, files=None):
    return await history_import.plan(session, user, parse_registers(files or registers()))


async def _series(session, employee_id) -> list[EmployeeSalary]:
    return list(
        (
            await session.scalars(
                select(EmployeeSalary)
                .where(EmployeeSalary.employee_id == employee_id)
                .order_by(EmployeeSalary.effective_from)
            )
        ).all()
    )


def _lines(payslip) -> dict[str, Decimal]:
    return {line["code"]: D(line["amount"]) for line in payslip.line_items}


def _inputs(payslip) -> dict[str, dict]:
    return {line["code"]: line for line in payslip.input_lines}


# ── The import ──────────────────────────────────────────────────────────────


async def test_the_full_import_writes_locked_months_payslips_and_dated_salaries(
    session: AsyncSession, hr_user: User, world
) -> None:
    people, live = world["people"], world["live"]
    plan = await _plan(session, hr_user)
    assert plan.blocking == 0, [i.message for i in plan.issues if i.level == "error"]
    report = history_import.preview(plan)
    assert [m["payslips"] for m in report["months"]] == [9, 9, 7]
    assert report["people"]["mapped_by_name"] == [
        {"register_code": "CDL24-565", "code": "AB0024", "name": "Kushal Sunaria"}
    ]
    await history_import.apply(session, hr_user, plan)

    runs = (
        await session.scalars(select(PayrollRun).order_by(PayrollRun.period_month))
    ).all()
    assert [(r.period_year, r.period_month) for r in runs] == [(2024, 1), (2024, 2), (2024, 3)]
    for run in runs:
        assert run.origin == "imported" and run.status is PayrollRunStatus.PAID
        assert run.run_type is PayrollRunType.REGULAR
        assert run.approved_by_user_id == run.published_by_user_id == run.paid_by_user_id
        assert run.approved_by_user_id == hr_user.id and run.computed_by_user_id is None
        assert run.processed_at is None and run.payment_date is None
        assert run.rule_set_snapshot == {} and run.variance_flags == []
        assert run.reason.startswith("Imported from factoHR: Salary Register")
        assert run.import_details["system"] == "factoHR"
        assert run.import_details["files"][0]["sha256"]
        slips = (await session.scalars(select(Payslip).where(Payslip.run_id == run.id))).all()
        assert run.employee_count == len(slips)
        assert D(str(run.total_gross)) == sum((D(str(p.gross_earnings)) for p in slips), D(0))
        assert D(str(run.total_net)) == sum((D(str(p.net_pay)) for p in slips), D(0))
        assert D(str(run.total_deductions)) == sum(
            (D(str(p.total_deductions)) for p in slips), D(0)
        )
        assert D(str(run.total_employer_cost)) == sum(
            (D(str(p.employer_cost)) for p in slips), D(0)
        )

    slips = (await session.scalars(select(Payslip))).all()
    assert len(slips) == 25
    for slip in slips:
        assert slip.is_published and slip.published_at is not None
        assert slip.tds_working is None and slip.review_flags == []
        assert all(line["code"] != "TDS" for line in slip.line_items)
        codes = [line["code"] for line in slip.line_items]
        assert len(codes) == len(set(codes))
        assert all(line["rule_set_id"] is None for line in slip.line_items)
        assert all(line["basis"].startswith("As paid in factoHR") for line in slip.line_items)
        # The arithmetic the payslip PDF refuses to render without.
        payslip_pdf.build_payslip_document(
            company_name="Acme", company_address=None, period_year=2024, period_month=1,
            employee_name="x", employee_code="x", designation=None, date_of_joining=None,
            pan=None, uan=None, payment_date=None, working_days=slip.working_days,
            paid_days=slip.paid_days, lop_days=slip.lop_days,
            not_employed_days=slip.not_employed_days, line_items=slip.line_items,
            input_lines=slip.input_lines, tds=slip.tds, gross_earnings=slip.gross_earnings,
            other_earnings=other_earnings_of(slip), total_deductions=slip.total_deductions,
            net_pay=slip.net_pay,
        )
        for value in ("BPNPP1086C", "870330987299", "132420613006", "HSBC0400002",
                      "101095435751", "someone@example.com"):
            assert value not in str(slip.import_details)

    by = {}
    for slip in slips:
        run = next(r for r in runs if r.id == slip.run_id)
        by[(slip.employee_id, run.period_month)] = slip

    # Hari, February: the "Hold" row and the F&F settlement are one payslip.
    hari_feb = by[(people["hari"].id, 2)]
    assert hari_feb.import_details["kind"] == "combined"
    assert hari_feb.import_details["sheets"] == ["Salary Register", "F&F Payment"]
    assert D(_inputs(hari_feb)["ADJUSTMENT"]["amount"]) == D("439.00")
    assert D(str(hari_feb.working_days)) == D("29") and D(str(hari_feb.paid_days)) == D("23")
    assert _lines(hari_feb)["INSURANCE"] == D("478.00")  # both rows' 239
    assert hari_feb.employee_salary_id == live["hari"].id  # in force on his exit day

    # Adjustments: positive ones are taxable salary, negative ones are not.
    bala_feb = _inputs(by[(people["bala"].id, 2)])["ADJUSTMENT"]
    assert D(bala_feb["amount"]) == D("500.00") and D(bala_feb["taxable_amount"]) == D("500.00")
    asha_feb = _inputs(by[(people["asha"].id, 2)])["ADJUSTMENT"]
    assert D(asha_feb["amount"]) == D("-200.00") and D(asha_feb["taxable_amount"]) == D("0.00")
    assert asha_feb["provider"] == "factohr_register" and asha_feb["kind"] == "earning"

    # March: ROUNDED, a hand-set net, a typed gross and total deduction, a zero row.
    chitra_mar = by[(people["chitra"].id, 3)]
    assert D(_inputs(chitra_mar)["ROUNDED_OFF"]["amount"]) == D("0.40")
    assert D(str(chitra_mar.tds)) == D("100.40")
    esha_mar = by[(people["esha"].id, 3)]
    assert D(str(esha_mar.net_pay)) == D("0.00")
    assert D(_inputs(esha_mar)["NET_DIFF"]["amount"]) == -(
        D(str(esha_mar.gross_earnings)) - D(str(esha_mar.total_deductions))
    )
    kushal_mar = by[(people["kushal"].id, 3)]
    assert _lines(kushal_mar)["GROSS_DIFF"] == D("10.00")
    assert _lines(kushal_mar)["DEDUCTION_DIFF"] == D("61.00")
    assert kushal_mar.import_details["register_code"] == "CDL24-565"
    dev_mar = by[(people["dev"].id, 3)]
    assert dev_mar.import_details["zero_row"] is True and dev_mar.line_items == []

    # January (unrounded cells): rounding to the paisa is a line of its own.
    asha_jan = by[(people["asha"].id, 1)]
    assert D(str(asha_jan.lop_days)) == D("1")
    assert abs(D(_inputs(asha_jan).get("PAISA_ROUNDING", {"amount": "0"})["amount"])) <= D(
        "0.05"
    )

    # Projects: "CGB Internal" is the Internal project; text no project matches stays text.
    internal = await session.scalar(select(Project.id).where(Project.code == "INTERNAL"))
    assert asha_jan.project_id == internal
    assert by[(people["dev"].id, 1)].project_id is None
    assert by[(people["dev"].id, 1)].import_details["client"] == "Transition"

    # ── Salary revisions ────────────────────────────────────────────────────
    asha = await _series(session, people["asha"].id)
    assert [(s.id, s.effective_from, s.effective_to) for s in asha] == [
        (live["asha"].id, date(2024, 1, 1), None)
    ]
    bala = await _series(session, people["bala"].id)
    assert [(s.effective_from, s.effective_to) for s in bala] == [
        (date(2024, 1, 1), date(2024, 1, 31)),
        (date(2024, 2, 1), None),
    ]
    assert bala[1].id == live["bala"].id
    assert D(str(bala[0].ctc_annual)) == D("480000.00")
    assert {c.component_code: D(str(c.monthly_amount)) for c in bala[0].components} == {
        "BASIC": D("16000.00"), "HRA": D("6400.00"), "SPECIAL": D("14000.00")
    }
    assert all(c.pay_basis == "monthly" for c in bala[0].components)
    assert bala[0].gross_basis == "monthly" and D(str(bala[0].gross_amount)) == D("36400.00")
    assert bala[0].revision_reason.startswith("factoHR payroll register Jan 2024")
    assert by[(people["bala"].id, 1)].employee_salary_id == bala[0].id
    assert by[(people["bala"].id, 2)].employee_salary_id == bala[1].id

    chitra = await _series(session, people["chitra"].id)
    assert [(s.effective_from, s.effective_to) for s in chitra] == [
        (date(2024, 1, 1), date(2024, 3, 31)),
        (date(2024, 4, 1), None),
    ]
    assert chitra[1].id == live["chitra"].id
    assert any(
        i.category == "increment_date_unknown" and i.code == "CHD23-103" for i in plan.issues
    )

    dev = await _series(session, people["dev"].id)
    assert len(dev) == 1 and dev[0].id != live["dev"].id
    assert (dev[0].effective_from, dev[0].effective_to) == (date(2024, 1, 1), None)
    assert by[(people["dev"].id, 2)].employee_salary_id == dev[0].id

    esha = await _series(session, people["esha"].id)
    assert [(s.effective_from, s.effective_to) for s in esha] == [(date(2024, 1, 10), None)]

    farah = await _series(session, people["farah"].id)
    assert [(s.effective_from, s.effective_to) for s in farah] == [
        (date(2024, 1, 1), date(2024, 4, 30)),
        (date(2024, 5, 1), None),
    ]
    assert farah[1].id == live["farah"].id

    gopal = await _series(session, people["gopal"].id)
    assert [(s.effective_from, s.effective_to, D(str(s.ctc_annual))) for s in gopal] == [
        (date(2024, 1, 1), date(2024, 2, 13), D("288000.00")),
        (date(2024, 2, 14), None, D("360000.00")),
    ]
    assert by[(people["gopal"].id, 2)].import_details["kind"] == "combined"

    hari = await _series(session, people["hari"].id)
    assert [(s.id, s.effective_from) for s in hari] == [(live["hari"].id, date(2024, 1, 1))]

    # Revisions payroll was paid on are locked in the Salary master now.
    locks = await revision_locks(session, hr_user, [live["asha"].id, bala[0].id])
    assert set(locks) == {live["asha"].id, bala[0].id}

    # ── Audit, in the same transaction ──────────────────────────────────────
    audits = (await session.scalars(select(AuditLog))).all()
    kinds = {(a.entity_type, a.action.value) for a in audits}
    assert ("payroll_run", "create") in kinds and ("payroll_run", "approve") in kinds
    assert ("payroll_history_import", "create") in kinds
    salary_rows = [a for a in audits if a.entity_type == "employee_salary"]
    redated = [
        a for a in salary_rows if a.entity_id == live["asha"].id and a.action.value == "update"
    ]
    assert redated and redated[-1].changes["before"]["effective_from"] == "2024-09-01"
    assert redated[-1].changes["after"]["effective_from"] == "2024-01-01"
    deleted = [
        a for a in salary_rows if a.entity_id == live["dev"].id and a.action.value == "delete"
    ]
    assert deleted and deleted[0].changes["replaced_by"] == "factoHR payroll registers"
    created = [a for a in salary_rows if a.entity_id == bala[0].id]
    assert created[0].action.value == "create"
    assert created[0].changes["source"]["months"] == ["2024-01"]
    approve = [a for a in audits if a.action.value == "approve"]
    assert approve[0].summary == "Locked on import: figures as paid on factoHR"


async def test_months_already_imported_or_run_are_refused_and_nothing_is_written(
    session: AsyncSession, hr_user: User, world
) -> None:
    await history_import.apply(session, hr_user, await _plan(session, hr_user))
    count = await session.scalar(select(func.count(Payslip.id)))

    again = await _plan(session, hr_user)
    refused = [i for i in again.issues if i.category == "month_already_run"]
    assert {i.month for i in refused} == {"2024-01", "2024-02", "2024-03"}
    result = await history_import.apply(session, hr_user, again)
    assert not result.applied
    assert await session.scalar(select(func.count(PayrollRun.id))) == 3
    assert await session.scalar(select(func.count(Payslip.id))) == count


async def test_a_month_with_a_computed_run_is_refused(
    session: AsyncSession, hr_user: User, tenant: Tenant, world
) -> None:
    session.add(
        PayrollRun(
            tenant_id=tenant.id,
            legal_entity_id=world["entity"].id,
            period_year=2024,
            period_month=2,
            effective_date=date(2024, 2, 29),
            run_type=PayrollRunType.REGULAR,
            status=PayrollRunStatus.DRAFT,
            employee_count=0,
            total_gross=0,
            total_deductions=0,
            total_net=0,
            total_employer_cost=0,
            variance_flags=[],
            review_flags=[],
            rule_set_snapshot={},
        )
    )
    await session.commit()
    plan = await _plan(session, hr_user)
    errors = [i for i in plan.issues if i.level == "error"]
    assert [(i.category, i.month) for i in errors] == [("month_already_run", "2024-02")]
    assert "computed in AskMira" in errors[0].message
    assert not (await history_import.apply(session, hr_user, plan)).applied
    assert await session.scalar(select(func.count(Payslip.id))) == 0


async def test_an_unknown_person_blocks_the_whole_import(
    session: AsyncSession, hr_user: User, world
) -> None:
    files = registers()
    extra = book([Sheet("Sheet1", DEC23, [person("CND26-779", "Chandra Pal", date(2026, 5, 18),
                                                  ctc=16000, basic=8000, hra=3200,
                                                  special=4000, workable=30)])])
    files.append(("Salary Register Apr-24.xlsx", extra))
    plan = await _plan(session, hr_user, files)
    assert [(i.category, i.code) for i in plan.issues if i.level == "error"] == [
        ("unmatched_person", "CND26-779")
    ]
    assert plan.people["unmatched"][0]["code"] == "CND26-779"
    assert not (await history_import.apply(session, hr_user, plan)).applied
    assert await session.scalar(select(func.count(PayrollRun.id))) == 0
    assert await session.scalar(select(func.count(EmployeeSalary.id))) == 6


async def test_the_open_salary_of_an_active_person_is_checked_strictly(
    session: AsyncSession, hr_user: User, world
) -> None:
    rows = [_row("CHD23-105", "Esha M", "esha", workable=JAN, a_special=15000)]
    plan = await _plan(session, hr_user, [("Salary Register Jan-24.xlsx",
                                           book([Sheet("Sheet1", DEC23, rows)]))])
    errors = [i for i in plan.issues if i.level == "error"]
    assert [(i.category, i.code) for i in errors] == [("salary_rule", "CHD23-105")]
    assert "more than CTC TOTAL" in errors[0].message


async def test_a_dry_run_writes_and_rolls_back(
    session: AsyncSession, hr_user: User, world
) -> None:
    plan = await _plan(session, hr_user)
    result = await history_import.apply(session, hr_user, plan, dry_run=True)
    assert result.dry_run and not result.applied
    assert result.stats["payslips"] == 25 and result.stats["runs"] == 3
    assert await session.scalar(select(func.count(PayrollRun.id))) == 0
    assert await session.scalar(select(func.count(Payslip.id))) == 0
    assert await session.scalar(select(func.count(EmployeeSalary.id))) == 6
    # The report is plain data and still reads after the rollback.
    assert history_import.preview(result)["payslips"] == 25


async def test_undo_removes_an_imported_month_only(
    session: AsyncSession, hr_user: User, other_hr_user: User, tenant: Tenant, world
) -> None:
    await history_import.apply(session, hr_user, await _plan(session, hr_user))
    run = await session.scalar(select(PayrollRun).where(PayrollRun.period_month == 3))
    revisions = await session.scalar(select(func.count(EmployeeSalary.id)))

    from fastapi import HTTPException

    with pytest.raises(HTTPException) as other:
        await history_import.undo_imported_run(session, other_hr_user, run.id)
    assert other.value.status_code == 404

    await history_import.undo_imported_run(session, hr_user, run.id)
    assert await session.scalar(select(PayrollRun).where(PayrollRun.id == run.id)) is None
    assert await session.scalar(select(func.count(Payslip.id)).where(Payslip.run_id == run.id)) == 0
    assert await session.scalar(select(func.count(EmployeeSalary.id))) == revisions
    audit = await session.scalar(
        select(AuditLog).where(
            AuditLog.entity_id == run.id, AuditLog.action == AuditAction.DELETE
        )
    )
    assert audit is not None and audit.changes["payslips_removed"] == 7

    computed = PayrollRun(
        tenant_id=tenant.id, legal_entity_id=world["entity"].id, period_year=2024,
        period_month=4, effective_date=date(2024, 4, 30), run_type=PayrollRunType.REGULAR,
        status=PayrollRunStatus.DRAFT, employee_count=0, total_gross=0, total_deductions=0,
        total_net=0, total_employer_cost=0, variance_flags=[], review_flags=[],
        rule_set_snapshot={},
    )
    session.add(computed)
    await session.commit()
    with pytest.raises(HTTPException) as refused:
        await history_import.undo_imported_run(session, hr_user, computed.id)
    assert refused.value.status_code == 409

    # The month can be imported again once undone.
    again = await _plan(session, hr_user)
    assert {i.month for i in again.issues if i.category == "month_already_run"} == {
        "2024-01", "2024-02"
    }


# ── The API ─────────────────────────────────────────────────────────────────


@pytest.fixture
async def client(session: AsyncSession) -> AsyncGenerator[httpx.AsyncClient, None]:
    async def override() -> AsyncGenerator[AsyncSession, None]:
        yield session

    app.dependency_overrides[get_session] = override
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://testserver"
        ) as c:
            yield c
    finally:
        app.dependency_overrides.clear()


def _auth(user: User) -> dict[str, str]:
    token = create_access_token(
        user_id=str(user.id), tenant_id=str(user.tenant_id), role=user.role.value
    )
    return {"Authorization": f"Bearer {token}"}


def _upload(files):
    return [("files", (name, data, "application/octet-stream")) for name, data in files]


async def test_the_admin_api_previews_imports_reports_and_undoes(
    client: httpx.AsyncClient, session: AsyncSession, hr_user: User, tenant: Tenant, world
) -> None:
    files = registers()
    preview = await client.post(
        "/api/v1/payroll/history-import/preview", headers=_auth(hr_user), files=_upload(files)
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["applied"] is False and body["blocking_errors"] == 0
    assert body["payslips"] == 25
    assert body["salary"]["increments_date_unknown"] == ["CHD23-103"]
    assert await session.scalar(select(func.count(PayrollRun.id))) == 0

    issues = await client.post(
        "/api/v1/payroll/history-import/issues", headers=_auth(hr_user), files=_upload(files)
    )
    assert issues.status_code == 200
    workbook = load_workbook(io.BytesIO(issues.content))
    assert workbook.sheetnames == ["Read me", "Errors", "Warnings", "Notes", "Salary changes",
                                   "Months"]

    wrong = await client.post(
        "/api/v1/payroll/history-import/preview", headers=_auth(hr_user),
        files=[("files", ("notes.txt", b"x", "text/plain"))],
    )
    assert wrong.status_code == 422

    applied = await client.post(
        "/api/v1/payroll/history-import/apply", headers=_auth(hr_user), files=_upload(files)
    )
    assert applied.status_code == 200, applied.text
    assert applied.json()["applied"] is True and len(applied.json()["run_ids"]) == 3

    run_id = applied.json()["run_ids"][0]
    undone = await client.delete(f"/api/v1/payroll/runs/{run_id}/import", headers=_auth(hr_user))
    assert undone.status_code == 204
    missing = await client.delete(
        f"/api/v1/payroll/runs/{uuid.uuid4()}/import", headers=_auth(hr_user)
    )
    assert missing.status_code == 404


async def test_only_an_admin_may_import(
    client: httpx.AsyncClient, session: AsyncSession, tenant: Tenant, world
) -> None:
    sub_admin = User(tenant_id=tenant.id, email="ops@acme.test", password_hash="x",
                     full_name="Ops", role=UserRole.SUB_ADMIN, is_active=True)
    session.add(sub_admin)
    await session.commit()
    response = await client.post(
        "/api/v1/payroll/history-import/preview", headers=_auth(sub_admin),
        files=_upload(registers()),
    )
    assert response.status_code == 403
