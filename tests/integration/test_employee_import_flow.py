"""Importing employees from a factoHR export, end to end.

1. An admin checks the file (nothing saved), then imports it: current staff
   and leavers, reporting lines inside the file and to people already here,
   new departments and locations, identifiers, bank account, addresses and
   family — each audited, no identifier in the audit trail. The same file
   imported again changes nothing.
2. A sub-admin's import leaves unknown departments and locations blank and
   can leave the leavers out.
3. Somebody already here under another code — the admin made with the
   workspace — is not added twice, nor are two rows with one email; the rows
   left out come back as a sheet that imports again as it was read.
"""

from __future__ import annotations

import io
import json
from datetime import date, datetime

from sqlalchemy import func, select

from app.models import AuditLog, Department, Employee, EmploymentStatus, LegalEntity, Location
from app.models.disbursement import EmployeeBankAccount
from app.models.employee import EmployeeProfile
from app.models.user import UserRole
from app.services import employee_identifiers as ids
from app.services.employee_import.sheet import WHY_COLUMN, read_table
from tests.integration import platform_support
from tests.integration.conftest import requires_db
from tests.integration.platform_support import bearer, make_person
from tests.people import factohr_sheet

pytestmark = requires_db

client = platform_support.client

API = "/api/v1/employees/import"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _aadhaar(base: str = "37877286149") -> str:
    """A twelve-digit number that passes the Verhoeff check."""
    return next(base + str(d) for d in range(10) if ids.verhoeff_valid(base + str(d)))


AADHAAR = _aadhaar()


def _people() -> list[dict]:
    return [
        {
            "Emp Code": "CAP22-001",
            "Full Name": "Asha Verma",
            "First Name": "Asha",
            "Last Name": "Verma",
            "Birth Date": datetime(1994, 3, 15),
            "Joining Date": datetime(2022, 4, 1),
            "Confirmation Date": datetime(2022, 10, 1),
            "Gender": "Female",
            "Status": "Active",
            "Employment Type": "Permanent",
            "Mobile No": "9876543210",
            "Email": "Asha.Verma@example.com",
            "Personal Email": "asha.v@gmail.com",
            "Blood Group": "O positive",
            "Marital Status": "Unmarried",
            "Father Name": "Ramesh Verma",
            "Mother Name": "Sunita Verma",
            "Pan No": "ABCPV1234F",
            "Reporting Manager": "CAP22-002 - Rahul Mehta",
            "Local Address > Address": "12, Lake View Road, Indiranagar",
            "Local Address > City": "Bangalore",
            "Local Address > State": "Karnataka",
            "Local Address > Pincode": "560038",
            "Permanent Same As Local Address": "YES",
            "Companyy": "Example Solutions Pvt Ltd",
            "Department": "Engineering",
            "Designation": "Software Engineer",
            "LOCATION": "Delhi",
            "Bank Detail > Bank Name": "HDFC BANK",
            "Account No": "123456789012",
            "IFSC Code": "HDFC0001234",
            "Name As Per Bank Records": "Asha Verma",
            "PF Detail > PF No": "KRBNG00123450000012345",
            "UA No": "100234223910",
            "ESIC Detail > ESIC No": "5040451052",
            "Aadhaar Card Number Details > Aadhaar No": AADHAAR,
        },
        {
            "Emp Code": "CAP22-002",
            "First Name": "Rahul",
            "Last Name": "Mehta",
            "Joining Date": datetime(2021, 1, 4),
            "Status": "Active",
            "Email": "rahul.mehta@example.com",
            "Pan No": "PANNOTAVBL",
            "Reporting Manager": "OLD-001 - Old Timer",
            "Department": "Talent",
            "LOCATION": "Bangalore office",
        },
        {
            "Emp Code": "CAP21-010",
            "First Name": "Lata",
            "Last Name": "Iyer",
            "Joining Date": datetime(2021, 6, 1),
            "Leaving Date": datetime(2024, 3, 31),
            "Employee Last working date": datetime(2024, 3, 31),
            "Status": "Inactive",
            "Employment Type": "Consultant",
            "Email": "lata.iyer@example.com",
            "Reporting Manager": "CAP22-002 - Rahul Mehta",
        },
        {
            "Emp Code": "OLD-001",
            "First Name": "Old",
            "Joining Date": datetime(2019, 1, 1),
            "Status": "Active",
            "Email": "old@example.com",
        },
        {"Emp Code": "CAP23-099", "First Name": "No", "Joining Date": datetime(2023, 5, 1)},
        {
            "Emp Code": "CAP23-100",
            "First Name": "Gone",
            "Joining Date": datetime(2023, 5, 1),
            "Status": "Inactive",
            "Email": "gone@example.com",
        },
        {
            "Emp Code": "CAP23-101",
            "First Name": "Zara",
            "Joining Date": datetime(2023, 7, 1),
            "Status": "Active",
            "Email": "zara@example.com",
            "LOCATION": "Dubai",
        },
    ]


async def _setup(session, tenant, *, role: UserRole):
    hr = await make_person(session, tenant, role=role, name="Priya Hr")
    entity = LegalEntity(tenant_id=tenant.id, name="Example Solutions Private Limited", code="EXS")
    session.add(entity)
    session.add(Department(tenant_id=tenant.id, name="Engineering", code="ENG"))
    session.add(
        Employee(
            tenant_id=tenant.id,
            employee_code="OLD-001",
            first_name="Old",
            last_name="Timer",
            work_email="old@example.com",
            status=EmploymentStatus.ACTIVE,
            date_of_joining=date(2019, 1, 1),
        )
    )
    await session.commit()
    return hr, entity


def _upload(data: bytes, **form):
    return {"files": {"file": ("Employee Detail 06 OCT 26.xlsx", data, XLSX)}, "data": form}


async def _employee(session, code: str) -> Employee:
    return await session.scalar(select(Employee).where(Employee.employee_code == code))


async def test_an_admin_checks_then_imports_a_factohr_export(client, session, tenant):
    hr, entity = await _setup(session, tenant, role=UserRole.ADMIN)
    headers = bearer(hr)
    data = factohr_sheet.build(_people())

    # ── The check writes nothing ────────────────────────────────────────────
    checked = await client.post(f"{API}/preview", headers=headers, **_upload(data))
    assert checked.status_code == 200, checked.text
    preview = checked.json()
    assert preview["source"] == "factoHR Employee Detail report"
    assert preview["total_rows"] == 7
    assert (preview["ready_current"], preview["ready_leavers"]) == (3, 1)
    assert (preview["skipped"], preview["errors"]) == (1, 2)
    assert preview["new_departments"] == ["Talent"]
    assert {(loc["name"], loc["state_code"]) for loc in preview["new_locations"]} == {
        ("Delhi", "DL"),
        ("Bangalore office", "KA"),
    }
    assert preview["can_create_org"] is True
    rows = {row["employee_code"]: row for row in preview["rows"]}
    assert rows["OLD-001"]["outcome"] == "skip"
    assert rows["CAP23-099"]["errors"] == ["No Email (work email) or Personal Email."]
    assert "no Leaving Date" in rows["CAP23-100"]["errors"][0]
    assert any(w.startswith("Pan No:") for w in rows["CAP22-002"]["warnings"])
    assert any("Dubai" in w for w in rows["CAP23-101"]["warnings"])
    assert rows["CAP21-010"]["status"] == "exited"
    used = {column["field"] for column in preview["columns_used"]}
    assert {"Emp Code", "Aadhaar No", "Local Pincode", "PF No"} <= used
    assert "Title" not in used
    assert await session.scalar(select(func.count(Employee.id))) == 2  # HR and OLD-001

    # ── The import ──────────────────────────────────────────────────────────
    imported = await client.post(API, headers=headers, **_upload(data))
    assert imported.status_code == 201, imported.text
    result = imported.json()
    assert (result["created"], result["created_leavers"]) == (4, 1)
    assert (result["skipped"], result["failed"]) == (1, 2)
    assert result["departments_created"] == ["Talent"]
    assert sorted(result["locations_created"]) == ["Bangalore office", "Delhi"]

    asha = await _employee(session, "CAP22-001")
    rahul = await _employee(session, "CAP22-002")
    lata = await _employee(session, "CAP21-010")
    old = await _employee(session, "OLD-001")
    zara = await _employee(session, "CAP23-101")
    assert asha.status is EmploymentStatus.ACTIVE
    assert asha.work_email == "asha.verma@example.com"
    assert asha.date_of_joining == date(2022, 4, 1)
    assert asha.probation_end_date == date(2022, 10, 1)
    assert asha.legal_entity_id == entity.id
    assert asha.manager_id == rahul.id
    assert rahul.manager_id == old.id
    assert lata.manager_id == rahul.id
    assert (lata.status, lata.date_of_exit, lata.is_active) == (
        EmploymentStatus.EXITED,
        date(2024, 3, 31),
        False,
    )
    assert lata.employment_type.value == "consultant"
    assert zara.location_id is None

    engineering = await session.scalar(select(Department).where(Department.code == "ENG"))
    assert asha.department_id == engineering.id
    delhi = await session.scalar(select(Location).where(Location.name == "Delhi"))
    assert (delhi.state_code, delhi.legal_entity_id) == ("DL", entity.id)
    assert asha.location_id == delhi.id
    talent = await session.scalar(select(Department).where(Department.name == "Talent"))
    assert rahul.department_id == talent.id

    # Identifiers on the record; only the Aadhaar's last four.
    assert (asha.pan, asha.uan, asha.esi_number) == ("ABCPV1234F", "100234223910", "5040451052")
    assert asha.pf_number == "KRBNG00123450000012345"
    assert asha.aadhaar_last4 == AADHAAR[-4:]
    assert rahul.pan is None
    account = await session.scalar(
        select(EmployeeBankAccount).where(EmployeeBankAccount.employee_id == asha.id)
    )
    assert (account.ifsc, account.bank_name, account.is_primary) == (
        "HDFC0001234",
        "HDFC BANK",
        True,
    )
    assert account.account_number != "123456789012"  # stored encrypted

    profile = await session.scalar(
        select(EmployeeProfile).where(EmployeeProfile.employee_id == asha.id)
    )
    assert profile.personal["marital_status"] == "single"
    assert profile.personal["blood_group"] == "O+"
    assert profile.personal["father_name"] == "Ramesh Verma"
    assert profile.addresses["current"]["state_code"] == "KA"
    assert profile.addresses["permanent_same_as_current"] is True
    assert profile.family_members == [
        {
            "name": "Sunita Verma",
            "relationship": "mother",
            "date_of_birth": None,
            "gender": None,
            "occupation": None,
            "is_dependent": False,
        }
    ]

    # Audited: one row per person and one for the import; never an identifier.
    audits = (await session.scalars(select(AuditLog).where(AuditLog.tenant_id == tenant.id))).all()
    imported_rows = [
        a for a in audits if a.entity_type == "employee" and a.action.value == "create"
    ]
    assert len(imported_rows) == 4
    assert any(a.entity_type == "employee_import" for a in audits)
    trail = json.dumps([a.changes for a in audits], default=str)
    for secret in ("ABCPV1234F", AADHAAR, "123456789012", "100234223910"):
        assert secret not in trail

    # ── The same file again: everybody is already here ─────────────────────
    again = await client.post(API, headers=headers, **_upload(data))
    assert again.status_code == 422
    assert "Nothing in this file" in again.json()["detail"]


async def test_a_sub_admin_import_leaves_unknown_places_blank(client, session, tenant):
    hr, _ = await _setup(session, tenant, role=UserRole.SUB_ADMIN)
    headers = bearer(hr)
    data = factohr_sheet.build(_people())

    checked = (await client.post(f"{API}/preview", headers=headers, **_upload(data))).json()
    assert checked["can_create_org"] is False
    assert checked["new_departments"] == []
    rows = {row["employee_code"]: row for row in checked["rows"]}
    assert any("Talent" in w and "admin" in w for w in rows["CAP22-002"]["warnings"])

    imported = await client.post(API, headers=headers, **_upload(data, include_leavers="false"))
    assert imported.status_code == 201, imported.text
    assert (imported.json()["created"], imported.json()["created_leavers"]) == (3, 0)
    assert await _employee(session, "CAP21-010") is None
    rahul = await _employee(session, "CAP22-002")
    assert rahul.department_id is None
    assert await session.scalar(select(func.count(Location.id))) == 0

    # The template comes with the workspace's own fields.
    template = await client.get(f"{API}/template", headers=headers)
    assert template.status_code == 200
    assert template.headers["content-type"] == XLSX


async def test_employees_cannot_import(client, session, tenant):
    staff = await make_person(session, tenant, role=UserRole.EMPLOYEE, name="Esha Staff")
    refused = await client.post(
        f"{API}/preview", headers=bearer(staff), **_upload(factohr_sheet.build(_people()))
    )
    assert refused.status_code == 403


async def test_an_email_already_here_is_not_added_twice(client, session, tenant):
    from openpyxl import load_workbook

    hr, _ = await _setup(session, tenant, role=UserRole.ADMIN)
    headers = bearer(hr)
    mine = await session.get(Employee, hr.employee_id)
    # The admin's own row in the old system's export: another code, the same email.
    data = factohr_sheet.build(
        [
            {
                "Emp Code": "CGB-007",
                "First Name": "Priya",
                "Last Name": "Hr",
                "Joining Date": datetime(2020, 1, 6),
                "Status": "Active",
                "Email": mine.work_email.upper(),
                "Pan No": "ABCPV1234F",
            },
            {
                "Emp Code": "CGB-008",
                "First Name": "Ravi",
                "Joining Date": datetime(2021, 2, 1),
                "Status": "Active",
                "Email": "ravi@example.com",
            },
            {
                "Emp Code": "CGB-009",
                "First Name": "Ravi",
                "Last Name": "Again",
                "Joining Date": datetime(2021, 2, 1),
                "Status": "Active",
                "Email": "Ravi@example.com",
            },
        ]
    )
    preview = (await client.post(f"{API}/preview", headers=headers, **_upload(data))).json()
    rows = {row["employee_code"]: row for row in preview["rows"]}
    assert rows["CGB-007"]["outcome"] == "error"
    assert f"already used by {mine.employee_code} (Priya Hr)" in rows["CGB-007"]["errors"][0]
    assert rows["CGB-008"]["outcome"] == "ready"
    line = rows["CGB-008"]["line"]
    assert rows["CGB-009"]["errors"] == [f"Email ravi@example.com is also on row {line}."]

    # The rows left out, with their own cells and why.
    sheet = await client.post(f"{API}/rejected", headers=headers, **_upload(data))
    assert sheet.status_code == 200, sheet.text
    assert sheet.headers["content-type"] == XLSX
    book = load_workbook(io.BytesIO(sheet.content))
    grid = [list(r) for r in book.active.iter_rows(values_only=True)]
    header = next(i for i, r in enumerate(grid) if r[0] == WHY_COLUMN)
    left_out = {r[grid[header].index("Emp Code")]: r for r in grid[header + 1 :]}
    assert set(left_out) == {"CGB-007", "CGB-009"}
    assert "already used by" in left_out["CGB-007"][0]
    assert left_out["CGB-007"][grid[header].index("Pan No")] == "ABCPV1234F"
    # Read back the way the original was: every field from the same column.
    original, again = read_table(data, "a.xlsx"), read_table(sheet.content, "b.xlsx")
    by_name = {c.index: c.header for c in original.columns}
    assert {k: by_name[i] for k, i in original.mapping.items()} == {
        k: {c.index: c.header for c in again.columns}[i] for k, i in again.mapping.items()
    }

    imported = (await client.post(API, headers=headers, **_upload(data))).json()
    assert (imported["created"], imported["failed"]) == (1, 2)
    assert await _employee(session, "CGB-007") is None
    assert await session.scalar(
        select(func.count(Employee.id)).where(Employee.work_email == mine.work_email)
    ) == 1

    # Nothing left out: nothing to download.
    clean = factohr_sheet.build(
        [{"Emp Code": "CGB-010", "First Name": "Neha", "Joining Date": datetime(2022, 1, 3),
          "Status": "Active", "Email": "neha@example.com"}]
    )
    nothing = await client.post(f"{API}/rejected", headers=headers, **_upload(clean))
    assert nothing.status_code == 422
