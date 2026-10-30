"""Importing an old HR system's export as it had it: nothing dropped, every difference listed.

1. A person's later rows (a second qualification, family member, nominee)
   fold into them; lists land in the profile; columns AskMira has no field for
   become custom fields on a "factoHR" tab; columns that say nothing (the same
   for everyone, worked out by the old system, a copy of an identifier) are
   listed with why. Values that fail AskMira's checks are kept as typed and
   listed. A leaver may lack or share an email; two people still working here
   may not share one. The issues list says it all, person by person.
2. What HR types into a profile afterwards is still checked as before.
3. A payslip opens with the date of birth when the PAN on file is the old
   system's "PANNOTAVBL".
4. A leaver's salary is a record of what they were paid: kept as typed (a
   negative amount, more in the CTC than CTC TOTAL, a CTC alone). Somebody
   still working here gets the usual checks.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models import Employee, EmploymentStatus, LegalEntity
from app.models.disbursement import EmployeeBankAccount
from app.models.employee import CustomFieldDefinition, EmployeeProfile
from app.models.payroll import EmployeeSalary, SalaryStructure
from app.models.tax import PreviousEmployerIncome, ProofStatus
from app.models.user import UserRole
from app.services import employee_identifiers as ids
from app.services.payroll import salary_import
from app.services.payroll.payslips import payslip_password
from tests.integration import platform_support
from tests.integration.conftest import requires_db
from tests.integration.platform_support import bearer, make_person
from tests.people import factohr_sheet

pytestmark = requires_db

client = platform_support.client

API = "/api/v1/employees/import"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: factoHR's layout, cut down to what these tests need. A group heading sits on
#: the first column of its group; the second emergency contact repeats the
#: names under the same heading, as the real export does.
COLUMNS: list[tuple[str, str | None]] = [
    ("S#", None),
    ("Emp Code", None),
    ("Full Name", None),
    ("First Name", None),
    ("Title", None),
    ("Joining Date", None),
    ("Leaving Date", None),
    ("Employee Last working date", None),
    ("Status", None),
    ("Email", None),
    ("Pan No", None),
    ("Currency", None),
    ("Age", None),
    ("Old Code", None),
    ("Machine Code", None),
    ("Companyy", None),
    ("Address", "Local Address"),
    ("City", None),
    ("Pincode", None),
    ("Account No", "Bank Detail"),
    ("IFSC Code", None),
    ("Name As Per Bank Records", None),
    ("Year", "Qualification Detail"),
    ("Qualifications", None),
    ("Course", None),
    ("Institute", None),
    ("From Date", None),
    ("To Date", None),
    ("Percentage", None),
    ("Financial Year", "Past Employee History"),
    ("Income After Exemption", None),
    ("TDS / Tax Paid", None),
    ("Aadhaar No", "Aadhaar Card Number Details"),
    ("PAN No", "Permanent Account Number Details"),
    ("Name", "Family Detail"),
    ("Relation", None),
    ("Date of Birth", None),
    ("Contact No", None),
    ("Approved", None),
    ("Name", "Dependant Detail"),
    ("NominationType", None),
    ("Relation", None),
    ("Status", None),
    ("Name", "Emergency Contact Detail"),
    ("Relationship", None),
    ("Mobile 1", None),
    ("Address", None),
    ("Name", None),
    ("Relationship", None),
    ("Mobile 1", None),
]


def _bad_aadhaar() -> str:
    """Twelve digits that fail the Verhoeff check."""
    return next(f"37877286149{d}" for d in range(10) if not ids.verhoeff_valid(f"37877286149{d}"))


AADHAAR = _bad_aadhaar()


def _people() -> list[dict]:
    always = {"Currency": "Indian Rupees"}
    return [
        {
            **always,
            "Emp Code": "EMP-1",
            "Full Name": "Asha Rao",
            "First Name": "Asha",
            "Title": "Ms",
            "Joining Date": datetime(2022, 4, 1),
            "Status": "Active",
            "Email": "asha@example.com",
            "Pan No": "PANNOTAVBL",
            "Age": 30,
            "Old Code": "IN-45",
            "Machine Code": "0042",
            "Companyy": "Acme Pvt Ltd",
            "Local Address > Address": "12 Lake Road, Indiranagar",
            "Account No": "0012345678",
            "IFSC Code": "HDFCO001234",
            "Name As Per Bank Records": "Asha Rao",
            "Year": "2020-2022",
            "Qualifications": "M.C.A.",
            "Course": "Post Graduate",
            "Institute": "SRM University",
            "From Date": datetime(2020, 8, 1),
            "To Date": datetime(2022, 5, 31),
            "Percentage": "70",
            "Financial Year": "2024-2025",
            "Income After Exemption": 1220120,
            "TDS / Tax Paid": 0,
            "Aadhaar No": AADHAAR,
            "Permanent Account Number Details > PAN No": "PANNOTAVBL",
            "Family Detail > Name": "Ravi Rao",
            "Family Detail > Relation": "Spouse",
            "Date of Birth": datetime(1990, 1, 1),
            "Contact No": "9876543210",
            "Approved": "NO",
            "Dependant Detail > Name": "Ravi Rao",
            "NominationType": "Pension",
            "Dependant Detail > Relation": "Spouse",
            "Dependant Detail > Status": "Unknown",
            "Emergency Contact Detail > Name": "Ravi Rao",
            "Emergency Contact Detail > Relationship": "Husband",
            "Emergency Contact Detail > Mobile 1": "9876543210",
            "Emergency Contact Detail > Address": "Same house",
            "Emergency Contact Detail > Name#2": "Home",
            "Emergency Contact Detail > Relationship#2": "Home",
        },
        {
            # Asha's second row: only the lists continue.
            "Emp Code": "EMP-1",
            "Full Name": "Asha Rao",
            "Qualifications": "B.Sc.",
            "Course": "Graduate",
            "Financial Year": "2023-2024",
            "Income After Exemption": 0,
            "TDS / Tax Paid": 0,
            "Family Detail > Name": "Kavi Rao",
            "Family Detail > Relation": "Son",
            "Approved": "NO",
            "Dependant Detail > Name": "Kavi Rao",
            "NominationType": "PF",
            "Dependant Detail > Relation": "Son",
            "Dependant Detail > Status": "Unknown",
        },
        {
            **always,
            "Emp Code": "EMP-2",
            "Full Name": "Bala",
            "First Name": "Bala",
            "Joining Date": datetime(2020, 1, 1),
            "Leaving Date": datetime(2023, 7, 12),
            "Employee Last working date": datetime(2023, 7, 23),
            "Status": "Inactive",
            "Age": 40,
            "Old Code": "IN-46",
            "Companyy": "Acme Pvt Ltd",
        },
        {
            **always,
            "Emp Code": "EMP-3",
            "Full Name": "Chitra",
            "First Name": "Chitra",
            "Joining Date": datetime(2024, 1, 1),
            "Status": "Active",
            "Email": "shared@example.com",
            "Age": 28,
            "Machine Code": "0043",
        },
        {
            **always,
            "Emp Code": "EMP-4",
            "Full Name": "Chitra (old code)",
            "First Name": "Chitra",
            "Joining Date": datetime(2019, 1, 1),
            "Leaving Date": datetime(2020, 1, 1),
            "Status": "Withdrawn",
            "Email": "shared@example.com",
            "Age": 28,
        },
        {
            **always,
            "Emp Code": "EMP-5",
            "Full Name": "Dev",
            "First Name": "Dev",
            "Joining Date": datetime(2021, 1, 1),
            "Status": "TemporaryDeactive",
            "Email": "dev@example.com",
            "Age": 35,
        },
        {
            **always,
            "Emp Code": "EMP-6",
            "Full Name": "Esha",
            "First Name": "Esha",
            "Joining Date": datetime(2025, 1, 1),
            "Status": "Active",
            "Email": "SHARED@example.com",
            "Age": 26,
        },
    ]


def _upload(data: bytes, **form):
    return {"files": {"file": ("Employee Detail.xlsx", data, XLSX)}, "data": form}


async def _setup(session, tenant):
    admin = await make_person(session, tenant, role=UserRole.ADMIN, name="Priya Hr")
    session.add(
        LegalEntity(tenant_id=tenant.id, name="Example Solutions Private Limited", code="EXS")
    )
    await session.commit()
    return admin


async def _employee(session, code: str) -> Employee:
    return await session.scalar(select(Employee).where(Employee.employee_code == code))


async def test_an_export_comes_in_as_typed_with_every_difference_listed(client, session, tenant):
    admin = await _setup(session, tenant)
    headers = bearer(admin)
    data = factohr_sheet.build(_people(), columns=COLUMNS, dates_as_serials=False)

    checked = await client.post(f"{API}/preview", headers=headers, **_upload(data))
    assert checked.status_code == 200, checked.text
    preview = checked.json()
    assert (preview["total_rows"], preview["continuation_rows"]) == (6, 1)
    assert (preview["ready_current"], preview["ready_leavers"], preview["errors"]) == (2, 3, 1)
    rows = {row["employee_code"]: row for row in preview["rows"]}
    assert rows["EMP-6"]["errors"] == [
        f"Email shared@example.com is also on row {rows['EMP-3']['line']}, and both are still "
        "working here."
    ]
    notes = " ".join(preview["notes"])
    assert "'Acme Pvt Ltd' (2) recorded under Example Solutions Private Limited" in notes
    assert "Leaving Date: differs from Employee Last working date for 1 person" in notes
    assert "Past Employee History: 1 financial year(s) with every amount 0" in notes
    assert {
        "Title",
        "Old Code",
        "Machine Code",
        "Company (as in the file)",
        "Leaving Date (where it differs)",
    } <= set(preview["new_custom_fields"])
    not_kept = {c["header"]: c["why"] for c in preview["columns_not_kept"]}
    assert "Indian Rupees" in not_kept["Currency"]
    assert "Calculated" in not_kept["Age"]
    assert "Pan No" in not_kept["Permanent Account Number Details > PAN No"]
    assert "NO (2)" in not_kept["Family Detail > Approved"]
    issues = " | ".join(rows["EMP-1"]["issues"])
    for expected in (
        "Pan No 'PANNOTAVBL'",
        f"Aadhaar No '{AADHAAR}'",
        "Bank account:",
        "Local address imported as typed, without its city and state and PIN code",
        "Emergency contact Home has no mobile number",
        "Nominee share not in the sheet",
    ):
        assert expected in issues, expected
    assert any("no Leaving Date" in i for i in rows["EMP-5"]["issues"])
    assert any("No Email or Personal Email" in i for i in rows["EMP-2"]["issues"])
    assert any("EMP-3" in i for i in rows["EMP-4"]["issues"])

    imported = await client.post(API, headers=headers, **_upload(data))
    assert imported.status_code == 201, imported.text
    result = imported.json()
    assert (result["created"], result["created_leavers"], result["failed"]) == (5, 3, 1)
    assert result["departments_created"] == []

    asha = await _employee(session, "EMP-1")
    definitions = {
        d.label: d
        for d in (
            await session.scalars(
                select(CustomFieldDefinition).where(CustomFieldDefinition.tenant_id == tenant.id)
            )
        ).all()
    }
    assert all(d.tab == "factoHR" and not d.is_required for d in definitions.values())
    custom = {label: asha.custom_fields.get(d.key) for label, d in definitions.items()}
    assert (custom["Title"], custom["Old Code"], custom["Machine Code"]) == ("Ms", "IN-45", "0042")
    assert custom["Company (as in the file)"] == "Acme Pvt Ltd"
    # As typed, each failing value listed above.
    assert (asha.pan, asha.aadhaar_number, asha.aadhaar_last4) == (
        "PANNOTAVBL",
        AADHAAR,
        AADHAAR[-4:],
    )
    account = await session.scalar(
        select(EmployeeBankAccount).where(EmployeeBankAccount.employee_id == asha.id)
    )
    assert account.ifsc == "HDFCO001234"

    profile = await session.scalar(
        select(EmployeeProfile).where(EmployeeProfile.employee_id == asha.id)
    )
    current = profile.addresses["current"]
    assert current["line1"] == "12 Lake Road, Indiranagar"
    assert (current["city"], current["state_code"], current["pincode"]) == (None, None, None)
    assert [(e["qualification"], e["year"], e["institution"]) for e in profile.education] == [
        ("M.C.A.", "2020-2022", "SRM University"),
        ("B.Sc.", None, None),
    ]
    assert profile.education[0]["course"] == "Post Graduate"
    assert profile.education[0]["level"] == "postgraduate"
    assert profile.education[0]["percentage"] == "70"
    assert [(m["name"], m["relationship"]) for m in profile.family_members] == [
        ("Ravi Rao", "spouse"),
        ("Kavi Rao", "child"),
    ]
    assert profile.family_members[0]["phone"] == "9876543210"
    assert [
        (n["name"], n["purpose"], n["purpose_text"], n["share_percent"]) for n in profile.nominees
    ] == [("Ravi Rao", "eps", "Pension", None), ("Kavi Rao", "pf", None, None)]
    assert profile.nominees[0]["remarks"] == "Status: Unknown"
    contacts = profile.emergency_contacts
    assert [
        (c["name"], c["relationship"], c["relationship_text"], c["phone"]) for c in contacts
    ] == [
        ("Ravi Rao", "spouse", None, "9876543210"),
        ("Home", "other", "Home", None),
    ]
    assert contacts[0]["address"] == "Same house"
    [income] = (
        await session.scalars(
            select(PreviousEmployerIncome).where(PreviousEmployerIncome.employee_id == asha.id)
        )
    ).all()
    assert (income.financial_year, Decimal(str(income.gross_salary)), income.status) == (
        "2024-25",
        Decimal("1220120.00"),
        ProofStatus.VERIFIED,
    )

    bala = await _employee(session, "EMP-2")
    assert (bala.work_email, bala.status, bala.date_of_exit) == (
        "",
        EmploymentStatus.EXITED,
        date(2023, 7, 23),
    )
    differs = definitions["Leaving Date (where it differs)"]
    assert bala.custom_fields[differs.key] == "2023-07-12"
    assert differs.key not in (await _employee(session, "EMP-4")).custom_fields
    dev = await _employee(session, "EMP-5")
    assert (dev.status, dev.date_of_exit) == (EmploymentStatus.EXITED, None)
    old = await _employee(session, "EMP-4")
    assert (old.work_email, old.status) == ("shared@example.com", EmploymentStatus.EXITED)
    assert await _employee(session, "EMP-6") is None

    # The issues list: one line per point, the rows not added too.
    listed = await client.post(f"{API}/issues", headers=headers, **_upload(data))
    assert listed.status_code == 200, listed.text
    book = load_workbook(io.BytesIO(listed.content))
    lines = [list(r) for r in book["Issues"].iter_rows(values_only=True)]
    assert lines[0] == ["Emp Code", "Name", "Row", "Status", "Kind", "Issue"]
    kinds = {(line[0], line[4]) for line in lines[1:]}
    assert ("EMP-6", "Not added") in kinds and ("EMP-1", "Imported as typed") in kinds
    assert "Columns not kept" in book.sheetnames and "About the file" in book.sheetnames


async def test_what_hr_types_into_a_profile_is_still_checked(client, session, tenant):
    admin = await _setup(session, tenant)
    person = await make_person(session, tenant, role=UserRole.EMPLOYEE, name="Ravi Kumar")
    url = f"/api/v1/employees/{person.employee_id}/profile"
    headers = bearer(admin)
    refused = {
        "addresses": {
            "current": {"line1": "12 Lake Road"},
            "permanent": None,
            "permanent_same_as_current": False,
        },
        "nominees": [{"name": "Sita", "relationship": "spouse", "purpose": "pf"}],
        "emergency_contacts": [{"name": "Sita", "relationship": "spouse", "is_primary": True}],
        "education": [{"level": "graduate", "qualification": "B.Sc."}],
    }
    for section, data in refused.items():
        response = await client.put(
            url, headers=headers, json={"change": {"section": section, "data": data}}
        )
        assert response.status_code == 422, (section, response.text)
    # The fields an import brings are kept when typed here too.
    saved = await client.put(
        url,
        headers=headers,
        json={
            "change": {
                "section": "family_members",
                "data": [
                    {
                        "name": "Sita",
                        "relationship": "spouse",
                        "phone": "9876543210",
                        "blood_group": "B positive",
                        "health_insurance": True,
                        "title": "Smt",
                    }
                ],
            }
        },
    )
    assert saved.status_code == 200, saved.text
    [member] = saved.json()["family_members"]
    assert (
        member["phone"],
        member["blood_group"],
        member["health_insurance"],
        member["title"],
    ) == (
        "9876543210",
        "B positive",
        True,
        "Smt",
    )


async def test_a_payslip_opens_with_the_birth_date_when_the_pan_is_a_placeholder(session, tenant):
    person = Employee(
        tenant_id=tenant.id,
        employee_code="EMP-9",
        first_name="Asha",
        work_email="asha9@example.com",
        date_of_birth=date(1994, 3, 15),
        pan="PANNOTAVBL",
        status=EmploymentStatus.ACTIVE,
    )
    assert payslip_password(person) == "15031994"
    person.pan = "abcpe1234f"
    assert payslip_password(person) == "ABCPE1234F"


# ── Salary: a leaver's record kept as typed ───────────────────────────────────


def _salary_sheet(rows: list[dict]) -> bytes:
    headers = list(salary_import.LAYOUT)
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(h) for h in headers])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def _salary_setup(client, session, tenant):
    admin = await _setup(session, tenant)
    response = await client.post("/api/v1/salary-components/standard", headers=bearer(admin))
    assert response.status_code == 200, response.text
    people = {}
    for code, status, left in (
        ("LFT-1", EmploymentStatus.EXITED, date(2024, 3, 31)),
        ("CUR-1", EmploymentStatus.ACTIVE, None),
    ):
        person = Employee(
            tenant_id=tenant.id,
            employee_code=code,
            first_name=code,
            work_email=f"{code.lower()}@example.com",
            status=status,
            date_of_joining=date(2022, 1, 1),
            date_of_exit=left,
        )
        session.add(person)
        people[code] = person
    await session.commit()
    return admin, people


async def test_a_leavers_salary_is_kept_as_typed_and_noted(client, session, tenant):
    admin, people = await _salary_setup(client, session, tenant)
    base = {"WEF": datetime(2023, 4, 1), "Pay Basis": "Yearly"}
    data = _salary_sheet(
        [
            # Negative special allowance, and more in the CTC than CTC TOTAL.
            {
                **base,
                "Employee Code": "LFT-1",
                "CTC TOTAL": 500400,
                "BASIC SALARY": 600000,
                "SPECIAL ALLOWANCE": -12000,
                "GRATUITY AMOUNT MANUAL": 28860,
                "GRATUITY PROVISION": 28860,
            },
        ]
    )
    result = await salary_import.upload(session, admin, data, "s.xlsx", apply=True)
    assert result.applied and result.errors == 0
    [row] = result.rows
    assert any("cannot be negative" in note and "Kept as typed" in note for note in row.notes)
    assert any("more than CTC TOTAL" in note for note in row.notes)
    salary = await session.scalar(
        select(EmployeeSalary).where(EmployeeSalary.employee_id == people["LFT-1"].id)
    )
    amounts = {c.component_code: Decimal(str(c.annual_amount)) for c in salary.components}
    assert amounts["SPECIAL"] == Decimal("-12000.00")
    assert Decimal(str(salary.ctc_annual)) == Decimal("500400.00")


async def test_a_leavers_ctc_alone_gets_no_breakdown(client, session, tenant):
    admin, people = await _salary_setup(client, session, tenant)
    data = _salary_sheet(
        [
            {
                "Employee Code": "LFT-1",
                "WEF": datetime(2023, 4, 1),
                "Pay Basis": "Yearly",
                "CTC TOTAL": 480000,
            }
        ]
    )
    result = await salary_import.upload(session, admin, data, "s.xlsx", apply=True)
    assert result.applied
    salary = await session.scalar(
        select(EmployeeSalary).where(EmployeeSalary.employee_id == people["LFT-1"].id)
    )
    assert (Decimal(str(salary.ctc_annual)), salary.components) == (Decimal("480000.00"), [])


async def test_somebody_still_working_here_gets_the_usual_checks(client, session, tenant):
    admin, _ = await _salary_setup(client, session, tenant)
    base = {"Employee Code": "CUR-1", "WEF": datetime(2026, 4, 1), "Pay Basis": "Yearly"}
    negative = _salary_sheet(
        [{**base, "CTC TOTAL": 600000, "BASIC SALARY": 300000, "SPECIAL ALLOWANCE": -1}]
    )
    over = _salary_sheet([{**base, "CTC TOTAL": 100000, "BASIC SALARY": 300000}])
    for data, message in ((negative, "cannot be negative"), (over, "more than CTC TOTAL")):
        result = await salary_import.upload(session, admin, data, "s.xlsx", apply=False)
        [row] = result.rows
        assert row.status == "error" and message in row.message, row.message
    # A CTC alone is split by the structure, as before.
    alone = _salary_sheet([{**base, "CTC TOTAL": 600000}])
    result = await salary_import.upload(session, admin, alone, "s.xlsx", apply=True)
    assert result.applied
    salary = await session.scalar(select(EmployeeSalary))
    assert salary.components, "split by the structure"
    assert await session.scalar(select(SalaryStructure)) is not None
