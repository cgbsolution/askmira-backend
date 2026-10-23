"""Salary master, second part: the standard setup, wage groups, the change trail, bulk upload.

Gates:
  * The standard setup only adds; it never changes a wage type a workspace has.
  * Every change to a person's salary shows on their trail, with who made it.
  * A bulk upload saves nothing while any row has an error, and saves every
    ready row in one transaction when none has.
"""

from __future__ import annotations

import io
from collections.abc import AsyncGenerator
from datetime import date
from decimal import Decimal

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.security import create_access_token
from app.db.session import get_session
from app.main import app
from app.models.employee import Employee
from app.models.org import LegalEntity
from app.models.payroll import (
    CalculationType,
    ComponentType,
    EmployeeSalary,
    PayrollRun,
    PayrollRunStatus,
    Payslip,
    SalaryComponent,
    SalaryStructure,
)
from app.models.tenant import Tenant
from app.models.user import User
from app.schemas.payroll import EmployeeSalaryCreate
from app.services.payroll import service as salary_service
from tests.integration.conftest import requires_db

pytestmark = requires_db

PDF = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n"


@pytest.fixture(autouse=True)
def _storage_root(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "FILE_STORAGE_ROOT", str(tmp_path / "store"))


@pytest.fixture
async def client(session: AsyncSession) -> AsyncGenerator[httpx.AsyncClient, None]:
    async def override() -> AsyncGenerator[AsyncSession, None]:
        yield session

    app.dependency_overrides[get_session] = override
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://testserver"
        ) as c:
            yield c
    finally:
        app.dependency_overrides.clear()


def _auth(user: User) -> dict[str, str]:
    token = create_access_token(
        user_id=str(user.id), tenant_id=str(user.tenant_id), role=user.role.value
    )
    return {"Authorization": f"Bearer {token}"}


async def _standard(client: httpx.AsyncClient, hr: User) -> dict:
    response = await client.post("/api/v1/salary-components/standard", headers=_auth(hr))
    assert response.status_code == 200, response.text
    return response.json()


def _sheet(rows: list[list], headers: list[str]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Salaries"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def _upload(client, hr, data: bytes, *, apply: bool = False) -> httpx.Response:
    return await client.post(
        "/api/v1/salaries/import",
        headers=_auth(hr),
        files={"file": ("salaries.xlsx", data, "application/octet-stream")},
        data={"apply": "true" if apply else "false"},
    )


# ═══════════════════════════════════════════════════════════════════════════
# Standard setup and wage groups
# ═══════════════════════════════════════════════════════════════════════════


async def test_the_standard_setup_creates_wage_types_and_a_working_structure(
    client: httpx.AsyncClient, session: AsyncSession, hr_user: User, employee: Employee
) -> None:
    result = await _standard(client, hr_user)
    assert "BASIC" in result["components_added"]
    assert "GRATUITY" in result["components_added"]
    assert result["structure_name"] == "Standard CTC"

    structure = await session.scalar(select(SalaryStructure))
    preview = await client.post(
        f"/api/v1/salary-structures/{structure.id}/preview",
        headers=_auth(hr_user),
        json={"ctc_annual": "620000"},
    )
    assert preview.status_code == 200, preview.text
    by_code = {c["code"]: c for c in preview.json()["components"]}
    assert by_code["BASIC"]["annual_amount"] == "310000"  # 50% of CTC
    assert by_code["HRA"]["annual_amount"] == "124000"  # 40% of basic
    assert by_code["GRATUITY"]["annual_amount"] == "14911"  # 4.81% of basic
    assert by_code["GRATUITY"]["wage_group"] == "provision"
    groups = {g["wage_group"]: g for g in preview.json()["group_totals"]}
    assert groups["regular_earning"]["annual"] == "605089.00"
    assert groups["provision"]["annual"] == "14911.00"

    # Bonus is a variable earning, typed per person.
    components = (await client.get("/api/v1/salary-components", headers=_auth(hr_user))).json()
    assert {c["code"]: c["wage_group"] for c in components}["BONUS"] == "variable_earning"


async def test_the_standard_setup_leaves_existing_wage_types_alone(
    client: httpx.AsyncClient, session: AsyncSession, tenant: Tenant, hr_user: User
) -> None:
    session.add(
        SalaryComponent(
            tenant_id=tenant.id,
            code="BASIC",
            name="Our basic",
            component_type=ComponentType.EARNING,
            calculation_type=CalculationType.FIXED,
            value=Decimal("1000"),
        )
    )
    await session.commit()

    result = await _standard(client, hr_user)
    assert "BASIC" not in result["components_added"]
    basic = await session.scalar(select(SalaryComponent).where(SalaryComponent.code == "BASIC"))
    await session.refresh(basic)
    assert basic.name == "Our basic"
    # A fixed BASIC cannot be "50% of CTC", so no structure, and HR is told why.
    assert result["structure_name"] is None
    assert "BASIC" in result["structure_skipped"]
    assert await session.scalar(select(SalaryStructure)) is None

    again = await _standard(client, hr_user)
    assert again["components_added"] == []


async def test_the_standard_structure_states_its_own_split_over_an_existing_basic(
    client: httpx.AsyncClient, session: AsyncSession, tenant: Tenant, hr_user: User
) -> None:
    """A workspace's BASIC at 40% of CTC keeps its own definition, but the standard
    structure built on it is the 50% it says it is."""
    session.add(
        SalaryComponent(
            tenant_id=tenant.id,
            code="BASIC",
            name="Basic",
            component_type=ComponentType.EARNING,
            calculation_type=CalculationType.PERCENTAGE_OF,
            base_component_code="CTC",
            value=Decimal("40"),
        )
    )
    await session.commit()
    result = await _standard(client, hr_user)
    assert result["structure_name"] == "Standard CTC"
    structure = await session.scalar(select(SalaryStructure))
    preview = await client.post(
        f"/api/v1/salary-structures/{structure.id}/preview",
        headers=_auth(hr_user),
        json={"ctc_annual": "1200000"},
    )
    by_code = {c["code"]: c for c in preview.json()["components"]}
    assert by_code["BASIC"]["annual_amount"] == "600000"


async def test_a_wage_group_must_fit_the_kind(client: httpx.AsyncClient, hr_user: User) -> None:
    response = await client.post(
        "/api/v1/salary-components",
        headers=_auth(hr_user),
        json={
            "code": "ODD",
            "name": "Odd",
            "component_type": "deduction",
            "wage_group": "variable_earning",
        },
    )
    assert response.status_code == 422


# ═══════════════════════════════════════════════════════════════════════════
# The trail
# ═══════════════════════════════════════════════════════════════════════════


async def test_every_change_shows_on_the_trail_with_who_made_it(
    client: httpx.AsyncClient, session: AsyncSession, hr_user: User, employee: Employee
) -> None:
    await _standard(client, hr_user)
    structure = await session.scalar(select(SalaryStructure))
    base = f"/api/v1/employees/{employee.id}/salary"
    created = await client.post(
        base,
        headers=_auth(hr_user),
        json={
            "structure_id": str(structure.id),
            "ctc_annual": "600000",
            "effective_from": "2026-04-01",
        },
    )
    assert created.status_code == 201, created.text
    salary_id = created.json()["id"]
    edited = await client.put(
        f"{base}/{salary_id}",
        headers=_auth(hr_user),
        json={
            "structure_id": str(structure.id),
            "ctc_annual": "660000",
            "effective_from": "2026-04-01",
            "components": [{"component_code": "TELEPHONE", "amount": "1000"}],
        },
    )
    assert edited.status_code == 200, edited.text
    attached = await client.post(
        f"{base}/{salary_id}/attachment",
        headers=_auth(hr_user),
        files={"file": ("letter.pdf", PDF, "application/pdf")},
    )
    assert attached.status_code == 200
    deleted = await client.delete(f"{base}/{salary_id}", headers=_auth(hr_user))
    assert deleted.status_code == 200

    trail = await client.get(f"{base}/trail", headers=_auth(hr_user))
    assert trail.status_code == 200
    entries = trail.json()
    assert [e["action"] for e in entries] == ["delete", "update", "update", "create"]
    assert all(e["actor_name"] == "Priya Ramesh" for e in entries)
    edit = entries[2]
    assert edit["changes"]["before"]["ctc_annual"] == "600000.00"
    assert edit["changes"]["after"]["ctc_annual"] == "660000.00"
    assert edit["changes"]["after"]["typed_amounts"] == {"TELEPHONE": "12000.00"}


async def test_the_trail_is_one_persons_only(
    client: httpx.AsyncClient,
    session: AsyncSession,
    tenant: Tenant,
    hr_user: User,
    employee: Employee,
) -> None:
    await _standard(client, hr_user)
    structure = await session.scalar(select(SalaryStructure))
    other = Employee(
        tenant_id=tenant.id, employee_code="ACM002", first_name="Ravi", work_email="r@acme.test"
    )
    session.add(other)
    await session.commit()
    await client.post(
        f"/api/v1/employees/{other.id}/salary",
        headers=_auth(hr_user),
        json={
            "structure_id": str(structure.id),
            "ctc_annual": "500000",
            "effective_from": "2026-04-01",
        },
    )
    trail = await client.get(
        f"/api/v1/employees/{employee.id}/salary/trail", headers=_auth(hr_user)
    )
    assert trail.json() == []


# ═══════════════════════════════════════════════════════════════════════════
# Bulk upload: the Employee Salary Master sheet
# ═══════════════════════════════════════════════════════════════════════════

#: The user's own template, column for column (EmployeeSalaryMaster.xlsx).
THEIR_HEADERS = [
    "Employee Code", "WEF", "Pay Basis", "Replace all future salary revisions",
    "CTC TOTAL", "MONTHLY GROSS", "GRATUITY AMOUNT MANUAL", "PROF TAX CTC",
    "BASIC SALARY", "HOUSE RENT ALLOWANCE", "CONVEYANCE ALLOWANCE", "SPECIAL ALLOWANCE",
    "EDUCATION ALLOWANCE", "BONUS", "LEAVE ENCASHMENT", "MEDICAL REIMBURSEMENT",
    "FUEL REIMBURSEMENT", "BONUS PROVISION", "GRATUITY PROVISION", "TDS MANUAL",
    "PROF. TAX MANUAL", "MPF CONTRIBUTION MANUAL", "ESIC EMPLOYEE CONTRIBUTION MANUAL",
    "LWF MANUAL", "INSURANCE", "Uniform Allowance", "Telephone Allowance",
]  # fmt: skip


def _row(**values) -> list:
    """A row of the user's sheet, by header."""
    return [values.get(h) for h in THEIR_HEADERS]


def _their_sheet(rows: list[list], headers: list[str] = THEIR_HEADERS) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sample = book.create_sheet("Sample")
    sample.append(headers)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def test_the_template_is_the_employee_salary_master_layout(
    client: httpx.AsyncClient, hr_user: User, employee: Employee
) -> None:
    await _standard(client, hr_user)
    response = await client.get("/api/v1/salaries/import/template", headers=_auth(hr_user))
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    sheet = book.worksheets[0]
    headers = [c.value for c in sheet[1]]
    assert headers[: len(THEIR_HEADERS)] == THEIR_HEADERS
    assert headers[len(THEIR_HEADERS) :] == ["Salary Structure", "Reason", "Employee Name"]
    assert sheet["A2"].value == "ACM001"
    assert {"Sample", "Instructions"} <= set(book.sheetnames)


async def test_their_empty_template_and_an_untouched_one_save_nothing(
    client: httpx.AsyncClient, hr_user: User, employee: Employee
) -> None:
    await _standard(client, hr_user)
    empty = (await _upload(client, hr_user, _their_sheet([]))).json()
    assert (empty["ready"], empty["errors"], empty["skipped"]) == (0, 0, 0)

    template = await client.get("/api/v1/salaries/import/template", headers=_auth(hr_user))
    untouched = (await _upload(client, hr_user, template.content)).json()
    assert (untouched["ready"], untouched["errors"], untouched["skipped"]) == (0, 0, 1)
    refused = await _upload(client, hr_user, template.content, apply=True)
    assert refused.status_code == 422


async def test_an_old_system_row_keeps_every_amount_and_says_where_the_ctc_differs(
    client: httpx.AsyncClient, session: AsyncSession, hr_user: User, employee: Employee
) -> None:
    """The row from the user's own Salary Master screenshot."""
    await _standard(client, hr_user)
    data = _their_sheet(
        [
            _row(
                **{
                    "Employee Code": "ACM001",
                    "WEF": "01-Apr-2026",
                    "Pay Basis": "Monthly",
                    "CTC TOTAL": "51666.67",
                    "MONTHLY GROSS": 48474,
                    "GRATUITY AMOUNT MANUAL": 1242,
                    "BASIC SALARY": 25833,
                    "HOUSE RENT ALLOWANCE": 10334,
                    "SPECIAL ALLOWANCE": 12307,
                    "PROF. TAX MANUAL": "0.01",
                    "INSURANCE": 239,
                }
            )
        ]
    )
    checked = await _upload(client, hr_user, data)
    assert checked.status_code == 200, checked.text
    [row] = checked.json()["rows"]
    assert row["status"] == "ready", row
    assert row["ctc_annual"] == "596592"
    assert row["gross_monthly"] == "48474"
    notes = " | ".join(row["notes"])
    assert "₹5,96,592" in notes and "₹6,20,000" in notes and "₹23,408" in notes
    assert "PROF. TAX MANUAL 0.01 not used" in notes
    assert "MONTHLY GROSS" not in notes  # it matches

    saved = await _upload(client, hr_user, data, apply=True)
    assert saved.status_code == 200, saved.text
    salary = await session.scalar(select(EmployeeSalary))
    monthly = {c.component_code: Decimal(str(c.monthly_amount)) for c in salary.components}
    assert monthly["BASIC"] == Decimal("25833")
    assert monthly["HRA"] == Decimal("10334")
    assert monthly["SPECIAL"] == Decimal("12307")
    assert monthly["GRATUITY"] == Decimal("1242")
    assert monthly["INSURANCE"] == Decimal("239")


async def test_a_ctc_alone_is_split_by_the_structure(
    client: httpx.AsyncClient, session: AsyncSession, hr_user: User, employee: Employee
) -> None:
    await _standard(client, hr_user)
    data = _their_sheet(
        [_row(**{"Employee Code": "acm001", "WEF": date(2026, 4, 1), "Pay Basis": "Yearly",
                 "CTC TOTAL": 1200000, "Telephone Allowance": 12000})]
    )  # fmt: skip
    body = (await _upload(client, hr_user, data)).json()
    [row] = body["rows"]
    assert row["status"] == "ready", row
    assert row["ctc_annual"] == "1200000"
    assert row["notes"] == []
    assert (await _upload(client, hr_user, data, apply=True)).status_code == 200
    salary = await session.scalar(select(EmployeeSalary))
    telephone = next(c for c in salary.components if c.component_code == "TELEPHONE")
    assert telephone.pay_basis == "annual" and telephone.is_manual
    assert Decimal(str(telephone.monthly_amount)) == Decimal("1000")


async def test_replace_all_future_revisions_puts_the_upload_in_their_place(
    client: httpx.AsyncClient, session: AsyncSession, hr_user: User, employee: Employee
) -> None:
    await _standard(client, hr_user)
    structure = await session.scalar(select(SalaryStructure))
    for on, ctc in ((date(2026, 4, 1), "600000"), (date(2026, 10, 1), "700000")):
        await salary_service.assign_salary(
            session,
            hr_user,
            employee.id,
            EmployeeSalaryCreate(
                structure_id=structure.id, ctc_annual=Decimal(ctc), effective_from=on
            ),
        )

    def sheet(replace: str | None) -> bytes:
        return _their_sheet(
            [_row(**{"Employee Code": "ACM001", "WEF": "01-07-2026", "Pay Basis": "Yearly",
                     "Replace all future salary revisions": replace, "CTC TOTAL": 660000})]
        )  # fmt: skip

    refused = (await _upload(client, hr_user, sheet(None))).json()["rows"][0]
    assert refused["status"] == "error"
    assert "Replace all future salary revisions" in refused["message"]

    checked = (await _upload(client, hr_user, sheet("Yes"))).json()["rows"][0]
    assert checked["status"] == "ready", checked
    assert any("Replaces 1 revision" in note for note in checked["notes"])
    assert (await _upload(client, hr_user, sheet("Yes"), apply=True)).status_code == 200

    series = (
        await session.scalars(select(EmployeeSalary).order_by(EmployeeSalary.effective_from))
    ).all()
    assert [(s.effective_from, s.effective_to) for s in series] == [
        (date(2026, 4, 1), date(2026, 6, 30)),
        (date(2026, 7, 1), None),
    ]
    trail = (
        await client.get(f"/api/v1/employees/{employee.id}/salary/trail", headers=_auth(hr_user))
    ).json()
    assert any(e["changes"].get("replaced_by_upload") for e in trail)


async def test_a_revision_payroll_paid_is_never_replaced(
    client: httpx.AsyncClient,
    session: AsyncSession,
    tenant: Tenant,
    hr_user: User,
    employee: Employee,
    legal_entity: LegalEntity,
) -> None:
    await _standard(client, hr_user)
    structure = await session.scalar(select(SalaryStructure))
    paid = await salary_service.assign_salary(
        session,
        hr_user,
        employee.id,
        EmployeeSalaryCreate(
            structure_id=structure.id, ctc_annual=Decimal("600000"), effective_from=date(2026, 4, 1)
        ),
    )
    run = PayrollRun(
        tenant_id=tenant.id,
        legal_entity_id=legal_entity.id,
        period_year=2026,
        period_month=4,
        effective_date=date(2026, 4, 30),
        status=PayrollRunStatus.PAID,
    )
    session.add(run)
    await session.flush()
    session.add(
        Payslip(
            tenant_id=tenant.id, run_id=run.id, employee_id=employee.id, employee_salary_id=paid.id
        )
    )
    await session.commit()

    data = _their_sheet(
        [_row(**{"Employee Code": "ACM001", "WEF": "01-03-2026", "Pay Basis": "Yearly",
                 "Replace all future salary revisions": "Yes", "CTC TOTAL": 660000})]
    )  # fmt: skip
    [row] = (await _upload(client, hr_user, data)).json()["rows"]
    assert row["status"] == "error"
    assert "cannot be replaced" in row["message"]


async def test_one_bad_row_blocks_the_whole_sheet(
    client: httpx.AsyncClient, session: AsyncSession, hr_user: User, employee: Employee
) -> None:
    await _standard(client, hr_user)
    data = _their_sheet(
        [
            _row(**{"Employee Code": "ACM001", "WEF": "01-04-2026", "Pay Basis": "Yearly",
                    "CTC TOTAL": 620000}),
            _row(**{"Employee Code": "NOPE1", "WEF": "01-04-2026", "Pay Basis": "Yearly",
                    "CTC TOTAL": 500000}),
            _row(**{"Employee Code": "ACM001", "WEF": "01-05-2026", "Pay Basis": "Weekly",
                    "CTC TOTAL": 10000}),
            _row(**{"Employee Code": "ACM001", "WEF": "01-06-2026", "Pay Basis": "Yearly",
                    "CTC TOTAL": 700000, "BASIC SALARY": "abc"}),
            _row(**{"Employee Code": "ACM001", "WEF": "01-07-2026", "CTC TOTAL": 700000}),
        ]
    )  # fmt: skip
    rows = (await _upload(client, hr_user, data)).json()["rows"]
    assert [r["status"] for r in rows] == ["ready", "error", "error", "error", "error"]
    assert "No employee with code NOPE1" in rows[1]["message"]
    assert "Yearly or Monthly" in rows[2]["message"]
    assert "not an amount" in rows[3]["message"]
    assert "Pay Basis is empty" in rows[4]["message"]

    refused = await _upload(client, hr_user, data, apply=True)
    assert refused.status_code == 422
    assert await session.scalar(select(EmployeeSalary)) is None


async def test_an_unknown_column_matters_only_where_it_holds_a_value(
    client: httpx.AsyncClient, hr_user: User, employee: Employee
) -> None:
    await _standard(client, hr_user)
    headers = [*THEIR_HEADERS, "SHOE ALLOWANCE"]
    base = {
        "Employee Code": "ACM001",
        "WEF": "01-04-2026",
        "Pay Basis": "Yearly",
        "CTC TOTAL": 620000,
    }
    blank = _their_sheet([[*_row(**base), None]], headers)
    assert (await _upload(client, hr_user, blank)).json()["rows"][0]["status"] == "ready"
    filled = _their_sheet([[*_row(**base), 500]], headers)
    [row] = (await _upload(client, hr_user, filled)).json()["rows"]
    assert row["status"] == "error" and "SHOE ALLOWANCE" in row["message"]
