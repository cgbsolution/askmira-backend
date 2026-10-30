"""The employee export: every column of the record, a column choice, Excel or CSV.

HR asked for an export that misses nothing they hold on a person, and for the
choice to leave columns out. These tests fill one person's record in every
section and read the file back.
"""

from __future__ import annotations

import csv
import io
from datetime import date

from cryptography.fernet import Fernet
from openpyxl import load_workbook
from sqlalchemy import select, text

from app.core.crypto import _FIELD_PREFIX, encrypt_secret
from app.models.access import UserAreaRights
from app.models.audit import AuditAction, AuditLog
from app.models.custom_tables import CustomTableDefinition, CustomTableRow
from app.models.disbursement import EmployeeBankAccount
from app.models.employee import CustomFieldDefinition, Employee, EmployeeProfile
from app.models.org import Department
from app.models.user import UserRole
from tests.integration import platform_support
from tests.integration.conftest import requires_db
from tests.integration.platform_support import bearer, make_person

pytestmark = requires_db

client = platform_support.client

API = "/api/v1"
URL = f"{API}/employees/export"


async def _filled(session, tenant, legal_entity, hr) -> Employee:
    """One person with something in every section of the record."""
    department = Department(tenant_id=tenant.id, name="Engineering", code="ENG")
    session.add(department)
    await session.flush()
    login = await make_person(session, tenant, role=UserRole.EMPLOYEE, name="Kavya Menon")
    person = await session.get(Employee, login.employee_id)
    person.employee_code = "EMP0042"
    person.designation = "=HYPERLINK(1)"
    person.legal_entity_id = legal_entity.id
    person.department_id = department.id
    person.manager_id = hr.employee_id
    person.phone = "09876543210"
    person.date_of_birth = date(1993, 3, 15)
    person.date_of_joining = date(2024, 4, 1)
    person.pan = "ABCPE1234F"
    person.aadhaar_number = "234567890123"
    person.aadhaar_last4 = "0123"
    person.uan = "100200300400"
    person.custom_fields = {"tshirt": "M", "old_code": "=SUM(1)"}
    session.add_all(
        [
            EmployeeProfile(
                tenant_id=tenant.id,
                employee_id=person.id,
                personal={"father_name": "Ravi Menon", "marital_status": "married"},
                addresses={
                    "current": {
                        "line1": "12 MG Road",
                        "city": "Bengaluru",
                        "state_code": "KA",
                        "pincode": "560001",
                    },
                    "permanent_same_as_current": True,
                    "permanent": {"line1": "12 MG Road", "city": "Bengaluru"},
                },
                emergency_contacts=[
                    {"name": "Arun", "relationship": "spouse", "phone": "9000000001"}
                ],
                family_members=[
                    {"name": "Arun", "relationship": "spouse", "date_of_birth": "1990-01-02"},
                    {"name": "Asha", "relationship": "other", "relationship_text": "Aunt"},
                ],
                nominees=[
                    {
                        "name": "Arun",
                        "relationship": "spouse",
                        "purpose": "pf",
                        "share_percent": "100",
                    }
                ],
                education=[{"level": "graduate", "qualification": "B.Tech", "is_highest": True}],
                previous_employment=[
                    {"employer": "Globex", "from_date": "2019-01-01", "to_date": "2024-03-31"}
                ],
                hiring_snapshot={"source": "Referral", "expected_ctc": "1200000", "fit_score": 87},
            ),
            EmployeeBankAccount(
                tenant_id=tenant.id,
                employee_id=person.id,
                account_number=encrypt_secret("001234567890", tenant.id),
                ifsc="HDFC0001234",
                account_holder_name="KAVYA MENON",
                bank_name="HDFC Bank",
            ),
            CustomFieldDefinition(
                tenant_id=tenant.id,
                key="tshirt",
                label="T-shirt size",
                field_type="select",
                options=["S", "M"],
            ),
            CustomFieldDefinition(
                tenant_id=tenant.id,
                key="old_code",
                label="Old Code",
                field_type="text",
                tab="factoHR",
            ),
            CustomFieldDefinition(
                tenant_id=tenant.id,
                key="retired",
                label="Retired field",
                field_type="text",
                is_active=False,
            ),
        ]
    )
    table = CustomTableDefinition(
        tenant_id=tenant.id,
        key="certs",
        label="Certification",
        columns=[
            {"key": "name", "label": "Certificate", "field_type": "text"},
            {"key": "earned", "label": "Earned on", "field_type": "date"},
        ],
    )
    session.add(table)
    await session.flush()
    session.add(
        CustomTableRow(
            tenant_id=tenant.id,
            table_id=table.id,
            employee_id=person.id,
            values={"name": "AWS SA", "earned": "2023-05-06"},
        )
    )
    await session.commit()
    return person


def _sheet(content: bytes) -> tuple[list[str], dict[str, dict[str, object]], object]:
    sheet = load_workbook(io.BytesIO(content)).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    by_code = {str(r[0]): dict(zip(headers, r, strict=True)) for r in rows[1:]}
    return headers, by_code, sheet


async def test_every_column_by_default_with_lists_custom_fields_and_identifiers(
    client, session, tenant, legal_entity
):
    hr = await make_person(session, tenant, role=UserRole.ADMIN, name="Priya Hr")
    await _filled(session, tenant, legal_entity, hr)

    catalogue = await client.get(f"{URL}/columns", headers=bearer(hr))
    assert catalogue.status_code == 200, catalogue.text
    sections = {s["key"]: s for s in catalogue.json()["sections"]}
    for key in (
        "basic", "job", "personal", "address", "statutory", "bank", "emergency", "family",
        "nominees", "education", "experience", "hiring", "cf", "cf:factoHR", "ct.certs",
    ):  # fmt: skip
        assert key in sections, key
    assert sections["family"]["repeats"] is True
    labels = [c["label"] for s in sections.values() for c in s["columns"]]
    # Switched-off custom fields and the AI fit score are not offered.
    assert "Retired field" not in labels
    assert not any("fit" in label.lower() for label in labels)

    response = await client.post(URL, json={}, headers=bearer(hr))
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert ".xlsx" in response.headers["content-disposition"]
    headers, rows, sheet = _sheet(response.content)

    # Every single-value column the picker offers is in the file.
    for section in sections.values():
        if not section["repeats"]:
            for column in section["columns"]:
                assert column["label"] in headers, column["label"]
    # Lists: as many column sets as the most anybody has.
    assert "Family member 1 - Name" in headers
    assert "Family member 2 - Relationship" in headers
    assert "Family member 3 - Name" not in headers
    assert "Previous job 1 - Employer" in headers
    assert "Certification 1 - Certificate" in headers

    row = rows["EMP0042"]
    assert row["Name"] == "Kavya Menon"
    assert row["Company"] == legal_entity.name
    assert row["Department"] == "Engineering"
    assert row["Manager"] == "Priya Hr"
    assert row["Portal role"] == "Employee"
    assert row["Login status"] == "Active"
    # Identifiers in full, as text: leading zeros and every digit kept.
    assert row["PAN"] == "ABCPE1234F"
    assert row["Aadhaar"] == "234567890123"
    assert row["Account number"] == "001234567890"
    assert row["Mobile"] == "09876543210"
    assert row["UAN"] == "100200300400"
    assert row["IFSC"] == "HDFC0001234"
    assert row["Name as per bank"] == "KAVYA MENON"
    assert row["Father's name"] == "Ravi Menon"
    assert row["Marital status"] == "Married"
    assert row["Current state"] == "Karnataka"
    assert row["Current PIN code"] == "560001"
    assert row["Permanent same as current"] == "Yes"
    assert row["Family member 1 - Relationship"] == "Spouse"
    assert row["Family member 2 - Relationship"] == "Aunt"
    assert row["Nominee 1 - Nominated for"] == "Provident Fund"
    assert row["Nominee 1 - Share (%)"] == 100
    assert row["Education 1 - Highest qualification"] == "Yes"
    assert row["Hiring source"] == "Referral"
    assert row["T-shirt size"] == "M"
    assert row["Certification 1 - Certificate"] == "AWS SA"
    # Dates are dates in Excel, shown day first.
    assert row["Date of joining"].date() == date(2024, 4, 1)
    assert row["Certification 1 - Earned on"].date() == date(2023, 5, 6)
    line = next(i for i in (2, 3) if sheet.cell(row=i, column=1).value == "EMP0042")
    joining = sheet.cell(row=line, column=headers.index("Date of joining") + 1)
    assert joining.number_format == "DD-MM-YYYY"
    # A value that looks like a formula is text, never a formula.
    assert row["Designation"] == "=HYPERLINK(1)"
    assert row["Old Code"] == "=SUM(1)"
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            assert cell.data_type != "f"

    # The audit row says what left: how many people, and that identifiers did.
    exported = await session.scalar(
        select(AuditLog)
        .where(AuditLog.action == AuditAction.EXPORT, AuditLog.entity_type == "employee_directory")
        .order_by(AuditLog.created_at.desc())
    )
    assert exported.changes["rows"] == 2
    assert exported.changes["all_columns"] is True
    assert exported.changes["identifiers"] == ["aadhaar", "account_number", "pan"]


async def test_hr_chooses_columns_and_the_ticked_people(client, session, tenant, legal_entity):
    hr = await make_person(session, tenant, role=UserRole.ADMIN, name="Priya Hr")
    person = await _filled(session, tenant, legal_entity, hr)

    # Columns in the picker's order, whatever order they were sent in.
    body = {"columns": ["pan", "code", "family.name", "name"], "employee_ids": [str(person.id)]}
    response = await client.post(URL, json=body, headers=bearer(hr))
    assert response.status_code == 200, response.text
    headers, rows, _ = _sheet(response.content)
    assert headers == ["Employee code", "Name", "PAN", "Family member 1 - Name",
                       "Family member 2 - Name"]  # fmt: skip
    assert list(rows) == ["EMP0042"]

    exported = await session.scalar(
        select(AuditLog)
        .where(AuditLog.action == AuditAction.EXPORT)
        .order_by(AuditLog.created_at.desc())
    )
    assert exported.changes["ticked_people"] == 1
    assert exported.changes["all_columns"] is False
    assert exported.changes["identifiers"] == ["pan"]

    # Without PAN, Aadhaar or the account, none of them is in the file.
    body = {"columns": ["code", "name", "bank_name"]}
    response = await client.post(URL, json=body, headers=bearer(hr))
    headers, rows, _ = _sheet(response.content)
    assert headers == ["Employee code", "Name", "Bank name"]
    assert rows["EMP0042"]["Bank name"] == "HDFC Bank"

    # Nothing it knows: refused with a sentence, not an empty file.
    response = await client.post(URL, json={"columns": ["nope"]}, headers=bearer(hr))
    assert response.status_code == 422
    assert "at least one column" in response.json()["detail"]
    assert (await client.post(URL, json={"columns": []}, headers=bearer(hr))).status_code == 422
    bad = {"file_type": "pdf"}
    assert (await client.post(URL, json=bad, headers=bearer(hr))).status_code == 422


async def test_csv_keeps_utf8_and_defuses_formulas(client, session, tenant, legal_entity):
    hr = await make_person(session, tenant, role=UserRole.ADMIN, name="Priya Hr")
    await _filled(session, tenant, legal_entity, hr)

    response = await client.post(
        URL, json={"file_type": "csv", "columns": ["code", "designation", "date_of_joining"]},
        headers=bearer(hr),
    )  # fmt: skip
    assert response.status_code == 200, response.text
    assert ".csv" in response.headers["content-disposition"]
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = {
        r["Employee code"]: r
        for r in csv.DictReader(io.StringIO(response.content.decode("utf-8-sig")))
    }
    assert rows["EMP0042"]["Designation"] == "'=HYPERLINK(1)"
    assert rows["EMP0042"]["Date of joining"] == "01-04-2024"

    # The old GET still works, filtered like the screen, with every column.
    old = await client.get(f"{URL}?search=Kavya", headers=bearer(hr))
    assert old.status_code == 200
    lines = old.content.decode("utf-8-sig").splitlines()
    assert len(lines) == 2
    assert "Account number" in lines[0] and "Family member 1 - Name" in lines[0]
    chosen = await client.get(f"{URL}?format=xlsx&columns=code,name", headers=bearer(hr))
    headers, _, _ = _sheet(chosen.content)
    assert headers == ["Employee code", "Name"]


async def test_an_unreadable_identifier_does_not_sink_the_file(
    client, session, tenant, legal_entity
):
    hr = await make_person(session, tenant, role=UserRole.ADMIN, name="Priya Hr")
    person = await _filled(session, tenant, legal_entity, hr)
    headers, person_id = bearer(hr), person.id
    junk = _FIELD_PREFIX + Fernet(Fernet.generate_key()).encrypt(b"ABCPV1234F").decode()
    await session.execute(
        text("update employees set pan = :junk, aadhaar_number = :junk where id = :id"),
        {"junk": junk, "id": person_id},
    )
    await session.execute(
        text(
            "update employee_bank_accounts set account_number = 'enc:v1:junk' "
            "where employee_id = :id"
        ),
        {"id": person_id},
    )
    await session.commit()
    session.expire_all()

    response = await client.post(URL, json={}, headers=headers)
    assert response.status_code == 200, response.text
    _, rows, _ = _sheet(response.content)
    assert rows["EMP0042"]["PAN"] == "Could not be read"
    assert rows["EMP0042"]["Aadhaar"] == "Could not be read"
    assert rows["EMP0042"]["Account number"] == "Could not be read"
    assert rows["EMP0042"]["Name"] == "Kavya Menon"


async def test_only_hr_with_the_export_right_and_only_their_own_people(
    client, session, tenant, other_tenant, legal_entity
):
    hr = await make_person(session, tenant, role=UserRole.ADMIN, name="Priya Hr")
    person = await _filled(session, tenant, legal_entity, hr)
    employee = await make_person(session, tenant, role=UserRole.EMPLOYEE, name="Esha Emp")
    assert (await client.post(URL, json={}, headers=bearer(employee))).status_code == 403
    assert (await client.get(f"{URL}/columns", headers=bearer(employee))).status_code == 403

    # A sub-admin: the Export right on People decides, not the Create right.
    exporter = await make_person(session, tenant, role=UserRole.SUB_ADMIN, name="Xena Export")
    viewer = await make_person(session, tenant, role=UserRole.SUB_ADMIN, name="Vik View")
    session.add_all(
        [
            UserAreaRights(
                tenant_id=tenant.id, user_id=exporter.id, area="people",
                can_create=False, can_edit=False, can_delete=False, can_export=True,
            ),
            UserAreaRights(
                tenant_id=tenant.id, user_id=viewer.id, area="people",
                can_create=True, can_edit=True, can_export=False,
            ),
        ]
    )  # fmt: skip
    await session.commit()
    assert (await client.post(URL, json={}, headers=bearer(exporter))).status_code == 200
    assert (await client.post(URL, json={}, headers=bearer(viewer))).status_code == 403
    assert (await client.get(f"{URL}/columns", headers=bearer(viewer))).status_code == 403

    # Another workspace's HR: neither our people nor our custom fields.
    other = await make_person(session, other_tenant, role=UserRole.ADMIN, name="Olga Other")
    columns = await client.get(f"{URL}/columns", headers=bearer(other))
    keys = {s["key"] for s in columns.json()["sections"]}
    assert "cf" not in keys and "ct.certs" not in keys
    body = {"employee_ids": [str(person.id)]}
    response = await client.post(URL, json=body, headers=bearer(other))
    assert response.status_code == 200
    _, rows, _ = _sheet(response.content)
    assert rows == {}
