"""A workbook laid out like factoHR's Employee Detail report, with made-up people.

The real export opens with a title, puts group headings on the two rows above
the header, merges them unevenly, and repeats header names under different
groups. This builds the same shape, so the importer is tested against it
without anybody's real data.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook

#: (header, group on the row above or None), in factoHR's column order.
COLUMNS: list[tuple[str, str | None]] = [
    ("S#", None),
    ("Emp Code", None),
    ("Full Name", None),
    ("First Name", None),
    ("Middle Name", None),
    ("Last Name", None),
    ("Title", None),
    ("Birth Date", None),
    ("Joining Date", None),
    ("Confirmation Date", None),
    ("Leaving Date", None),
    ("Employee Last working date", None),
    ("Gender", None),
    ("Status", None),
    ("Employment Type", None),
    ("Mobile No", None),
    ("Mobile No2", None),
    ("Email", None),
    ("Personal Email", None),
    ("Blood Group", None),
    ("Pan No", None),
    ("Nationality", None),
    ("Marital Status", None),
    ("Father Name", None),
    ("Mother Name", None),
    ("Spouse Name", None),
    ("Reporting Manager", None),
    ("Approving Manager", None),
    ("Address", "Local Address"),
    ("Address 2", "Local Address"),
    ("Country", "Local Address"),
    ("State", "Local Address"),
    ("City", "Local Address"),
    ("Pincode", "Local Address"),
    ("Permanent Same As Local Address", "Local Address"),
    ("Address", "Permanent Address"),
    ("Address 2", "Permanent Address"),
    ("Country", "Permanent Address"),
    # Outside the merged heading in the real export, as here.
    ("State", None),
    ("City", None),
    ("Pincode", None),
    ("Probation Period In Days", None),
    ("Notice Period For Employee", None),
    ("Attendance Cycle", "Category Types"),
    ("Branch", None),
    ("Companyy", None),
    ("Department", None),
    ("Designation", None),
    ("Grade", None),
    ("LOCATION", None),
    ("Bank Name", "Bank Detail"),
    ("Account No", None),
    ("IFSC Code", None),
    ("Name As Per Bank Records", None),
    ("PF Eligibility", "PF Detail"),
    ("PF No", None),
    ("Joining Date", None),
    ("UA No", None),
    ("ESIC Eligibility", "ESIC Detail"),
    ("ESIC No", None),
    ("Joining Date", None),
    ("Aadhaar No", "Aadhaar Card Number Details"),
    ("Name in Aadhaar", None),
]

#: Top-row headings, over the groups below them.
TOP_GROUPS = {"PF Eligibility": "PF / ESIC / EPS Detail", "Aadhaar No": "Employee Identity Detail"}


def build(people: list[dict[str, Any]], *, dates_as_serials: bool = True) -> bytes:
    """An .xlsx with one row per person; keys are header names (the first of a
    repeated name, or "Group > Name" for a grouped one)."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.cell(row=2, column=1, value="Employee Detail Report")
    group_start: dict[str, int] = {}
    for index, (name, group) in enumerate(COLUMNS, start=1):
        sheet.cell(row=6, column=index, value=name)
        if group and group not in group_start:
            group_start[group] = index
            sheet.cell(row=5, column=index, value=group)
        if name in TOP_GROUPS:
            sheet.cell(row=4, column=index, value=TOP_GROUPS[name])
    # The uneven merge: "Permanent Address" spans only Address to Country.
    start = group_start["Permanent Address"]
    sheet.merge_cells(start_row=5, start_column=start, end_row=5, end_column=start + 2)

    keys = []
    current_group = None
    for name, group in COLUMNS:
        current_group = group or current_group
        keys.append((name, current_group))
    for row_number, person in enumerate(people, start=7):
        sheet.cell(row=row_number, column=1, value=row_number - 6)
        seen: set[str] = set()
        for index, (name, group) in enumerate(keys, start=1):
            qualified = f"{group} > {name}" if group else name
            value = person.get(qualified)
            if value is None and name not in seen:
                value = person.get(name)
            seen.add(name)
            if isinstance(value, datetime) and dates_as_serials:
                value = (value - datetime(1899, 12, 30)).days
            if value is not None:
                sheet.cell(row=row_number, column=index, value=value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
